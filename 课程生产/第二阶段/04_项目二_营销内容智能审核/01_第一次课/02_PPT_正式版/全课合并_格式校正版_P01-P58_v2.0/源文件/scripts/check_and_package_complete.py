from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
from pypdf import PdfReader
from pptx import Presentation
from PIL import Image
import json,re,shutil,html,posixpath,hashlib
ROOT=Path('/Users/didi/Documents/Codex/2026-10-18/wo-l');p=ROOT/'work/ppt/lesson1_complete_review_ppt169_20261018';stage=p/'delivery';stage.mkdir(exist_ok=True)
base='第一次课_完整授课PPT_格式校正版_v2.0';ppt=p/'exports'/(base+'.pptx');pdf=p/'validation/render'/(base+'.pdf')
P='http://schemas.openxmlformats.org/presentationml/2006/main';A='http://schemas.openxmlformats.org/drawingml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';ns={'p':P,'a':A,'r':R}
parts=json.loads((p/'analysis/slide_order.json').read_text())
chapters={1,7,13,18,27,34,49};bounds=[(1,6,'开场与第一节·业务背景'),(7,12,'第二节·项目目标'),(13,17,'第三节·业务方案'),(18,26,'第四节·审核体系'),(27,33,'第五节·方案选择'),(34,48,'第六节·产品方案'),(49,58,'第七节·落地与验证')]
with ZipFile(ppt) as z:data={k:z.read(k) for k in z.namelist()}
if 'docProps/app.xml' in data:
 r=E.fromstring(data['docProps/app.xml'])
 for x in r:
  if E.QName(x).localname=='Slides':x.text='58'
 data['docProps/app.xml']=E.tostring(r,encoding='UTF-8',xml_declaration=True,standalone=True)
 with ZipFile(ppt,'w',ZIP_DEFLATED) as z:
  for k,v in data.items():z.writestr(k,v)
r=Presentation(ppt);assert len(r.slides)==58;assert all(s.has_notes_slide for s in r.slides)
checks=[];pictures=0;native=0;totalrels=0;outline=[]
for i,part in enumerate(parts,1):
 root=E.fromstring(data[part]);t=''.join(root.xpath('//a:t/text()',namespaces=ns));native+=len(root.xpath('//p:sp',namespaces=ns));pictures+=len(root.xpath('//p:pic',namespaces=ns))
 if i not in chapters:
  nav=[];page=None
  for sp in root.xpath('//p:sp',namespaces=ns):
   off=sp.find('p:spPr/a:xfrm/a:off',ns);s=''.join(sp.xpath('.//a:t/text()',namespaces=ns))
   if off is None or int(off.get('y'))<6286500:continue
   if re.match(r'^[1-7] ',s):nav.append(s)
   if int(off.get('x'))>1100*9525:page=s
  assert len(nav)==7,(i,nav)
  assert page==f'{i:02}',(i,page)
 if 39<=i<=43:
  assert not root.xpath('//p:pic',namespaces=ns),(i,'old screenshot remains')
  assert root.xpath('//p:cNvPr[contains(@name,"截图预留区")]',namespaces=ns)
 # Keep exact slide text and ordering for review/reproduction.
 titletexts=root.xpath('//a:t/text()',namespaces=ns)
 outline.append({'page':i,'heading':titletexts[:2],'text':titletexts})
for name,b in data.items():
 if not name.endswith('.rels'):continue
 own='' if name=='_rels/.rels' else posixpath.join(posixpath.dirname(posixpath.dirname(name)),posixpath.basename(name)[:-5])
 for rel in E.fromstring(b):
  if rel.get('TargetMode')=='External':continue
  t=rel.get('Target');dest=posixpath.normpath(posixpath.join(posixpath.dirname(own),t)) if not t.startswith('/') else t.lstrip('/')
  assert dest in data,(name,dest);totalrels+=1
pages=PdfReader(pdf).pages;assert len(pages)==58
for i,page in enumerate(pages,1):
 assert len(re.findall(r'[\u4e00-\u9fff]',page.extract_text()))>20,i
 assert abs(float(page.mediabox.width)/float(page.mediabox.height)-16/9)<.001
assert len([k for k in data if k.startswith('ppt/notesSlides/') and re.search(r'\.xml$',k)])==58
report={'slides':58,'sections':7,'notes':58,'size':'16:9','native_shapes':native,'pictures':pictures,'internal_relationships_verified':totalrels,'screenshot_blank_pages':[39,40,41,42,43],'all_content_pages_have_7_chapter_navigation':True,'continuous_page_numbers':True,'all_pdf_pages_have_readable_CJK':True,'visual_review':'全部58页逐页缩略图检查；P01、P03、P05、P07、P08、P09、P12、P40、P47等重点页复查。','fixed':'前12页配色与内容标题层级；全课内容页导航和页码；旧样稿重叠、页底遮挡和局部断行。','limitations':'预览通过本机字体替换渲染；其他电脑缺少主题字体时仍可能重排。历史分节版本保留，v2.0全课合并版为本次校正结果。'}
(p/'validation/format_review.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));(p/'analysis/slide_text.json').write_text(json.dumps(outline,ensure_ascii=False,indent=2))
shutil.copy2(ppt,stage/(base+'.pptx'));shutil.copy2(pdf,stage/(base+'_浏览版.pdf'))
review='''# 第一次课PPT格式检查说明\n\n本次交付：7节，共58页；含逐页教师备注。以“完整授课PPT·格式校正版v2.0”为本次统一授课文件，原分节版本保留备查。\n\n| 检查项 | 结果 |\n|---|---|\n| 页面比例 | 全部16:9，58页导出正常 |\n| 配色与字体 | 全课采用绿色模板配色和Microsoft YaHei主题字体，统一前12页内容标题层级 |\n| 结构 | 从第二节起保留章节标题页；内容页统一7项章节导航与当前节高亮 |\n| 页码与顺序 | P01—P58连续，无重复或缺页 |\n| 内容排版 | 已检查58页实际渲染画面，修复P03、P05文字重叠及多处底部遮挡、断行问题 |\n| 截图区 | P39—P43保留5处空白截图框，旧图已清除 |\n| 教师备注 | 全部58页保留备注 |\n| 可编辑与完整性 | 原生文本及图形可编辑；内部关系与媒体引用检查通过 |\n| 教学数据 | 第七节200→205分钟明确标为课堂计算示例 |\n\n使用时直接在P39—P43粘贴截图，可覆盖或删除虚线框。预览使用本机中文字体替换渲染，正式PPT保留Microsoft YaHei主题字体；更换授课电脑后建议先打开确认字体。\n\n## 章节页码\n\n'''+''.join(f'- {name}：P{a:02}—P{b:02}\n' for a,b,name in bounds)
(stage/(base+'_格式检查说明.md')).write_text(review)
imgs=stage/'页面预览';imgs.mkdir(exist_ok=True)
for i in range(1,59):shutil.copy2(p/f'validation/render/page-{i:02}.png',imgs/f'P{i:02}.png')
for start in range(0,58,12):
 n=min(12,58-start);o=Image.new('RGB',(1936,((n+2)//3)*376-16),'#D3DFDA')
 for j in range(n):o.paste(Image.open(p/f'validation/render/page-{start+j+1:02}.png').resize((640,360)),((j%3)*648,(j//3)*376))
 o.save(stage/f'全课页面总览_P{start+1:02}-P{start+n:02}.png')
sections=''.join('<section><h2>'+name+'</h2>'+''.join(f'<figure id="p{i:02}"><figcaption>P{i:02} · {html.escape(outline[i-1]["heading"][1] if len(outline[i-1]["heading"])>1 else name)}</figcaption><img loading="lazy" src="页面预览/P{i:02}.png" alt="P{i:02}"></figure>' for i in range(a,b+1))+'</section>' for a,b,name in bounds)
nav=''.join(f'<a href="#p{a:02}">{name}</a>' for a,b,name in bounds)
page='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>第一次课·完整授课PPT·格式校正版v2.0</title><style>body{margin:0;background:#edf2ef;color:#17302f;font:18px/1.7 -apple-system,BlinkMacSystemFont,"PingFang SC",sans-serif}main{max-width:1280px;margin:32px auto;padding:0 24px}h1{font-size:32px}h2{font-size:28px;padding-top:30px}nav{display:flex;flex-wrap:wrap;gap:18px;margin:20px 0}a{color:#126b63}figure{margin:28px 0 42px;scroll-margin-top:30px}figcaption{font-size:23px;font-weight:600;margin-bottom:12px}img{display:block;width:100%;height:auto;box-shadow:0 8px 26px #17302f15}</style><main><h1>第一次课 · 完整授课PPT</h1><p>格式校正版v2.0 · 7节58页 · 可编辑文字与图形 · 含逐页教师备注。P39—P43为空白截图区域。</p><nav><a href="'''+base+'.pptx">下载可编辑PPT</a><a href="'+base+'_浏览版.pdf">浏览版PDF</a><a href="'+base+'_格式检查说明.md">格式检查说明</a></nav><nav>'+nav+'</nav>'+sections+'</main></html>'
(stage/(base+'_逐页预览.html')).write_text(page)
(p/'analysis/delivery_meta.json').write_text(json.dumps({'base':base,'stage':str(stage),'folder':'全课合并_格式校正版_P01-P58_v2.0','pages':58},ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
