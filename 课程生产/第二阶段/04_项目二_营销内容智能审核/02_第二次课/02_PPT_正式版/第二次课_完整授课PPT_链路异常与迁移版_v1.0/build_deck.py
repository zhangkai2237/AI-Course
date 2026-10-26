#!/usr/bin/env python3
"""Build the complete 46-slide second-lesson deck.

Build shape: solo — subagent dispatch not authorized in this run.
Visual lock: follows the first-lesson deck (deep teal, warm off-white, restrained coral/gold).
Motion: static throughout; user opted out of click builds.
"""
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[6]
SKILL = Path("/Users/keivn/.codex/skills/slide-maker")
sys.path.insert(0, str(SKILL / "scripts"))

import deckkit as dk  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR  # noqa: E402

OUTDIR = Path(__file__).resolve().parent
OUT = OUTDIR / "第二次课_完整授课PPT_链路异常与迁移版_v1.0.pptx"
IMG = ROOT / "课程生产/第二阶段/04_项目二_营销内容智能审核/01_第一次课/02_PPT_正式版/源文件/images/prototype-selfcheck.jpg"

# Palette locked to lesson one.
BG = RGBColor(0xF7, 0xF8, 0xF4)
PALE = RGBColor(0xE8, 0xF2, 0xEF)
TEAL = RGBColor(0x12, 0x6B, 0x63)
TEAL2 = RGBColor(0x18, 0x7A, 0x70)
CORAL = RGBColor(0xD8, 0x5C, 0x3D)
GOLD = RGBColor(0xDB, 0xAB, 0x4A)
INK = RGBColor(0x17, 0x30, 0x2F)
MUTE = RGBColor(0x61, 0x76, 0x73)
GRID = RGBColor(0xD3, 0xDF, 0xDA)
DARK = RGBColor(0x10, 0x2F, 0x30)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT_CORAL = RGBColor(0xF8, 0xE8, 0xE2)
SOFT_GOLD = RGBColor(0xF6, 0xEF, 0xDA)

dk.set_palette(deep=DARK, blue=TEAL, teal=TEAL2, magenta=CORAL, slate=INK,
               mute=MUTE, mono="Arial", font="Arial", display="Arial",
               eafont="Noto Sans CJK SC", eadisplay="Noto Sans CJK SC",
               accents=[TEAL, CORAL, GOLD, MUTE])
dk.set_ground(BG)

W, H = 10.0, 5.625
SECTIONS = ["1  全链路", "2  异常处理", "3  场景迁移"]


def tx(s, x, y, w, h, value, size=17, color=INK, bold=False,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Noto Sans CJK SC"):
    return dk.text(s, x, y, w, h, [[(value, size, color, bold, False, font)]],
                   align=align, anchor=anchor, space_after=0, line_spacing=1.0)


def paras(s, x, y, w, h, values, size=15, color=INK, gap=8, bold_first=False):
    runs = []
    for i, v in enumerate(values):
        runs.append([(v, size, color, bold_first and i == 0, False, "Noto Sans CJK SC")])
    return dk.text(s, x, y, w, h, runs, space_after=gap, line_spacing=1.0)


def nav(s, page, sec=1, tag="第二次课 · 营销内容智能审核"):
    dk.box(s, 0, 5.27, W, 0.01, fill=GRID, line=GRID)
    for i, (label, x) in enumerate(zip(SECTIONS, [0.38, 2.00, 3.62]), start=1):
        c = TEAL if i == sec else MUTE
        tx(s, x, 5.31, 1.55, 0.18, label, 7.8, c, i == sec, anchor=MSO_ANCHOR.MIDDLE)
    tx(s, 7.45, 5.31, 1.75, 0.18, tag, 7.2, MUTE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    tx(s, 9.48, 5.31, 0.18, 0.18, str(page), 8, MUTE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def base(prs, title, kicker, page, sec, notes, accent=TEAL):
    s = dk.add_slide(prs)
    dk.box(s, 0, 0, W, H, fill=BG, line=BG)
    tx(s, 0.4, 0.27, 8.6, 0.22, kicker, 9.5, accent)
    tx(s, 0.4, 0.56, 9.1, 0.55, title, 25, INK, True)
    dk.register_mark(s, "arcs", corner="tr", color=PALE, size=0.75)
    nav(s, page, sec)
    dk.speaker_notes(s, notes)
    return s


def divider(prs, num, title, question, agenda, page, sec, notes):
    s = dk.add_slide(prs)
    dk.box(s, 0, 0, W, H, fill=DARK, line=DARK)
    tx(s, 0.4, 0.38, 4.8, 0.24, f"第二次课 / 第{num}节", 10, GOLD)
    tx(s, 0.4, 0.86, 4.6, 0.62, title, 34, WHITE, True)
    tx(s, 0.4, 1.77, 3.7, 1.05, question, 24, WHITE, True)
    dk.box(s, 4.45, 1.7, 0.015, 1.95, fill=MUTE, line=MUTE)
    for item, y in zip(agenda, [1.82, 2.51, 3.20]):
        tx(s, 4.72, y, 4.55, 0.42, item, 15, WHITE)
    dk.box(s, 0.4, 4.28, 9.15, 0.68, fill=TEAL2, line=TEAL2, round=True)
    tx(s, 0.62, 4.43, 8.7, 0.32, "本节重点  ·  " + agenda[-1], 13.5, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    tx(s, 0.4, 5.25, 2.0, 0.18, "链路 → 判断 → 迁移", 7.5, GRID)
    tx(s, 9.45, 5.25, 0.2, 0.18, str(page), 8, GRID, align=PP_ALIGN.RIGHT)
    dk.speaker_notes(s, notes)
    return s


def card(s, x, y, w, h, title, lines=(), fill=WHITE, title_c=TEAL, hub=False, num=None):
    f = TEAL2 if hub else fill
    ink = WHITE if hub else INK
    dk.box(s, x, y, w, h, fill=f, line=f if hub else GRID, round=True)
    if num is not None:
        dk.box(s, x + 0.16, y + 0.15, 0.34, 0.34, fill=GOLD if hub else PALE, line=GOLD if hub else PALE, round=True)
        tx(s, x + 0.16, y + 0.145, 0.34, 0.34, str(num), 10, DARK, True,
           PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, "Arial")
        title_x = x + 0.60
        title_w = w - 0.76
    else:
        title_x = x + 0.18
        title_w = w - 0.36
    compact = h < 1.0
    if not lines:
        tx(s, title_x, y + 0.04, title_w, h - 0.08, title, 12 if compact else 15,
           WHITE if hub else title_c, True, anchor=MSO_ANCHOR.MIDDLE)
    else:
        tx(s, title_x, y + (0.08 if compact else 0.17), title_w, 0.28 if compact else 0.34,
           title, 12 if compact else 15, WHITE if hub else title_c, True)
        paras(s, x + 0.18, y + (0.39 if compact else 0.63), w - 0.36,
              max(0.24, h - (0.48 if compact else 0.75)), list(lines),
              9.5 if compact else 11.5, ink, 3 if compact else 4)
    return (x, y, w, h)


def takeaway(s, text, color=TEAL):
    dk.box(s, 0.4, 4.59, 9.15, 0.48, fill=PALE if color == TEAL else SOFT_CORAL, line=PALE if color == TEAL else SOFT_CORAL, round=True)
    tx(s, 0.58, 4.70, 8.8, 0.24, text, 12.5, color, True, anchor=MSO_ANCHOR.MIDDLE)


def flow(s, labels, subs=None, y=2.0, hub=-1, colors=None):
    n = len(labels)
    gap = 0.28
    w = (9.1 - gap * (n - 1)) / n
    rects = []
    for i, label in enumerate(labels):
        x = 0.45 + i * (w + gap)
        if i < n - 1:
            dk.arrow(s, x + w + 0.05, y + 0.48, gap - 0.10, 0.20, color=TEAL)
        fill = colors[i] if colors else (TEAL2 if i == hub else PALE)
        is_hub = i == hub
        rects.append(card(s, x, y, w, 1.45, label,
                          (subs[i],) if subs else (), fill=fill, hub=is_hub, num=i + 1))
    return rects


def slide_01(prs):
    """role=cover | form=statement + audit-trail motif | static: user opted out | takeaway='从完整审核链路到判断型AI产品迁移'"""
    s = dk.add_slide(prs)
    dk.box(s, 0, 0, W, H, fill=BG, line=BG)
    dk.box(s, 0, 0, 0.16, H, fill=TEAL, line=TEAL)
    tx(s, 0.62, 0.55, 4.0, 0.22, "第二阶段 · 项目二 · 第二次课", 10, TEAL)
    tx(s, 0.62, 1.25, 8.7, 1.15, "从完整审核链路\n到判断型 AI 产品迁移", 34, DARK, True)
    tx(s, 0.64, 2.82, 7.8, 0.45, "页面连起来 · 异常接得住 · 方法迁得走", 18, MUTE)
    # Signature trail
    xs = [0.82, 3.15, 5.48, 7.81]
    labels = ["版本", "任务", "判断", "处理"]
    for i, (x, lab) in enumerate(zip(xs, labels)):
        if i < 3:
            dk.box(s, x + 0.38, 4.17, 1.88, 0.035, fill=TEAL, line=TEAL)
        dk.box(s, x, 3.90, 0.58, 0.58, fill=TEAL if i < 3 else CORAL, line=TEAL if i < 3 else CORAL, round=True)
        tx(s, x, 3.91, 0.58, 0.55, str(i + 1), 14, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, "Arial")
        tx(s, x - 0.15, 4.56, 0.88, 0.25, lab, 11, INK, True, PP_ALIGN.CENTER)
    tx(s, 0.62, 5.15, 4.0, 0.18, "180分钟 · 项目深化课", 8.5, MUTE)
    dk.speaker_notes(s, "今天不重新讲项目背景，而是检验产品能否真正跑起来。全课围绕三件事：链路、异常、迁移。")
    return s


def slide_02(prs):
    """role=map | form=three-stage route | static: user opted out | takeaway='三段课最终汇成一页判断型AI产品方案'"""
    s = base(prs, "今天要完成的不是三块知识，而是一条进阶路线", "全课地图", 2, 1,
             "先从已有产品出发，检查运行链路；再进入异常决策；最后把方法迁移到新业务。")
    flow(s, ["链路跑通", "异常接住", "方法迁移"],
         ["对象、状态、角色", "分流、权限、兜底", "合同 → 客服 → 新场景"], y=1.72, hub=1)
    takeaway(s, "课堂终点：独立写出一页判断型 AI 产品设计。")
    return s


def slide_03(prs):
    """role=hook | form=dependency statement | static: user opted out | takeaway='通过不是一个孤立标签'"""
    s = base(prs, "“审核通过”只在三个条件同时成立时才有效", "开场判断", 3, 1,
             "让学员先判断：页面显示通过，是否就能发布？强调结论必须绑定版本、资料和角色。")
    card(s, 0.55, 1.55, 2.45, 2.25, "对应版本", ["结论属于哪一稿", "是否已被新稿替代"], fill=PALE, num="V")
    card(s, 3.28, 1.55, 2.45, 2.25, "材料有效", ["引用是否可用", "条件是否仍成立"], fill=WHITE, num="D")
    card(s, 6.01, 1.55, 2.45, 2.25, "角色有权", ["谁可以确认", "谁承担下一步"], fill=PALE, num="R")
    tx(s, 8.72, 2.13, 0.6, 0.6, "＝", 28, TEAL, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    dk.box(s, 9.24, 1.78, 0.36, 1.75, fill=TEAL2, line=TEAL2, round=True)
    tx(s, 9.24, 1.95, 0.36, 1.35, "可\n执\n行", 11.5, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    takeaway(s, "没有版本、材料与角色约束，“通过”只是页面文案。")
    return s


def slide_04(prs):
    """role=divider | form=dark divider | static: user opted out | takeaway='页面动作接成业务链'"""
    return divider(prs, "一", "完整链路", "页面怎么从“能点”\n变成“能跑”？",
                   ["四类业务对象", "三段演示链路", "状态、版本与责任人"], 4, 1,
                   "第一节从页面出发，但不止讲页面。要把动作、对象、状态和承接角色完整连接。")


def slide_05(prs):
    """role=concept | form=before/after comparison | static: user opted out | takeaway='页面集合不等于业务链'"""
    s = base(prs, "几张页面放在一起，不等于产品已经能运行", "第一节 · 业务链", 5, 1,
             "左边是常见误区：只列页面。右边是产品链路：每个动作都产生对象、改变状态并交给下一位角色。")
    card(s, 0.55, 1.48, 3.75, 2.75, "页面集合", ["编辑页", "任务列表", "审核详情", "巡检批次"], fill=WHITE)
    tx(s, 4.42, 2.45, 1.0, 0.5, "≠", 36, CORAL, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    card(s, 5.55, 1.48, 3.9, 2.75, "业务链", ["动作改变对象", "对象进入新状态", "状态触发下一角色", "结果绑定实际版本"], fill=PALE, title_c=TEAL)
    takeaway(s, "真正的链路，要回答：发生了什么变化，下一步交给谁。")
    return s


def slide_06(prs):
    """role=framework | form=concept equation | static: user opted out | takeaway='每个页面动作都要回答三件事'"""
    s = base(prs, "一个可运行的页面动作，要同时改变三件事", "第一节 · 动作定义", 6, 1,
             "讲清动作不是按钮名，而是业务变化。用保存新稿举例：生成新版本、改变状态、把后续责任交给作者或审核员。")
    tx(s, 0.7, 1.65, 2.2, 0.65, "对象", 31, TEAL, True, PP_ALIGN.CENTER)
    tx(s, 3.05, 1.65, 0.4, 0.65, "+", 30, GOLD, True, PP_ALIGN.CENTER)
    tx(s, 3.6, 1.65, 2.2, 0.65, "状态", 31, TEAL, True, PP_ALIGN.CENTER)
    tx(s, 5.95, 1.65, 0.4, 0.65, "+", 30, GOLD, True, PP_ALIGN.CENTER)
    tx(s, 6.5, 1.65, 2.2, 0.65, "责任人", 31, TEAL, True, PP_ALIGN.CENTER)
    dk.box(s, 0.72, 2.58, 7.95, 0.06, fill=GRID, line=GRID)
    tx(s, 0.75, 2.94, 7.9, 0.6, "谁被创建或修改 · 现在处于什么状态 · 谁要继续处理", 16, INK, True, PP_ALIGN.CENTER)
    takeaway(s, "按钮只描述交互；对象、状态、责任人才描述业务。")
    return s


def slide_07(prs):
    """role=framework | form=four-object grid | static: user opted out | takeaway='四类对象承担不同业务职责'"""
    s = base(prs, "审核链路中，至少要区分四类业务对象", "第一节 · 四类对象", 7, 1,
             "四类对象不能混成一个结果字段。内容版本承载正文；检查记录承载一次机器执行；正式任务承载业务决定；巡检任务承载上线后的维护。")
    card(s, 0.45, 1.48, 2.12, 2.66, "内容版本", ["V1 / V2 / V3", "正文与业务上下文", "被审核的真实对象"], fill=PALE, num=1)
    card(s, 2.75, 1.48, 2.12, 2.66, "检查记录", ["一次执行", "原句、结论、依据", "失败也要留状态"], fill=WHITE, num=2)
    card(s, 5.05, 1.48, 2.12, 2.66, "正式任务", ["提交版本", "复核意见", "正式决定"], fill=PALE, num=3)
    card(s, 7.35, 1.48, 2.12, 2.66, "巡检任务", ["在用内容", "变化事件", "实际更新"], fill=WHITE, num=4)
    takeaway(s, "对象分开，版本、权限与历史才不会互相污染。")
    return s


def slide_08(prs):
    """role=concept | form=layered state map | static: user opted out | takeaway='不同对象有自己的状态机'"""
    s = base(prs, "同一篇内容，可以同时处在四种不同状态里", "第一节 · 状态并存", 8, 1,
             "状态必须挂在正确对象上。内容版本已保存，不代表检查已完成；检查完成，不代表正式审核通过；审核通过，也不代表上线内容永远有效。")
    rows = [
        ("内容版本", "V2 已保存", TEAL, 1.35),
        ("检查记录", "运行中 / 已完成 / 失败", GOLD, 2.18),
        ("正式任务", "待复核 / 退回 / 通过", CORAL, 3.01),
        ("巡检任务", "待分派 / 处理中 / 已更新", TEAL2, 3.84),
    ]
    for name, states, c, y in rows:
        tx(s, 0.55, y + 0.14, 1.45, 0.34, name, 13, INK, True)
        dk.box(s, 2.05, y, 7.2, 0.62, fill=WHITE, line=GRID, round=True)
        dk.box(s, 2.05, y, 1.05, 0.62, fill=c, line=c, round=True)
        tx(s, 3.30, y + 0.13, 5.6, 0.35, states, 14, INK, True)
    takeaway(s, "不要问“这篇内容是什么状态”，先问“哪个对象是什么状态”。")
    return s


def slide_09(prs):
    """role=demo | form=product screenshot + callouts | static: user opted out | takeaway='自检服务于提交前修改'"""
    s = base(prs, "课堂演示一：从问题稿到自检修改", "第一节 · Demo 01", 9, 1,
             "演示时依次展示：固定当前版本与上下文、发起检查、查看原句与建议、修改为新版本。强调自检是作者工作入口。")
    dk.box(s, 0.42, 1.30, 6.35, 3.12, fill=WHITE, line=GRID, round=True)
    dk.picture(s, str(IMG), 0.55, 1.43, 6.08, 2.86, fit="contain", alt="营销内容自检原型页面")
    card(s, 7.05, 1.45, 2.35, 0.76, "① 固定输入", ["版本 + 上下文"], fill=PALE)
    card(s, 7.05, 2.34, 2.35, 0.76, "② 执行检查", ["规则与模型"], fill=WHITE)
    card(s, 7.05, 3.23, 2.35, 0.76, "③ 保存新稿", ["V1 → V2"], fill=PALE)
    takeaway(s, "自检的完成条件不是“看过结果”，而是“问题被修改或明确保留”。")
    return s


def slide_10(prs):
    """role=decision | form=branching interpretation | static: user opted out | takeaway='未发现不等于正式通过'"""
    s = base(prs, "机器“未发现问题”，不等于业务“正式通过”", "第一节 · 结论边界", 10, 1,
             "讲清两个结论的责任主体不同。机器输出表示在当前输入和能力范围内未发现；正式通过是有权限角色做出的业务决定。")
    card(s, 0.65, 1.55, 3.55, 2.42, "未发现问题", ["当前输入", "当前检查范围", "当前系统能力"], fill=PALE, title_c=TEAL)
    tx(s, 4.45, 2.28, 1.0, 0.5, "≠", 38, CORAL, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    card(s, 5.72, 1.55, 3.55, 2.42, "正式审核通过", ["有权角色确认", "责任边界明确", "允许进入下一业务环节"], fill=TEAL2, hub=True)
    takeaway(s, "系统给判断材料；业务角色作正式决定。")
    return s


def slide_11(prs):
    """role=demo | form=review workflow | static: user opted out | takeaway='正式审核围绕提交版本流转'"""
    s = base(prs, "课堂演示二：正式提交、人工复核与再修改", "第一节 · Demo 02", 11, 1,
             "从作者提交V2开始演示。审核员核对原句、依据与上下文，退回后作者形成V3，再由审核员复核并作正式决定。")
    flow(s, ["提交 V2", "审核判断", "退回修改", "复核决定"],
         ["正文与材料", "核对原句与来源", "形成 V3", "结论绑定 V3"], y=1.72, hub=1)
    takeaway(s, "正式审核围绕提交版本流转，不能只围绕一条机器结果。")
    return s


def slide_12(prs):
    """role=comparison | form=two-column comparison | static: user opted out | takeaway='自检与正式审核不是同一功能换皮'"""
    s = base(prs, "自检与正式审核，目标、角色和产出都不同", "第一节 · 功能分工", 12, 1,
             "让学员不要把两个流程做成同一页面改个标题。自检服务于作者改稿；正式审核服务于有权限角色作决定。")
    card(s, 0.55, 1.38, 4.25, 2.95, "提交前自检", ["角色：作者", "目标：尽早发现并修改", "产出：检查记录 + 新版本", "可反复执行"], fill=PALE)
    card(s, 5.18, 1.38, 4.25, 2.95, "正式审核", ["角色：审核员", "目标：作业务决定", "产出：审核任务 + 正式结论", "需要权限与责任"], fill=WHITE, title_c=CORAL)
    takeaway(s, "前者帮助准备；后者完成交接与决定。")
    return s


def slide_13(prs):
    """role=versioning | form=version timeline | static: user opted out | takeaway='V1结果不能自动覆盖V2'"""
    s = base(prs, "V1 检查通过，也不能自动覆盖后来的 V2", "第一节 · 版本有效性", 13, 1,
             "用时间线讲清旧结果迟到。V1发起检查后，作者保存V2；V1结果返回时只能归档到V1，V2仍需按规则重新检查。")
    xs = [0.75, 3.15, 5.55, 7.95]
    labs = [("发起 V1 检查", TEAL), ("保存 V2 新稿", CORAL), ("V1 结果返回", GOLD), ("V2 重新检查", TEAL2)]
    dk.box(s, 0.95, 2.34, 7.5, 0.04, fill=GRID, line=GRID)
    for i, (lab, c) in enumerate(labs):
        dk.box(s, xs[i], 2.08, 0.56, 0.56, fill=c, line=c, round=True)
        tx(s, xs[i]-0.38, 1.48, 1.32, 0.42, lab, 12, INK, True, PP_ALIGN.CENTER)
    tx(s, 0.75, 3.05, 8.1, 0.55, "归属规则：检查记录永远绑定发起时的版本与上下文", 17, TEAL, True, PP_ALIGN.CENTER)
    takeaway(s, "旧结果可以保留，但不能冒充新版本结论。", CORAL)
    return s


def slide_14(prs):
    """role=demo | form=inspection trigger chain | static: user opted out | takeaway='巡检从业务变化开始'"""
    s = base(prs, "课堂演示三：巡检从变化事件开始", "第一节 · Demo 03", 14, 1,
             "演示资料失效、规则更新或活动条件变化后，系统怎样圈定受影响内容、生成批次、分派任务并跟踪处理。")
    flow(s, ["业务变化", "圈定范围", "建立批次", "分派处理"],
         ["资料 / 规则 / 条件", "在用内容", "问题与待办", "负责人承接"], y=1.72, hub=2)
    takeaway(s, "巡检不是重新跑全库，而是从变化反推受影响范围。")
    return s


def slide_15(prs):
    """role=signature | form=end-to-end closure | static: user opted out | takeaway='巡检必须走到真实渠道更新'"""
    s = base(prs, "巡检只有走到真实渠道更新，业务才真正闭环", "第一节 · 闭环", 15, 1,
             "这是第一节的关键页。强调发现问题不等于完成：必须明确在用位置、处理人、更新动作与复核结果。")
    # signature audit trail
    stages = [("变化发生", "资料失效"), ("范围圈定", "找到在用内容"), ("任务处理", "修改并复核"), ("实际更新", "渠道内容生效")]
    for i, ((a, b), x) in enumerate(zip(stages, [0.62, 2.90, 5.18, 7.46])):
        if i < 3:
            dk.box(s, x + 0.58, 2.28, 1.58, 0.045, fill=TEAL, line=TEAL)
        dk.box(s, x, 1.96, 0.66, 0.66, fill=CORAL if i == 3 else TEAL, line=CORAL if i == 3 else TEAL, round=True)
        tx(s, x, 1.97, 0.66, 0.63, str(i+1), 15, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, "Arial")
        tx(s, x-0.12, 2.83, 1.42, 0.3, a, 13, INK, True, PP_ALIGN.CENTER)
        tx(s, x-0.28, 3.24, 1.74, 0.45, b, 10.5, MUTE, False, PP_ALIGN.CENTER)
    takeaway(s, "业务完成点不在“发现”，而在“实际更新并可追溯”。", CORAL)
    return s


def slide_16(prs):
    """role=reference | form=action-state-owner table | static: user opted out | takeaway='动作必须对应状态和责任人'"""
    s = base(prs, "把关键页面动作整理成一张运行表", "第一节 · 状态表", 16, 1,
             "这张表用于检查原型是否可运行。每个动作都要写清对象、状态变化与下一责任人。")
    rows = [
        ["页面动作", "作用对象", "状态变化", "下一责任人"],
        ["保存新稿", "内容版本", "草稿 V1 → V2", "作者"],
        ["发起自检", "检查记录", "未开始 → 运行中", "系统 / 作者"],
        ["提交正式审核", "正式任务", "草稿 → 待复核", "审核员"],
        ["创建巡检批次", "巡检任务", "待创建 → 待分派", "内容负责人"],
        ["确认更新", "巡检任务", "处理中 → 已完成", "负责人 / 审核员"],
    ]
    dk.table(s, 0.48, 1.30, 9.05, rows, col_w=[1.9, 1.75, 2.75, 2.65], size=11.5, row_h=0.52,
             head_c=DARK, body_c=INK, rule_c=GRID, hi_fill=PALE, hi_c=TEAL)
    return s


def slide_17(prs):
    """role=exercise | form=scenario board | static: user opted out | takeaway='先保护版本，再决定重跑'"""
    s = base(prs, "小组练习：V3 已保存，V2 结果刚刚返回", "第一节 · 练习", 17, 1,
             "给两分钟讨论。参考处理：V2结果归档到V2；页面明确它不是当前版本；V3根据改动范围决定重跑全部或受影响项。")
    card(s, 0.55, 1.45, 2.6, 2.75, "现场条件", ["V2 已发起检查", "作者随后保存 V3", "V2 结果迟到"], fill=PALE)
    card(s, 3.45, 1.45, 2.6, 2.75, "请回答", ["结果显示在哪里？", "能否更新 V3 状态？", "需要全部重跑吗？"], fill=WHITE, title_c=CORAL)
    card(s, 6.35, 1.45, 3.05, 2.75, "判断顺序", ["先看归属版本", "再看有效范围", "最后决定重跑策略"], fill=TEAL2, hub=True)
    takeaway(s, "先保护版本边界，再讨论效率优化。")
    return s


def slide_18(prs):
    """role=recap | form=three-question recap | static: user opted out | takeaway='页面动作翻译成对象状态责任人'"""
    s = base(prs, "第一节只需要记住三个追问", "第一节 · 小结", 18, 1,
             "用三个问题收束：动作改了什么对象？对象进入什么状态？下一步由谁承接？如果答不出来，链路仍然断着。")
    card(s, 0.65, 1.55, 2.65, 2.35, "改了什么对象？", ["版本 / 记录 / 任务"], fill=PALE, num=1)
    card(s, 3.68, 1.55, 2.65, 2.35, "进入什么状态？", ["运行 / 待办 / 完成"], fill=WHITE, num=2)
    card(s, 6.70, 1.55, 2.65, 2.35, "交给谁继续？", ["作者 / 审核员 / 负责人"], fill=PALE, num=3)
    takeaway(s, "能回答这三问，页面才开始具备业务含义。")
    return s


def slide_19(prs):
    """role=break | form=dark pause card | static: user opted out | takeaway='休息十分钟'"""
    s = dk.add_slide(prs)
    dk.box(s, 0, 0, W, H, fill=DARK, line=DARK)
    tx(s, 0.7, 1.35, 8.6, 0.9, "休息 10 分钟", 40, WHITE, True, PP_ALIGN.CENTER)
    tx(s, 0.7, 2.55, 8.6, 0.45, "回来以后，我们专门处理“正常流程之外”的产品决策。", 17, GRID, False, PP_ALIGN.CENTER)
    dk.box(s, 3.95, 3.55, 2.1, 0.58, fill=TEAL2, line=TEAL2, round=True)
    tx(s, 3.95, 3.56, 2.1, 0.56, "65 分钟继续", 13, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    dk.speaker_notes(s, "休息十分钟。回来进入异常处理部分。")
    return s


def slide_20(prs):
    """role=divider | form=dark divider | static: user opted out | takeaway='异常是下一步分流设计'"""
    return divider(prs, "二", "异常决策", "系统出问题时，\n下一步到底怎么走？",
                   ["八部分异常框架", "五类真实异常", "继续、降级、人工与停止"], 20, 2,
                   "第二节不讲报错文案，而是讲系统遇到不确定性时怎样分流、怎样交接、怎样保护业务。")


def slide_21(prs):
    """role=complex | form=eight-part exception framework | static: user opted out | takeaway='异常设计要覆盖八个问题'"""
    s = base(prs, "一个真实异常，至少要回答八个问题", "第二节 · 异常框架", 21, 2,
             "这八项是课堂统一分析框架：触发、影响、可见、状态、动作、责任、记录、恢复。后续五个案例都按它分析。")
    items = ["触发条件", "影响范围", "用户可见", "系统状态", "可选动作", "承接角色", "保留材料", "恢复条件"]
    coords = [(0.55,1.35),(2.85,1.35),(5.15,1.35),(7.45,1.35),(0.55,2.78),(2.85,2.78),(5.15,2.78),(7.45,2.78)]
    for i, (lab, (x,y)) in enumerate(zip(items, coords), start=1):
        card(s, x, y, 2.0, 1.08, lab, (), fill=PALE if i%2 else WHITE, num=i)
    takeaway(s, "异常不是一句提示，而是一条新的业务分支。")
    return s


def scenario(s, situation, wrong, decision, owner):
    card(s, 0.5, 1.35, 2.05, 2.83, "发生了什么", situation, fill=PALE)
    card(s, 2.78, 1.35, 2.05, 2.83, "危险做法", wrong, fill=SOFT_CORAL, title_c=CORAL)
    card(s, 5.06, 1.35, 2.05, 2.83, "产品处理", decision, fill=WHITE, title_c=TEAL)
    card(s, 7.34, 1.35, 2.16, 2.83, "谁来承接", owner, fill=TEAL2, hub=True)


def slide_22(prs):
    """role=scenario | form=four-part decision board | static: user opted out | takeaway='材料缺失时先限定结论'"""
    s = base(prs, "异常一：业务材料缺失，先限定结论范围", "第二节 · 异常 01", 22, 2,
             "缺资料时不应编造完整结论。系统要明确缺什么、影响哪些检查项，并把补充或确认动作交给对应角色。")
    scenario(s, ["商品资料过期", "活动规则缺页"], ["仍输出完整通过", "只弹“检查失败”"],
             ["标出缺项", "受影响项不作结论", "其余项可继续"], ["作者补充", "资料负责人确认"])
    takeaway(s, "信息不足时，先说明影响范围，再请求补充或确认。", CORAL)
    return s


def slide_23(prs):
    """role=scenario | form=partial-failure route | static: user opted out | takeaway='局部失败不应抹掉已完成部分'"""
    s = base(prs, "异常二：超时或局部失败，要保留已完成部分", "第二节 · 异常 02", 23, 2,
             "模型超时不代表整次任务毫无价值。可保留已完成项，对失败项标记状态，并提供重试、降级或人工补核。")
    flow(s, ["发起检查", "3 项完成", "1 项超时", "分流处理"],
         ["固定输入", "保留可用结果", "单项失败", "重试 / 降级 / 人工"], y=1.62, hub=3,
         colors=[PALE, PALE, SOFT_CORAL, TEAL2])
    takeaway(s, "失败要细化到检查项，不能把整次任务一键抹掉。", CORAL)
    return s


def slide_24(prs):
    """role=scenario | form=version collision timeline | static: user opted out | takeaway='迟到结果只回到原版本'"""
    s = base(prs, "异常三：旧版本结果迟到，只能回到原版本", "第二节 · 异常 03", 24, 2,
             "这页把第一节练习升级为异常策略。页面要显示版本标签和过期提示，状态层不能把旧结果写到当前稿。")
    card(s, 0.6, 1.42, 2.2, 2.85, "T1 · V2", ["发起检查", "输入已冻结"], fill=PALE)
    card(s, 3.18, 1.42, 2.2, 2.85, "T2 · V3", ["作者保存新稿", "当前版本改变"], fill=SOFT_CORAL, title_c=CORAL)
    card(s, 5.76, 1.42, 3.62, 2.85, "T3 · V2 结果返回", ["归档到 V2", "标注“非当前版本”", "V3 按影响范围重跑"], fill=WHITE, title_c=TEAL)
    takeaway(s, "归属先于展示：结果写回哪个对象，比显示在哪里更重要。")
    return s


def slide_25(prs):
    """role=scenario | form=conflict matrix | static: user opted out | takeaway='规则与模型冲突要显式分级'"""
    s = base(prs, "异常四：规则与模型冲突，不能静默取一个答案", "第二节 · 异常 04", 25, 2,
             "规则适合明确条件，模型适合理解表达。冲突时先区分规则级别，再决定是否阻断、提示或转人工。")
    rows = [
        ["规则级别", "模型判断", "产品动作", "结果呈现"],
        ["硬性禁止", "认为可通过", "阻断 + 人工确认", "规则冲突"],
        ["业务条件", "认为有风险", "展示双方理由", "待确认"],
        ["表达建议", "认为无需修改", "允许保留", "建议项"],
    ]
    dk.table(s, 0.55, 1.42, 8.9, rows, col_w=[1.65, 2.15, 2.65, 2.45], size=12, row_h=0.66,
             head_c=DARK, body_c=INK, rule_c=GRID, hi_fill=SOFT_CORAL, hi_c=CORAL)
    takeaway(s, "冲突不是系统失灵，而是需要更清晰的决定权设计。")
    return s


def slide_26(prs):
    """role=decision | form=conflict decision tree | static: user opted out | takeaway='先看规则级别再分流'"""
    s = base(prs, "处理冲突时，先问规则是否具有强制性", "第二节 · 冲突处理", 26, 2,
             "带学员按树判断：硬规则优先阻断；条件性规则展示冲突并转确认；建议型规则保留双方说明。")
    card(s, 3.65, 1.28, 2.7, 0.82, "规则具有强制性吗？", (), fill=TEAL2, hub=True)
    dk.arrow(s, 2.35, 2.18, 1.2, 0.20, color=TEAL, direction="left")
    dk.arrow(s, 6.45, 2.18, 1.2, 0.20, color=TEAL, direction="right")
    card(s, 0.55, 2.48, 3.3, 1.55, "是 · 硬规则", ["阻断自动通过", "交人工确认"], fill=SOFT_CORAL, title_c=CORAL)
    card(s, 6.15, 2.48, 3.3, 1.55, "否 · 条件或建议", ["并列展示理由", "允许确认或保留"], fill=PALE)
    takeaway(s, "系统不能替业务掩盖冲突；要把冲突变成可判断的信息。")
    return s


def slide_27(prs):
    """role=scenario | form=three-role disputes | static: user opted out | takeaway='争议本质是决定权与承接关系'"""
    s = base(prs, "异常五：角色争议，本质是决定权与承接关系", "第二节 · 异常 05", 27, 2,
             "分别讨论作者不同意、审核员怀疑资料、巡检无人处理。每种争议都要明确谁能提出、谁能确认、谁必须继续。")
    card(s, 0.55, 1.42, 2.75, 2.8, "作者不同意", ["允许说明理由", "不能自行改正式结论", "提交复核"], fill=PALE)
    card(s, 3.63, 1.42, 2.75, 2.8, "审核员质疑资料", ["暂停相关结论", "请求资料负责人确认", "保留争议状态"], fill=WHITE, title_c=CORAL)
    card(s, 6.70, 1.42, 2.75, 2.8, "巡检无人承接", ["升级负责人", "设置超时与提醒", "必要时暂停使用"], fill=PALE)
    takeaway(s, "争议页面要给下一步，不只给“不同意”按钮。")
    return s


def slide_28(prs):
    """role=reference | form=role-authority matrix | static: user opted out | takeaway='操作确认负责三种权力分开'"""
    s = base(prs, "把“谁操作、谁确认、谁负责”分开设计", "第二节 · 角色权限", 28, 2,
             "用矩阵确认角色边界。作者能改稿和发起；审核员能退回与正式决定；资料负责人确认资料；内容负责人处理在用内容。")
    rows = [
        ["角色", "可操作", "可确认", "主要责任"],
        ["作者", "编辑、发起自检", "修改说明", "提交可审版本"],
        ["审核员", "复核、退回", "正式结论", "决定是否进入下一环节"],
        ["资料负责人", "更新材料", "资料有效性", "维护业务来源"],
        ["内容负责人", "处理巡检", "实际更新", "维护在用内容"],
    ]
    dk.table(s, 0.50, 1.35, 9.0, rows, col_w=[1.45, 2.45, 2.05, 3.05], size=11.5, row_h=0.58,
             head_c=DARK, body_c=INK, rule_c=GRID, hi_fill=PALE, hi_c=TEAL)
    return s


def slide_29(prs):
    """role=framework | form=five-way response menu | static: user opted out | takeaway='异常处理不是只有重试'"""
    s = base(prs, "异常发生后，系统至少有五种处理方向", "第二节 · 分流菜单", 29, 2,
             "异常处理不是默认重试。根据影响与可恢复性选择继续、重试、降级、转人工或停止。")
    labels = [("继续", "不影响其余项", PALE), ("重试", "临时性失败", WHITE), ("降级", "使用较弱方案", SOFT_GOLD), ("转人工", "需要业务判断", PALE), ("停止", "风险不可接受", SOFT_CORAL)]
    for i, ((a,b,c), x) in enumerate(zip(labels, [0.42, 2.33, 4.24, 6.15, 8.06])):
        card(s, x, 1.55, 1.68, 2.42, a, [b], fill=c, title_c=CORAL if a=="停止" else TEAL, num=i+1)
    takeaway(s, "自动化程度不是越高越好，关键是风险是否仍可控。")
    return s


def slide_30(prs):
    """role=exercise | form=role-lens exercise | static: user opted out | takeaway='同一争议为不同角色提供不同动作'"""
    s = base(prs, "角色练习：同一条争议，三类用户分别需要什么？", "第二节 · 练习", 30, 2,
             "题目：系统认为“65℃热水直接拖地”对象不一致，作者不同意。作者需要说明与改稿；审核员需要比较材料并决定；资料负责人需要确认原始说法。")
    card(s, 0.55, 1.42, 2.75, 2.82, "作者视角", ["原句与问题范围", "修改或说明", "再次提交"], fill=PALE)
    card(s, 3.63, 1.42, 2.75, 2.82, "审核员视角", ["原句、材料与历史", "确认或退回", "正式决定"], fill=TEAL2, hub=True)
    card(s, 6.70, 1.42, 2.75, 2.82, "资料负责人视角", ["被引用材料", "有效性与适用范围", "更新资料"], fill=WHITE)
    takeaway(s, "不要给所有角色同一套按钮；给他们完成职责所需的信息与动作。")
    return s


def slide_31(prs):
    """role=recap | form=decision loop | static: user opted out | takeaway='异常处理是可恢复的决策链'"""
    s = base(prs, "第二节的核心：把报错变成可恢复的决策链", "第二节 · 小结", 31, 2,
             "异常页面要让用户看见影响、知道下一步、找到承接人，并让系统留下状态。这样才能恢复，而不是卡死。")
    flow(s, ["看见影响", "选择动作", "交给角色", "确认恢复"],
         ["范围可解释", "继续或停止", "责任明确", "状态可追踪"], y=1.65, hub=1)
    takeaway(s, "异常设计的终点不是提示成功，而是业务恢复。")
    return s


def slide_32(prs):
    """role=break | form=dark pause card | static: user opted out | takeaway='休息十分钟'"""
    s = dk.add_slide(prs)
    dk.box(s, 0, 0, W, H, fill=DARK, line=DARK)
    tx(s, 0.7, 1.35, 8.6, 0.9, "休息 10 分钟", 40, WHITE, True, PP_ALIGN.CENTER)
    tx(s, 0.7, 2.55, 8.6, 0.45, "下一段：把同一套方法迁移到合同、客服与其他业务。", 17, GRID, False, PP_ALIGN.CENTER)
    dk.box(s, 3.95, 3.55, 2.1, 0.58, fill=CORAL, line=CORAL, round=True)
    tx(s, 3.95, 3.56, 2.1, 0.56, "120 分钟继续", 13, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    dk.speaker_notes(s, "休息十分钟。回来进入跨场景迁移。")
    return s


def slide_33(prs):
    """role=divider | form=dark divider | static: user opted out | takeaway='迁移的是判断结构不是行业名词'"""
    return divider(prs, "三", "方法迁移", "怎样把一个案例，\n迁成另一类 AI 产品？",
                   ["合同审核母案例", "十一项迁移框架", "客服与三个新场景"], 33, 3,
                   "第三节用合同审核作为母案例，提取通用方法，再现场迁移到客服质检、招聘、报销和金融营销。")


def slide_34(prs):
    """role=concept | form=anti-pattern comparison | static: user opted out | takeaway='迁移不是换行业名词'"""
    s = base(prs, "真正的迁移，不是把“营销”两个字替换掉", "第三节 · 迁移定义", 34, 3,
             "左边是表面换皮，右边是结构迁移。要重新定义最终决定、所需输入、人工分工、边界条件与企业连接。")
    card(s, 0.55, 1.48, 3.78, 2.78, "表面换皮", ["审核对象改名", "页面字段改名", "提示词改名", "流程仍照搬"], fill=SOFT_CORAL, title_c=CORAL)
    tx(s, 4.52, 2.38, 0.95, 0.52, "→", 34, TEAL, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    card(s, 5.68, 1.48, 3.78, 2.78, "结构迁移", ["重新定义最终输出", "反推输入与分工", "补齐边界与系统连接", "验证能力是否适用"], fill=PALE)
    takeaway(s, "迁移的是判断结构，而不是行业名词。")
    return s


def slide_35(prs):
    """role=scope | form=scope funnel | static: user opted out | takeaway='合同审核首版必须收窄'"""
    s = base(prs, "合同审核第一步：把“审合同”收窄成可验证范围", "第三节 · 合同母案例", 35, 3,
             "不要做万能合同助手。首版可限定合同类型、审查角色、条款类别和输出方式，例如销售合同中的付款、交付与违约条款。")
    widths = [8.6, 6.8, 5.0, 3.3]
    labels = ["所有合同、所有风险", "销售合同", "付款 / 交付 / 违约", "给法务初审"]
    fills = [PALE, WHITE, SOFT_GOLD, TEAL2]
    for i, (w, lab, f, y0, y1) in enumerate(zip(widths, labels, fills, [1.28, 2.04, 2.80, 3.56], [1.39, 2.15, 2.91, 3.67])):
        x = (W-w)/2
        dk.box(s, x, y0, w, 0.58, fill=f, line=GRID if i<3 else TEAL2, round=True)
        tx(s, x, y1, w, 0.32, lab, 13, WHITE if i==3 else INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    takeaway(s, "范围越清楚，材料、验证与责任边界越容易设计。")
    return s


def slide_36(prs):
    """role=method | form=reverse design chain | static: user opted out | takeaway='从最终输出反推输入'"""
    s = base(prs, "先定义最后要交付什么，再反推需要什么输入", "第三节 · 反向设计", 36, 3,
             "合同案例最终不是给一个风险分数，而是给问题条款、风险类型、参考规则、修改建议与人工确认项。由此反推合同正文、模板、规则与业务背景。")
    flow(s, ["最终输出", "判断任务", "所需输入"],
         ["问题条款 + 建议 + 待确认项", "定位、分类、对照、解释", "正文 + 模板 + 规则 + 背景"], y=1.72, hub=0)
    takeaway(s, "先定输出，才能知道模型究竟要完成哪些判断。")
    return s


def slide_37(prs):
    """role=positioning | form=human-AI division | static: user opted out | takeaway='模型是勤奋法务实习生'"""
    s = base(prs, "模型更像勤奋的法务实习生，而不是最终决定者", "第三节 · 能力定位", 37, 3,
             "这个比喻用于限定能力。模型可快速定位与对照，但企业事实、非标交易安排和最终风险接受仍需有权角色确认。")
    card(s, 0.65, 1.45, 3.95, 2.85, "模型擅长", ["快速阅读", "定位相关条款", "对照规则与模板", "生成结构化初稿"], fill=PALE)
    card(s, 5.40, 1.45, 3.95, 2.85, "人必须负责", ["确认业务事实", "处理非标安排", "权衡风险与交易", "作最终决定"], fill=WHITE, title_c=CORAL)
    takeaway(s, "把模型放在合适岗位，比夸大“智能”更重要。")
    return s


def slide_38(prs):
    """role=workflow | form=parallel contract workflow | static: user opted out | takeaway='按人工分工拆成并行流程'"""
    s = base(prs, "合同审核：把人工分工拆成三条并行流程", "第三节 · 合同流程", 38, 3,
             "一条处理条款识别，一条处理规则或模板对照，一条处理业务条件。三路结果汇合后交给法务确认。")
    card(s, 0.5, 1.30, 2.0, 2.92, "共同输入", ["合同正文", "合同类型", "交易背景"], fill=TEAL2, hub=True)
    card(s, 2.86, 1.30, 1.84, 0.78, "条款识别", ["定位与分类"], fill=PALE)
    card(s, 2.86, 2.28, 1.84, 0.78, "规则对照", ["模板与制度"], fill=WHITE)
    card(s, 2.86, 3.26, 1.84, 0.78, "业务核对", ["条件与例外"], fill=PALE)
    dk.arrow(s, 4.94, 2.48, 0.55, 0.22, color=TEAL)
    card(s, 5.65, 1.30, 1.95, 2.92, "汇总问题", ["风险条款", "修改建议", "待确认项"], fill=SOFT_GOLD)
    dk.arrow(s, 7.82, 2.48, 0.45, 0.22, color=TEAL)
    card(s, 8.35, 1.30, 1.16, 2.92, "法务\n确认", (), fill=TEAL2, hub=True)
    return s


def slide_39(prs):
    """role=architecture | form=boundary branches + connections | static: user opted out | takeaway='通用能力要接入企业边界'"""
    s = base(prs, "从通用能力升级为企业产品，要补齐边界与连接", "第三节 · 企业化", 39, 3,
             "边界包括扫描件、附件、跨条款引用、未知模板与复杂交易。连接包括合同库、模板库、权限、流程与历史版本。")
    card(s, 0.50, 1.30, 2.65, 3.00, "边界分支", ["扫描件识别失败", "附件缺失", "跨条款冲突", "未知合同类型", "复杂交易转人工"], fill=SOFT_CORAL, title_c=CORAL)
    card(s, 3.55, 1.30, 2.90, 3.00, "审核能力", ["识别 → 对照 → 解释", "结果附定位与来源", "明确待确认项"], fill=TEAL2, hub=True)
    card(s, 6.85, 1.30, 2.65, 3.00, "企业连接", ["合同与模板库", "组织权限", "审批流程", "版本与历史", "处理结果回写"], fill=PALE)
    return s


def slide_40(prs):
    """role=peak | form=eleven-step migration framework | static: user opted out | takeaway='十一项完成一次完整迁移'"""
    s = base(prs, "十一项迁移框架：从业务问题走到可运行产品", "第三节 · 核心框架", 40, 3,
             "这是全课的峰值页。按顺序讲十一项：业务问题、首版范围、最终输出、人工流程、输入材料、判断任务、能力验证、系统流程、边界条件、角色权限、企业连接。")
    items = ["业务问题", "首版范围", "最终输出", "人工流程", "输入材料", "判断任务", "能力验证", "系统流程", "边界条件", "角色权限", "企业连接"]
    # 3 phase lanes
    phases = [("A 定义", items[:4], PALE), ("B 构建", items[4:8], WHITE), ("C 落地", items[8:], SOFT_GOLD)]
    idx = 1
    phase_layouts = [
        (1.28, [1.93, 3.89, 5.85, 7.81], 1.74),
        (2.36, [1.93, 3.89, 5.85, 7.81], 1.74),
        (3.44, [1.93, 4.48, 7.03], 2.32),
    ]
    for (phase, arr, fill), (y, xs, step_w) in zip(phases, phase_layouts):
        dk.box(s, 0.46, y, 1.25, 0.84, fill=TEAL2, line=TEAL2, round=True)
        tx(s, 0.46, y, 1.25, 0.82, phase, 12.5, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        for lab, x in zip(arr, xs):
            card(s, x, y, step_w, 0.84, lab, (), fill=fill, num=idx)
            idx += 1
    takeaway(s, "迁移顺序：先定义决定，再设计能力，最后接入真实组织。")
    return s


def slide_41(prs):
    """role=mapping | form=contract-to-service mapping table | static: user opted out | takeaway='同一结构可映射到客服质检'"""
    s = base(prs, "把合同审核结构，映射到客服会话质检", "第三节 · 现场迁移", 41, 3,
             "逐列映射而非直接复制。合同条款对应客服对话片段；法务规则对应服务规范；法务确认对应质检员复核；合同流程回写对应坐席整改与复训。")
    rows = [
        ["迁移维度", "合同审核", "客服会话质检"],
        ["被判断对象", "合同条款", "会话片段 / 整段会话"],
        ["业务材料", "模板、制度、交易背景", "服务规范、产品政策、客户上下文"],
        ["系统输出", "风险条款、建议、待确认", "问题定位、原因、整改建议"],
        ["人工角色", "法务", "质检员 / 班组长"],
        ["后续动作", "修改、审批、归档", "申诉、整改、复训"],
    ]
    dk.table(s, 0.48, 1.24, 9.04, rows, col_w=[1.65, 3.35, 4.04], size=11.1, row_h=0.55,
             head_c=DARK, body_c=INK, rule_c=GRID, hi_fill=PALE, hi_c=TEAL)
    return s


def slide_42(prs):
    """role=workflow | form=parallel customer-service workflow | static: user opted out | takeaway='客服质检也要形成闭环'"""
    s = base(prs, "客服会话质检：从识别问题走到整改复核", "第三节 · 客服流程", 42, 3,
             "展示完整链路：会话输入后，分别检查服务规范、事实准确和沟通质量，汇总问题后交质检员确认，再进入申诉、整改或复训。")
    flow(s, ["会话输入", "并行检查", "人工复核", "整改闭环"],
         ["文本 + 客户上下文", "规范 / 事实 / 沟通", "确认与申诉", "整改 / 复训 / 复核"], y=1.65, hub=2)
    takeaway(s, "质检产品的完成点不是打分，而是问题被承接和改进。")
    return s


def slide_43(prs):
    """role=exercise | form=one-page group canvas | static: user opted out | takeaway='按十一项完成一次迁移'"""
    s = base(prs, "小组迁移练习：用同一框架设计一个新场景", "第三节 · 练习", 43, 3,
             "每组选择招聘信息、发票报销或金融营销材料，在十二分钟内完成六个关键问题，并用两分钟说明首版范围与人工接点。")
    prompts = [("① 业务问题", "写清要解决的判断"), ("② 首版范围", "限定对象与角色"),
               ("③ 最终输出", "列出决定所需结果"), ("④ 输入材料", "说明业务来源"),
               ("⑤ 边界与人工", "标出转人工条件"), ("⑥ 企业连接", "写明系统回写位置")]
    for i, (p, hint) in enumerate(prompts):
        x = 0.48 + (i%3)*3.08
        y = 1.34 + (i//3)*1.43
        card(s, x, y, 2.78, 1.13, p, [hint], fill=PALE if i%2==0 else WHITE)
    takeaway(s, "先把首版做窄，再把闭环做完整。")
    return s


def slide_44(prs):
    """role=examples | form=three mini-cases | static: user opted out | takeaway='不同场景共享判断结构'"""
    s = base(prs, "三个可迁移场景：对象不同，判断结构相似", "第三节 · 三个小案例", 44, 3,
             "快速展开三个案例。招聘信息关注歧视性与承诺；发票报销关注真伪、制度与重复；金融营销关注适当性、风险披露与限制性表述。")
    card(s, 0.48, 1.30, 2.84, 3.02, "招聘信息审核", ["对象：职位文案", "材料：招聘规范", "输出：问题句 + 修改建议", "人工：HR / 法务"], fill=PALE, num="A")
    card(s, 3.58, 1.30, 2.84, 3.02, "发票与报销审核", ["对象：票据 + 报销单", "材料：费用制度", "输出：异常项 + 待补材料", "人工：财务"], fill=WHITE, num="B")
    card(s, 6.68, 1.30, 2.84, 3.02, "金融营销材料", ["对象：宣传文案", "材料：产品与合规规则", "输出：风险点 + 处理动作", "人工：合规"], fill=PALE, num="C")
    return s


def slide_45(prs):
    """role=critique | form=five-criterion review | static: user opted out | takeaway='用五个问题判断迁移质量'"""
    s = base(prs, "点评一份迁移方案，只看五个关键问题", "第三节 · 点评标准", 45, 3,
             "点评不比页面数量。看业务范围是否清楚、输入是否支持输出、人工决定点是否合理、异常是否可恢复、结果是否进入真实流程。")
    labels = [("范围", "是否足够清楚"), ("输入", "能否支撑输出"), ("人工", "决定点是否合理"), ("异常", "能否恢复"), ("闭环", "是否进入真实流程")]
    for i, ((a,b), x) in enumerate(zip(labels, [0.42, 2.33, 4.24, 6.15, 8.06])):
        card(s, x, 1.48, 1.68, 2.62, a, [b], fill=TEAL2 if i==4 else (PALE if i%2==0 else WHITE), hub=i==4, num=i+1)
    takeaway(s, "好方案不是“AI 功能多”，而是判断与行动接得上。")
    return s


def slide_46(prs):
    """role=close | form=one-page product blueprint | static: user opted out | takeaway='把一页方案作为全课输出'"""
    s = dk.add_slide(prs)
    dk.box(s, 0, 0, W, H, fill=DARK, line=DARK)
    tx(s, 0.46, 0.35, 8.8, 0.25, "全课收束 · 课堂输出", 10, GOLD)
    tx(s, 0.46, 0.78, 8.9, 0.62, "用一页纸，说明一个判断型 AI 产品", 31, WHITE, True)
    cols = [("1 定义", ["业务问题", "首版范围", "最终输出"]),
            ("2 构建", ["输入材料", "判断任务", "能力验证"]),
            ("3 落地", ["角色权限", "异常分流", "企业连接"])]
    for i, ((head, items), x) in enumerate(zip(cols, [0.48, 3.60, 6.72])):
        dk.box(s, x, 1.72, 2.84, 2.55, fill=TEAL2 if i==1 else RGBColor(0x16,0x42,0x42), line=TEAL2 if i==1 else MUTE, round=True)
        tx(s, x+0.18, 1.94, 2.48, 0.32, head, 16, WHITE, True)
        paras(s, x+0.18, 2.48, 2.48, 1.45, items, 13, WHITE, 9)
    tx(s, 0.48, 4.65, 9.0, 0.42, "最终检验：对象说得清 · 状态接得上 · 异常有人管 · 结果进业务", 17, GOLD, True, PP_ALIGN.CENTER)
    tx(s, 0.48, 5.28, 2.5, 0.18, "第二次课 · 完", 8, GRID)
    tx(s, 9.46, 5.28, 0.2, 0.18, "46", 8, GRID, align=PP_ALIGN.RIGHT)
    dk.speaker_notes(s, "请学员把一页方案作为课后产出。最后用四个条件自检：对象、状态、异常承接和业务闭环。")
    return s


SLIDES = [
    slide_01, slide_02, slide_03, slide_04, slide_05, slide_06, slide_07, slide_08,
    slide_09, slide_10, slide_11, slide_12, slide_13, slide_14, slide_15, slide_16,
    slide_17, slide_18, slide_19, slide_20, slide_21, slide_22, slide_23, slide_24,
    slide_25, slide_26, slide_27, slide_28, slide_29, slide_30, slide_31, slide_32,
    slide_33, slide_34, slide_35, slide_36, slide_37, slide_38, slide_39, slide_40,
    slide_41, slide_42, slide_43, slide_44, slide_45, slide_46,
]


def main():
    OUTDIR.mkdir(parents=True, exist_ok=True)
    prs = dk.blank_deck()
    for fn in SLIDES:
        fn(prs)
    dk.lint_layout(prs, strict=True)
    prs.save(OUT)
    dk.declare_delivery(str(OUT), "presented")
    print(f"saved -> {OUT} | slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
