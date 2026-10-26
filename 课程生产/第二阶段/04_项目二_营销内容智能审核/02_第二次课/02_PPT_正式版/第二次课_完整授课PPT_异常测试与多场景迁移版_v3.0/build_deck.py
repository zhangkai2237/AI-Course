#!/usr/bin/env python3
"""Build the corrected 48-slide second-lesson deck.

Build shape: solo — subagent dispatch not authorized in this run.
Visual lock: follows the existing project deck.
Motion: static throughout.
"""
from pathlib import Path
import importlib.util
from copy import deepcopy

from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = Path(__file__).resolve().parent
SOURCE = HERE.parent / "第二次课_完整授课PPT_审核体系方案选择与多场景迁移版_v2.0" / "build_deck.py"
spec = importlib.util.spec_from_file_location("prior_deck", SOURCE)
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)

m.SECTIONS = ["4  异常取舍", "5  测试判断", "7  场景迁移"]
OUT = HERE / "第二次课_完整授课PPT_异常测试与多场景迁移版_v3.0.pptx"

# Use a macOS system Chinese font so both PowerPoint and the local renderer can display CJK.
_prior_tx = m.tx


def _tx_cjk(s, x, y, w, h, value, size=17, color=m.INK, bold=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Hiragino Sans GB"):
    return _prior_tx(s, x, y, w, h, value, size, color, bold, align, anchor, font)


def _paras_cjk(s, x, y, w, h, values, size=12, color=m.INK, gap=5, bullet=False):
    rows = []
    for value in values:
        prefix = "•  " if bullet else ""
        rows.append([(prefix + value, size, color, False, False, "Hiragino Sans GB")])
    return m.dk.text(s, x, y, w, h, rows, space_after=gap, line_spacing=1.0)


m.tx = _tx_cjk
m.paras = _paras_cjk
m.dk.set_palette(deep=m.DARK, blue=m.TEAL, teal=m.TEAL2, magenta=m.CORAL, slate=m.INK,
                 mute=m.MUTE, mono="Arial", font="Arial", display="Arial",
                 eafont="Hiragino Sans GB", eadisplay="Hiragino Sans GB",
                 accents=[m.TEAL, m.CORAL, m.GOLD, m.MUTE])


def cover(prs):
    s = m.dk.add_slide(prs)
    m.dk.box(s, 0, 0, m.W, m.H, fill=m.DARK, line=m.DARK)
    m.tx(s, 0.48, 0.38, 3.5, 0.24, "项目二 · 第二次课", 10, m.GOLD, True)
    m.tx(s, 0.48, 1.03, 8.9, 0.88, "真实异常、测试判断\n与多场景迁移", 34, m.WHITE, True)
    m.tx(s, 0.50, 2.63, 7.9, 0.42, "从“能跑”走向“敢试点”，再走向可迁移的方法", 17, m.WHITE)
    for i, (label, x) in enumerate(zip(["异常取舍", "测试判断", "场景迁移"], [0.50, 3.38, 6.26]), start=1):
        m.dk.box(s, x, 3.58, 2.55, 0.82, fill=m.TEAL2 if i == 2 else m.PALE, line=m.TEAL2 if i == 2 else m.PALE, round=True)
        m.tx(s, x, 3.58, 2.55, 0.82, label, 15, m.WHITE if i == 2 else m.TEAL, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if i < 3:
            m.dk.arrow(s, x + 2.60, 3.88, 0.55, 0.22, color=m.GOLD)
    m.tx(s, 0.50, 5.22, 3.8, 0.18, "AI 产品经理项目实战", 8, m.GRID)
    m.tx(s, 9.45, 5.22, 0.2, 0.18, "1", 8, m.GRID, align=PP_ALIGN.RIGHT, font="Arial")
    m.dk.speaker_notes(s, "开场只说明今天的三道判断门：异常中能否安全取舍、测试后能否决定下一步、换业务后能否重新完成产品分析。")


def map_slide(prs):
    s = m.base(prs, "今天不是继续堆功能，而是连续过三道判断门", "课程地图", 2, 1,
               "用三道门建立全课地图。第四节回答异常中如何继续；第五节回答测试后是否继续；第七节回答换场景后方法是否继续成立。")
    items = [
        ("异常门", ["流程被打断时", "能否安全停下、恢复与交接"]),
        ("测试门", ["结果不完美时", "进入试点、补齐、收缩或暂停"]),
        ("迁移门", ["换一个业务后", "重新定义对象、材料、角色与边界"]),
    ]
    for i, ((head, lines), x) in enumerate(zip(items, [0.48, 3.57, 6.66]), start=1):
        m.card(s, x, 1.45, 2.86, 2.75, head, lines, fill=m.PALE if i != 2 else m.WHITE, hub=i == 2, num=i)
        if i < 3:
            m.dk.arrow(s, x + 2.91, 2.68, 0.22, 0.22, color=m.GOLD)
    m.takeaway(s, "能跑通只是起点；真实产品必须能处理异常、接受测试并跨场景重做判断。")


def opening(prs):
    s = m.base(prs, "一个产品能运行，不等于能交给真实业务", "开场 · 从演示走向试点", 3, 1,
               "用一个正常演示与五个真实打断做对照。不要展开技术细节，先让学员感受到今天的判断难度。")
    m.card(s, 0.48, 1.36, 3.20, 2.85, "演示中的顺利路径", ["提交文案", "机器返回问题", "运营修改", "再次提交"], fill=m.PALE, num=1)
    m.card(s, 4.05, 1.36, 5.47, 2.85, "真实业务中的打断", ["活动材料未同步", "某一路模型超时", "检查途中修改文案", "规则与模型判断不同", "发现问题却无人负责"], fill=m.WHITE, title_c=m.CORAL, num=2)
    m.takeaway(s, "今天关注的不是系统顺利时多聪明，而是不顺利时怎样做出安全、可解释的取舍。", danger=True)


def section_divider(prs, num, title, question, agenda, page, sec, notes):
    s = m.dk.add_slide(prs)
    m.dk.box(s, 0, 0, m.W, m.H, fill=m.DARK, line=m.DARK)
    m.tx(s, 0.42, 0.38, 4.8, 0.24, f"第二次课 / 第{num}节", 10, m.GOLD)
    m.tx(s, 0.42, 0.86, 5.45, 0.52, title, 29, m.WHITE, True)
    m.tx(s, 0.42, 1.80, 4.25, 1.08, question, 21, m.WHITE, True)
    m.dk.box(s, 5.02, 1.67, 0.015, 2.02, fill=m.MUTE, line=m.MUTE)
    for item, y in zip(agenda, [1.83, 2.52, 3.21]):
        m.tx(s, 5.28, y, 4.15, 0.42, item, 14.2, m.WHITE)
    m.dk.box(s, 0.42, 4.28, 9.12, 0.68, fill=m.TEAL2, line=m.TEAL2, round=True)
    m.tx(s, 0.64, 4.43, 8.68, 0.32, "本节重点  ·  " + agenda[-1], 13.0, m.WHITE,
         True, anchor=MSO_ANCHOR.MIDDLE)
    m.tx(s, 0.42, 5.25, 3.2, 0.18, "异常取舍 → 测试判断 → 场景迁移", 7.5, m.GRID)
    m.tx(s, 9.45, 5.25, 0.2, 0.18, str(page), 8, m.GRID, align=PP_ALIGN.RIGHT, font="Arial")
    m.dk.speaker_notes(s, notes)
    return s


def section4(prs):
    section_divider(prs, "四", "真实异常中的产品取舍", "正常流程被打断，\n产品怎样安全继续？",
                    ["把异常设计成业务状态", "处理五类真实异常", "选择继续、重试、降级、人工或终止"], 4, 1,
                    "第四节讲异常中的产品取舍。关键不是列异常，而是说明每次取舍背后的业务条件。")

    s = m.base(prs, "异常不是一句报错，而是一次未完成的业务任务", "第四节 · 异常定义", 5, 1,
               "把技术报错翻译成业务任务。用户必须知道未完成的位置、影响、动作和负责人。")
    m.card(s, 0.52, 1.45, 3.45, 2.55, "低质量表达", ["系统异常", "请稍后再试", "未发现问题"], fill=m.SOFT_CORAL, title_c=m.CORAL, num=1)
    m.tx(s, 4.22, 2.40, 1.05, 0.42, "→", 30, m.TEAL, True, PP_ALIGN.CENTER)
    m.card(s, 5.55, 1.45, 3.88, 2.55, "可处理的业务状态", ["哪里没完成", "影响哪些判断", "当前还能做什么", "由谁继续"], fill=m.PALE, num=2)
    m.takeaway(s, "异常设计的目标，是让任务可以被理解、恢复和交接。")

    s = m.base(prs, "一张异常卡要回答八个问题", "第四节 · 八项框架", 6, 1,
               "逐项解释八个字段。强调这是一张完整异常卡，而不是八条互不相关的功能。")
    fields = ["触发条件", "影响范围", "当前状态", "页面表达", "可执行动作", "处理角色", "恢复方式", "过程留存"]
    for i, label in enumerate(fields):
        x = 0.50 + (i % 4) * 2.28
        y = 1.30 + (i // 4) * 1.28
        m.card(s, x, y, 2.05, 0.98, label, (), fill=m.PALE if i % 2 == 0 else m.WHITE, num=i + 1)
    m.takeaway(s, "异常设计 = 状态说明 + 可执行动作 + 责任交接 + 恢复路径。")

    s = m.base(prs, "材料缺失时，不要把“无法判断”伪装成业务结论", "第四节 · 材料缺失", 7, 1,
               "讲赠品资格材料缺失案例。区分内容错误与材料不足；说明只重跑受影响分支。")
    m.card(s, 0.48, 1.34, 2.54, 2.90, "场景", ["文案：下单即送滤芯", "缺少资格、数量和有效期"], fill=m.SOFT_GOLD, num=1)
    m.card(s, 3.36, 1.34, 2.54, 2.90, "错误动作", ["直接判违规", "直接判无问题", "阻塞所有检查"], fill=m.SOFT_CORAL, title_c=m.CORAL, num=2)
    m.card(s, 6.24, 1.34, 3.22, 2.90, "正确动作", ["该检查点未完成", "其他分支继续", "提示补充哪类材料", "补齐后重跑受影响分支"], fill=m.PALE, num=3)
    m.takeaway(s, "材料不足说明信息不够；内容错误说明现有信息已经足以作出判断。")

    s = m.base(prs, "局部失败必须显示为“部分完成”", "第四节 · 模型超时", 8, 1,
               "说明三路检查只失败一路时的状态。是否允许业务继续，取决于失败分支是否为当前场景必检项。")
    m.flow(s, ["功能对象", "价格条件", "赠品语义", "整次运行"],
           ["完成", "完成", "超时待处理", "部分完成"], y=1.52, hub=2)
    m.card(s, 0.70, 3.38, 4.10, 0.96, "非必检分支", ["可保留结果并继续后续动作"], fill=m.PALE)
    m.card(s, 5.20, 3.38, 4.10, 0.96, "必检分支", ["不得自动通过；重试或人工补核"], fill=m.SOFT_CORAL, title_c=m.CORAL)
    m.takeaway(s, "不能用统一成功率替代业务判断；要看缺失分支是否影响当前决定。", danger=True)

    s = m.base(prs, "结果正确，也不能覆盖错误的版本", "第四节 · 版本错位", 9, 1,
               "演示V1发起检查、V2保存、V1结果迟到。强调输出与输入快照绑定。")
    m.flow(s, ["V1 发起", "保存 V2", "V1 迟到", "V2 重跑"],
           ["绑定内容与材料快照", "当前正文已变化", "只回到 V1 运行记录", "形成新的有效结果"], y=1.50, hub=2)
    m.takeaway(s, "AI 输出必须绑定内容、材料和规则版本；迟到结果保留，但不覆盖当前版本。")

    s = m.base(prs, "规则与模型不是两位投票者", "第四节 · 结论冲突", 10, 1,
               "讲两个案例：最后作为顺序词；1799计算成立但资格条件缺失。先判断两个结果是否回答同一问题。")
    m.card(s, 0.48, 1.35, 2.72, 2.95, "规则 / 程序", ["确定表达", "字段和值", "计算与格式"], fill=m.PALE, num=1)
    m.card(s, 3.52, 1.35, 2.72, 2.95, "模型", ["上下文", "对象关系", "语义条件"], hub=True, num=2)
    m.card(s, 6.56, 1.35, 2.96, 2.95, "业务材料 / 人工", ["可接受范围", "有效版本", "无法自动确定的争议"], fill=m.WHITE, num=3)
    m.takeaway(s, "先判断是否在回答同一个问题，再决定合并、复核还是交人工。")

    s = m.base(prs, "发现问题以后，还必须有人拥有下一步", "第四节 · 责任承接", 11, 1,
               "依次讲作者异议、材料过期和线上巡检无人处理。重点是权限与责任不能被机器结论替代。")
    rows = [
        ["角色", "可以做", "不能直接做"],
        ["作者", "修改、补充、提交异议", "修改正式结论"],
        ["审核员", "确认、退回、人工补核", "复制旧版本通过状态"],
        ["材料负责人", "更新受控材料、确认有效期", "静默改写历史材料"],
        ["内容负责人", "认领、分配、升级", "未经复核关闭高风险问题"],
    ]
    m.dk.table(s, 0.52, 1.32, 8.98, rows, col_w=[1.55, 3.25, 4.18], size=11, row_h=0.60,
               head_c=m.DARK, body_c=m.INK, rule_c=m.GRID, hi_fill=m.PALE, hi_c=m.TEAL)
    m.takeaway(s, "没有负责人、时限和升级路径，产品只是把问题从线上转移到列表。", danger=True)

    s = m.base(prs, "五种动作代表五种不同的产品判断", "第四节 · 决策阶梯", 12, 1,
               "逐个解释继续、重试、降级、转人工和终止。强调降级不是把失败改成成功。")
    labels = ["继续", "重试", "降级", "转人工", "终止"]
    subs = ["不影响当前决定", "输入有效，执行暂时失败", "换用边界明确的替代能力", "需要判断、权限或补充信息", "无法安全继续"]
    for i, (label, sub) in enumerate(zip(labels, subs)):
        x = 0.42 + i * 1.87
        y = 3.60 - i * 0.48
        m.card(s, x, y, 1.67, 1.05, label, [sub], fill=m.PALE if i < 3 else m.SOFT_CORAL, title_c=m.CORAL if i >= 3 else m.TEAL, num=i + 1)
    m.takeaway(s, "选择动作的依据，是对业务决定的影响，而不是技术团队觉得哪种处理方便。")

    s = m.base(prs, "同一个异常，对不同角色意味着不同动作", "第四节 · 角色视角", 13, 1,
               "用1799优惠条件不完整做教师推演，不组织角色练习。分别说明作者、审核员与材料负责人需要看到和执行什么。")
    items = [
        ("作者", ["看原句与缺失条件", "补正文或上下文", "提交异议"]),
        ("审核员", ["看提交快照与材料版本", "退回或人工补核", "记录最终原因"]),
        ("材料负责人", ["看旧版与新版生效时间", "确认有效版本", "触发受影响分支重跑"]),
    ]
    for i, ((head, lines), x) in enumerate(zip(items, [0.48, 3.57, 6.66]), start=1):
        m.card(s, x, 1.42, 2.86, 2.86, head, lines, fill=m.PALE if i % 2 else m.WHITE, num=i)
    m.takeaway(s, "角色不是展示字段，而是决定谁能看什么、做什么、把任务交给谁。")

    s = m.base(prs, "正常流程决定能不能演示，异常流程决定能不能上线", "第四节 · 小结", 14, 1,
               "用四个词收束第四节：可解释、可执行、可交接、可恢复。为第五节的测试判断做过渡。")
    for i, (head, sub, x) in enumerate([
        ("可解释", "哪里没完成", 0.55), ("可执行", "现在能做什么", 2.84),
        ("可交接", "由谁继续", 5.13), ("可恢复", "条件补齐后怎样继续", 7.42)
    ], start=1):
        m.card(s, x, 1.55, 2.02, 2.55, head, [sub], fill=m.PALE if i % 2 else m.WHITE, num=i)
    m.takeaway(s, "异常处理不是附加功能，而是业务闭环的一部分。")


def break_slide(prs, page, next_text):
    s = m.dk.add_slide(prs)
    m.dk.box(s, 0, 0, m.W, m.H, fill=m.DARK, line=m.DARK)
    m.tx(s, 0.55, 1.32, 8.9, 0.55, "休息 10 分钟", 34, m.WHITE, True, PP_ALIGN.CENTER)
    m.tx(s, 0.55, 2.30, 8.9, 0.42, next_text, 18, m.GOLD, True, PP_ALIGN.CENTER)
    m.tx(s, 9.45, 5.25, 0.2, 0.18, str(page), 8, m.GRID, align=PP_ALIGN.RIGHT, font="Arial")
    m.dk.speaker_notes(s, "休息十分钟。")


def section5(prs):
    section_divider(prs, "五", "用测试结果判断方案是否成立", "一个漂亮分数，\n能不能支持进入试点？",
                    ["建立可复现测试集", "同时看效果、稳定性、业务与高风险", "决定试点、补齐、收缩或暂停"], 16, 2,
                    "第五节不重新讲指标公式，而是讲产品经理怎样把测试结果转成方案判断。")

    s = m.base(prs, "测试的不是模型聪不聪明，而是整套方案能否完成任务", "第五节 · 测试对象", 17, 2,
               "强调规则、模型、材料、合并、页面和人工共同构成被测试方案。")
    labels = ["规则", "模型", "业务材料", "结果合并", "页面动作", "人工承接"]
    for i, label in enumerate(labels):
        x = 0.48 + (i % 3) * 3.05
        y = 1.34 + (i // 3) * 1.38
        m.card(s, x, y, 2.82, 1.12, label, (), fill=m.PALE if i % 2 == 0 else m.WHITE, num=i + 1)
    m.takeaway(s, "模型分数高，不代表产品流程稳定，也不代表业务总投入下降。")

    s = m.base(prs, "测试样本要来自真实任务、典型案例与边界情况", "第五节 · 样本来源", 18, 2,
               "讲四种来源。提醒脱敏与版本管理；构造样本只能补齐覆盖，不能替代真实分布。")
    items = [
        ("历史内容", ["脱敏后的真实审核任务"]),
        ("典型案例", ["审核员整理的常见与高风险问题"]),
        ("边界样本", ["围绕重点规则主动构造"]),
        ("线上抽检", ["已上线内容的实际表现"]),
    ]
    for i, ((head, lines), x) in enumerate(zip(items, [0.50, 2.80, 5.10, 7.40]), start=1):
        m.card(s, x, 1.52, 2.05, 2.52, head, lines, fill=m.PALE if i % 2 else m.WHITE, num=i)
    m.takeaway(s, "构造样本用来补齐覆盖，真实样本用来保持业务分布。")

    s = m.base(prs, "一条测试样本必须能够被复现和复核", "第五节 · 样本字段", 19, 2,
               "逐项解释六个字段。特别强调标注人和版本，否则同一内容会随着规则变化失去可比性。")
    fields = ["内容与场景", "必要上下文", "问题位置", "类型与等级", "期望动作", "标注人与版本"]
    for i, label in enumerate(fields):
        x = 0.48 + (i % 3) * 3.06
        y = 1.33 + (i // 3) * 1.42
        m.card(s, x, y, 2.82, 1.16, label, (), fill=m.PALE if i % 2 == 0 else m.WHITE, num=i + 1)
    m.takeaway(s, "没有上下文和版本的测试结果，无法解释差异来自能力变化还是标注变化。")

    s = m.base(prs, "测试集不能只放明显问题", "第五节 · 样本结构", 20, 2,
               "说明四类样本各自防止什么偏差。正常样本防误报，问题样本测识别，边界样本测条件，异常样本测安全承接。")
    items = [
        ("正常样本", ["观察误报与额外工作量"]),
        ("问题样本", ["观察识别、定位与分级"]),
        ("边界样本", ["观察否定、引用、范围与条件"]),
        ("异常样本", ["观察失败、冲突与人工接管"]),
    ]
    for i, ((head, lines), x) in enumerate(zip(items, [0.50, 2.80, 5.10, 7.40]), start=1):
        m.card(s, x, 1.46, 2.05, 2.65, head, lines, fill=m.PALE if i % 2 else m.WHITE, num=i)
    m.takeaway(s, "只用明显违规样本测试，会得到虚假的高分与过度乐观的结论。", danger=True)

    s = m.base(prs, "四组指标必须一起看", "第五节 · 指标体系", 21, 2,
               "不讲公式推导。用四组指标分别回答识别、稳定、业务和高风险问题。")
    items = [
        ("识别效果", ["分类型准确率、召回率、F1", "定位正确率"]),
        ("工程稳定性", ["结构化输出成功率", "超时率、依赖失败率"]),
        ("业务处理", ["人工复核比例、处理时长", "一次修改通过率"]),
        ("高风险表现", ["严重漏检数量", "人工调整高风险结论比例"]),
    ]
    for i, ((head, lines), x) in enumerate(zip(items, [0.50, 2.80, 5.10, 7.40]), start=1):
        m.card(s, x, 1.35, 2.05, 2.88, head, lines, fill=m.PALE if i % 2 else m.WHITE, hub=i == 4, num=i)
    m.takeaway(s, "平均分掩盖不了高风险漏检；高风险门槛需要单独判断。")

    s = m.base(prs, "高风险问题要按成本和责任单独设门槛", "第五节 · 高风险判断", 22, 2,
               "用两类同样数量的错误做对照。说明误报增加工作量，严重漏检可能直接改变是否试点的结论。")
    m.card(s, 0.60, 1.42, 3.65, 2.76, "十条普通误报", ["增加运营修改与复核", "通常可以通过规则或合并优化", "影响效率判断"], fill=m.PALE, num=1)
    m.card(s, 5.75, 1.42, 3.65, 2.76, "一条严重漏检", ["可能直接进入线上", "影响责任与业务风险", "可能触发暂停试点"], fill=m.SOFT_CORAL, title_c=m.CORAL, num=2)
    m.takeaway(s, "错误数量相同，业务后果可能完全不同。", danger=True)

    s = m.base(prs, "机器更快，不代表新流程更省", "第五节 · 人工基线", 23, 2,
               "用8分钟人工、40秒机器、60%复核、6分钟复核举例。不要计算精确ROI，只强调总投入。")
    m.flow(s, ["机器耗时", "人工复核", "失败返工", "维护成本"],
           ["每篇运行时间", "比例 × 单次耗时", "重试、补材料、重复操作", "规则、材料与运营"], y=1.40, hub=1)
    m.dk.box(s, 1.60, 3.60, 6.80, 0.62, fill=m.DARK, line=m.DARK, round=True)
    m.tx(s, 1.60, 3.60, 6.80, 0.62, "四项相加，才是新流程的真实总投入", 15, m.WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    m.takeaway(s, "比较对象不是模型速度，而是新流程总投入与原人工基线。")

    s = m.base(prs, "课堂数值只用于讨论，不能直接变成项目门槛", "第五节 · 门槛设定", 24, 2,
               "展示90%、85%、2分钟三个讨论值，并立即说明真实门槛取决于问题成本、样本分布与人工基线。")
    values = [("≥ 90%", "准确率讨论值"), ("≥ 85%", "召回率讨论值"), ("≤ 2分钟", "单篇耗时讨论值")]
    for i, ((value, label), x) in enumerate(zip(values, [0.62, 3.70, 6.78]), start=1):
        m.dk.box(s, x, 1.55, 2.60, 1.58, fill=m.PALE if i != 2 else m.TEAL2, line=m.PALE if i != 2 else m.TEAL2, round=True)
        m.tx(s, x, 1.78, 2.60, 0.48, value, 25, m.WHITE if i == 2 else m.TEAL, True, PP_ALIGN.CENTER)
        m.tx(s, x, 2.45, 2.60, 0.30, label, 11, m.WHITE if i == 2 else m.MUTE, False, PP_ALIGN.CENTER)
    m.takeaway(s, "真实门槛 = 问题成本 + 样本分布 + 人工基线 + 组织可承受范围。", danger=True)

    s = m.base(prs, "测试最终要导向四种产品判断", "第五节 · 决策矩阵", 25, 2,
               "这是第五节峰值页。用四个象限说明进入试点、补齐能力、收缩范围和暂停。")
    m.dk.box(s, 0.70, 1.25, 4.15, 1.40, fill=m.PALE, line=m.GRID, round=True)
    m.tx(s, 0.94, 1.45, 3.70, 0.30, "进入限定试点", 17, m.TEAL, True)
    m.tx(s, 0.94, 1.92, 3.70, 0.34, "高风险可接受 · 流程稳定 · 总投入下降", 11.5, m.INK)
    m.dk.box(s, 5.15, 1.25, 4.15, 1.40, fill=m.WHITE, line=m.GRID, round=True)
    m.tx(s, 5.39, 1.45, 3.70, 0.30, "补齐能力后回归", 17, m.TEAL, True)
    m.tx(s, 5.39, 1.92, 3.70, 0.34, "问题集中在少数可修复环节", 11.5, m.INK)
    m.dk.box(s, 0.70, 2.95, 4.15, 1.40, fill=m.SOFT_GOLD, line=m.GRID, round=True)
    m.tx(s, 0.94, 3.15, 3.70, 0.30, "收缩场景或拆分流程", 17, m.TEAL, True)
    m.tx(s, 0.94, 3.62, 3.70, 0.34, "某类场景持续拖累整体表现", 11.5, m.INK)
    m.dk.box(s, 5.15, 2.95, 4.15, 1.40, fill=m.SOFT_CORAL, line=m.CORAL, round=True)
    m.tx(s, 5.39, 3.15, 3.70, 0.30, "暂停并重新设计", 17, m.CORAL, True)
    m.tx(s, 5.39, 3.62, 3.70, 0.34, "高风险漏检或失败承接不可接受", 11.5, m.INK)
    m.takeaway(s, "测试不是证明原方案正确，而是决定下一步怎么做。")

    s = m.base(prs, "坏案例一：关键词命中，却忽略否定关系", "第五节 · 否定句误报", 26, 2,
               "读出输入。先定位根因，再给修改动作和回归样本。")
    m.card(s, 0.55, 1.38, 2.80, 2.85, "输入", ["我们不承诺100%有效", "实际体验因人而异"], fill=m.SOFT_GOLD, num=1)
    m.card(s, 3.60, 1.38, 2.80, 2.85, "根因", ["关键词只负责命中", "没有判断否定关系"], fill=m.SOFT_CORAL, title_c=m.CORAL, num=2)
    m.card(s, 6.65, 1.38, 2.80, 2.85, "改进", ["增加上下文语义确认", "加入否定、引用、对比样本"], fill=m.PALE, num=3)
    m.takeaway(s, "先判断错误发生在哪一层，再决定改规则、模型还是测试集。")

    s = m.base(prs, "坏案例二：上下文不足，却强行给出确定结论", "第五节 · 材料不足", 27, 2,
               "输入是连续三年销量第一。说明缺少报告范围、时间和来源时，应返回需补材料或转人工。")
    m.flow(s, ["宣传表达", "检查所需材料", "当前缺口", "正确输出"],
           ["连续三年销量第一", "报告范围、时间、来源", "无法确认材料有效性", "需补材料 / 转人工"], y=1.55, hub=2)
    m.takeaway(s, "不确定时能停下，是判断型 AI 产品的重要能力。", danger=True)

    s = m.base(prs, "坏案例三：一个问题被展示了很多次", "第五节 · 结果重复", 28, 2,
               "输入是全网第一、排名第一。说明合并逻辑不只是视觉去重，而是确定一个可处理业务对象。")
    m.card(s, 0.55, 1.36, 2.70, 2.86, "多路原始结果", ["规则A命中", "规则B命中", "模型也识别到同一问题"], fill=m.WHITE, num=1)
    m.tx(s, 3.55, 2.38, 0.72, 0.48, "→", 30, m.TEAL, True, PP_ALIGN.CENTER)
    m.card(s, 4.55, 1.36, 4.90, 2.86, "合并后的业务问题", ["按原句位置、类型与语义相似度合并", "保留最高等级与完整说明", "仍可回看各分支原始记录"], fill=m.PALE, num=2)
    m.takeaway(s, "用户要处理的是一个问题，不是三个模型输出。")

    s = m.base(prs, "坏案例要落到具体责任环节", "第五节 · 根因地图", 29, 2,
               "展示九类根因。讲解时要求每个坏案例先归类，再决定修改动作。")
    causes = ["误报", "漏报", "分级错误", "定位错误", "材料缺失", "输出校验失败", "结果合并失败", "页面动作不可用", "人工承接失败"]
    for i, label in enumerate(causes):
        x = 0.45 + (i % 3) * 3.08
        y = 1.24 + (i // 3) * 0.98
        m.card(s, x, y, 2.83, 0.75, label, (), fill=m.PALE if i % 2 == 0 else m.WHITE, num=i + 1)
    m.takeaway(s, "错误归类越准确，修改动作越可能真正解决问题。")

    s = m.base(prs, "每次修改都要回到样本里重新验证", "第五节 · 迭代闭环", 30, 2,
               "讲完整闭环，并要求记录改动原因、影响范围、旧问题与新问题。")
    m.flow(s, ["线上任务", "结果对比", "问题归类", "修改回归", "发布观察"],
           ["真实输入", "机器与人工", "定位责任环节", "规则·提示·材料·页面", "新版本继续观察"], y=1.60, hub=3)
    m.takeaway(s, "回归测试既确认旧问题消失，也检查是否引入了新问题。")

    s = m.base(prs, "同一组结果，会对应不同的下一步", "第五节 · 判断练习", 31, 2,
               "教师口头给出四类结果，让学员快速说出下一步。此页不安排分组产出。")
    cases = [
        ("高风险可接受", "限定试点"),
        ("局部集中出错", "补齐后回归"),
        ("单类持续偏差", "收缩或拆分"),
        ("严重漏检", "暂停重做"),
    ]
    for i, ((a, b), x) in enumerate(zip(cases, [0.50, 2.80, 5.10, 7.40]), start=1):
        m.card(s, x, 1.42, 2.05, 2.72, a, ["→  " + b], fill=m.SOFT_CORAL if i == 4 else (m.PALE if i % 2 else m.WHITE), title_c=m.CORAL if i == 4 else m.TEAL, num=i)
    m.takeaway(s, "产品经理的判断体现在：敢推进，也敢收缩和暂停。")

    s = m.base(prs, "一个平均分，不能替代产品判断", "第五节 · 小结", 32, 2,
               "用四个问题收束：识别、稳定、业务和高风险。为多场景迁移做过渡。")
    qs = ["找得准不准？", "流程稳不稳定？", "人工总投入是否下降？", "高风险是否可接受？"]
    for i, (q, x) in enumerate(zip(qs, [0.55, 2.84, 5.13, 7.42]), start=1):
        m.card(s, x, 1.52, 2.02, 2.58, q, (), fill=m.PALE if i % 2 else m.WHITE, num=i)
    m.takeaway(s, "四个问题一起成立，方案才有资格进入限定范围试点。")


def close_slide(prs):
    s = m.dk.add_slide(prs)
    m.dk.box(s, 0, 0, m.W, m.H, fill=m.DARK, line=m.DARK)
    m.tx(s, 0.46, 0.35, 8.8, 0.25, "全课收束", 10, m.GOLD)
    m.tx(s, 0.46, 0.78, 8.9, 0.62, "专业性来自三次判断", 31, m.WHITE, True)
    cols = [
        ("异常中的判断", ["继续还是停下", "重试、降级还是人工", "谁承担下一步"]),
        ("测试后的判断", ["进入试点", "补齐或收缩", "必要时暂停"]),
        ("迁移时的判断", ["重新理解对象", "重新确定材料和角色", "重新设计边界"]),
    ]
    for i, ((head, items), x) in enumerate(zip(cols, [0.48, 3.60, 6.72])):
        fill = m.TEAL2 if i == 1 else m.RGBColor(0x16, 0x42, 0x42)
        m.dk.box(s, x, 1.72, 2.84, 2.55, fill=fill, line=fill, round=True)
        m.tx(s, x + 0.18, 1.94, 2.48, 0.32, head, 16, m.WHITE, True)
        m.paras(s, x + 0.18, 2.48, 2.48, 1.45, items, 12.5, m.WHITE, 7)
    m.tx(s, 0.48, 4.66, 9.0, 0.42, "异常能接住 · 测试能决策 · 方法能迁移", 17, m.GOLD, True, PP_ALIGN.CENTER)
    m.tx(s, 0.48, 5.28, 2.5, 0.18, "第二次课 · 完", 8, m.GRID)
    m.tx(s, 9.46, 5.28, 0.2, 0.18, "48", 8, m.GRID, align=PP_ALIGN.RIGHT, font="Arial")
    m.dk.speaker_notes(s, "结束时回到三次判断。AI可以快速生成界面和方案草稿，但产品经理需要对异常、测试结果和跨场景迁移承担判断。")


def clone_slide(prs, source_slide):
    """Clone a native-shape slide from the prior approved deck."""
    target = m.dk.add_slide(prs)
    for shape in source_slide.shapes:
        target.shapes._spTree.insert_element_before(deepcopy(shape.element), "p:extLst")
    notes = "本页沿用已确认的正式案例讲解。案例用于教师完整推演，不组织分组练习。"
    try:
        source_notes = source_slide.notes_slide.notes_text_frame.text.strip()
        if source_notes:
            notes = source_notes
    except Exception:
        pass
    m.dk.speaker_notes(target, notes)
    return target


def build(prs):
    cover(prs)
    map_slide(prs)
    opening(prs)
    section4(prs)
    break_slide(prs, 15, "下一部分：用测试结果决定方案是否成立")
    section5(prs)
    break_slide(prs, 33, "下一部分：把同一方法迁移到五个完整业务案例")
    # Reuse the previously approved, formal case-treatment section. These are taught cases, not exercises.
    prior = m.dk.blank_deck()
    m.build(prior)
    for source_slide in list(prior.slides)[33:47]:
        clone_slide(prs, source_slide)
    close_slide(prs)


def main():
    HERE.mkdir(parents=True, exist_ok=True)
    prs = m.dk.blank_deck()
    build(prs)
    m.dk.lint_layout(prs, strict=True)
    prs.save(OUT)
    m.dk.declare_delivery(str(OUT), "presented")
    print(f"saved -> {OUT} | slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
