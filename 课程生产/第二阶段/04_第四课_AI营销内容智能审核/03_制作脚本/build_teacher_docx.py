from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


NAVY = "10262F"
TEAL = "12B8A6"
ORANGE = "F06B45"
LIGHT_TEAL = "E8F7F5"
LIGHT_GRAY = "F2F5F6"
MID_GRAY = "5B6870"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Hiragino Sans GB", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_rich_text(paragraph, text: str, size=10.5, color=NAVY, bold=False) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Menlo", size=max(size - 0.5, 8), bold=False, color=ORANGE)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold, color=color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("AI营销内容智能审核项目实战  ·  ")
    set_run_font(run, size=8, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Hiragino Sans GB"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.32

    for style_name, size, color in (
        ("Title", 30, WHITE),
        ("Heading 1", 22, NAVY),
        ("Heading 2", 15, TEAL),
        ("Heading 3", 12, ORANGE),
    ):
        style = styles[style_name]
        style.font.name = "Hiragino Sans GB"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    styles["Heading 1"].paragraph_format.space_before = Pt(8)
    styles["Heading 1"].paragraph_format.space_after = Pt(12)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(6)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    styles["Heading 3"].paragraph_format.space_before = Pt(7)
    styles["Heading 3"].paragraph_format.space_after = Pt(4)
    styles["Heading 3"].paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    sect_pr = section._sectPr
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "30" if edge == "left" else "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), TEAL if edge == "left" else NAVY)
        pg_borders.append(border)
    sect_pr.append(pg_borders)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    r = p.add_run("SECOND STAGE · PROJECT COURSE 04")
    set_run_font(r, size=10, bold=True, color=TEAL)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("AI营销内容智能审核")
    set_run_font(r, size=30, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("项目实战")
    set_run_font(r, size=30, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("老师版完整讲义")
    set_run_font(r, size=16, bold=True, color=ORANGE)

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(3.5)
    table.columns[1].width = Cm(11.5)
    rows = [
        ("课程定位", "第二阶段公开项目课"),
        ("课程时长", "180 分钟（含两次休息）"),
        ("项目产出", "AI营销内容智能审核平台方案"),
        ("授课主线", "业务 → 方案 → 规则 → 产品 → 评测"),
    ]
    for i, (label, value) in enumerate(rows):
        for j, txt in enumerate((label, value)):
            cell = table.cell(i, j)
            set_cell_shading(cell, LIGHT_TEAL if j == 0 else LIGHT_GRAY)
            set_cell_margins(cell, 140, 150, 140, 150)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            add_rich_text(p, txt, size=10, color=TEAL if j == 0 else NAVY, bold=(j == 0))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    r = p.add_run("核心目标")
    set_run_font(r, size=10, bold=True, color=TEAL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("让学员完成一条可试点、可展示、可复盘的智能审核闭环。")
    set_run_font(r, size=14, bold=True, color=NAVY)
    new_section = doc.add_section(WD_SECTION.NEW_PAGE)
    new_section.page_width = Cm(21)
    new_section.page_height = Cm(29.7)
    new_section.top_margin = Cm(2)
    new_section.bottom_margin = Cm(1.8)
    new_section.left_margin = Cm(2.2)
    new_section.right_margin = Cm(2.2)
    new_section.header_distance = Cm(0.8)
    new_section.footer_distance = Cm(0.8)


def add_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        parts = [p.strip() for p in line.strip().strip("|").split("|")]
        rows.append(parts)
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", p.replace(" ", "")) for p in rows[1]):
        rows.pop(1)
    if not rows:
        return
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(width):
            cell = table.cell(i, j)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                set_cell_shading(cell, NAVY)
            elif i % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            txt = row[j] if j < len(row) else ""
            p = cell.paragraphs[0]
            add_rich_text(p, txt, size=9, color=WHITE if i == 0 else NAVY, bold=(i == 0))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, 180, 220, 180, 220)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("\n".join(lines))
    set_run_font(r, name="Menlo", size=9, color=WHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_quote(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_TEAL)
    set_cell_margins(cell, 160, 220, 160, 220)
    p = cell.paragraphs[0]
    p.paragraph_format.left_indent = Cm(0.1)
    add_rich_text(p, text, size=10.5, color=NAVY)


def add_figure(doc: Document, image_path: Path, alt_text: str = "") -> None:
    if not image_path.exists():
        raise FileNotFoundError(f"Figure not found: {image_path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(16.4))
    if alt_text:
        doc_pr = run._r.xpath(".//wp:docPr")
        if doc_pr:
            doc_pr[0].set("descr", alt_text)


def build(markdown_path: Path, output_path: Path) -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    # The custom cover already contains the title and course metadata.
    # Skip the Markdown front matter through its first horizontal rule.
    i = 0
    while i < len(lines) and lines[i].strip() != "---":
        i += 1
    i += 1
    first_h1 = True
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            add_code_block(doc, block)
            i += 1
            continue
        image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", stripped)
        if image_match:
            alt_text, image_ref = image_match.groups()
            add_figure(doc, (markdown_path.parent / image_ref).resolve(), alt_text)
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue
        if stripped.startswith("# "):
            start_on_new_page = not first_h1
            first_h1 = False
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.page_break_before = start_on_new_page
            add_rich_text(p, stripped[2:], size=22, color=NAVY, bold=True)
            i += 1
            continue
        if stripped.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            add_rich_text(p, stripped[3:], size=15, color=TEAL, bold=True)
            i += 1
            continue
        if stripped.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            add_rich_text(p, stripped[4:], size=12, color=ORANGE, bold=True)
            i += 1
            continue
        if stripped.startswith("> "):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_quote(doc, " ".join(quote_lines))
            continue
        bullet = re.match(r"^[-*] (.+)$", stripped)
        numbered = re.match(r"^\d+\. (.+)$", stripped)
        if bullet or numbered:
            style = "List Bullet" if bullet else "List Number"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            add_rich_text(p, (bullet or numbered).group(1), size=10.5)
            i += 1
            continue

        para_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                break
            if nxt.startswith(("#", "|", ">", "```", "- ", "![")) or re.match(r"^\d+\. ", nxt):
                break
            para_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        add_rich_text(p, " ".join(para_lines), size=10.5)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.input, args.output)
