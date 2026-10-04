#!/usr/bin/env python3
"""Extend the approved 33-slide light deck through sections five to eight."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
COURSE = ROOT.parent
STAGE = ROOT.parents[1]
SOURCE_DECK = STAGE / "99_历史归档" / "课件" / "第二课_电商智能客服项目实战_过渡版本" / "第二课_电商智能客服项目实战_33页过渡版.pptx"
LECTURE = COURSE / "01_讲义与材料" / "04_第二课_电商智能客服项目实战_老师版完整讲义_v1.0.md"
MATERIAL = COURSE / "01_讲义与材料" / "05_第二课_项目案例与作业材料包_v1.0.md"
OUTPUT = ROOT / "第二课_电商智能客服项目实战_浅色完整PPT_v1.0.pptx"
LEDGER_PATH = ROOT / ".append-ledger.json"

W, H = 13.333, 7.5
TOTAL = 59
FONT = "Microsoft YaHei"
MONO = "Menlo"
WHITE = "FFFFFF"
BG = "F7F9FC"
INK = "1A2230"
TEXT = "24292E"
MUTED = "59677A"
FAINT = "8B97A8"
NAVY = "123A78"
BLUE = "2F6FD0"
PALE_BLUE = "E9F0FA"
SOFT_BLUE = "F2F6FC"
LINE = "D6DCE5"
RED = "C0392B"
PALE_RED = "FDF2F0"
GREEN = "2E7D5B"
PALE_GREEN = "EDF7F2"
AMBER = "B36A16"
PALE_AMBER = "FFF6E8"

LEDGER: list[dict] = []
lecture_text = LECTURE.read_text(encoding="utf-8")


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_text(slide, x, y, w, h, text, size=16, color=TEXT, bold=False,
             font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.02, line_spacing=1.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(0.8)
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except (IndexError, ValueError):
            pass
    return shape


def add_line(slide, x1, y1, x2, y2, color=LINE, width=1.2):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2),
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def add_arrow(slide, x, y, w=0.28, color=BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(0.28)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def set_notes(slide, section: str, headings: tuple[str, ...], extra=""):
    start = lecture_text.index(f"# {section}：")
    end = lecture_text.find("\n# ", start + 3)
    body = lecture_text[start:end if end >= 0 else len(lecture_text)]
    pieces = []
    for heading in headings:
        marker = f"## {heading}"
        pos = body.find(marker)
        if pos < 0:
            continue
        stop = body.find("\n## ", pos + len(marker))
        pieces.append(body[pos:stop if stop >= 0 else len(body)])
    notes = "\n\n".join(pieces) if pieces else body.split("\n## ", 1)[0]
    if section == "第七节":
        notes = notes.replace("2026 年 9 月 5 日", "2026 年 9 月 4 日")
    notes += f"\n\n{extra}" if extra else ""
    notes += f"\n\n讲义来源：{LECTURE.name}\n配套材料：{MATERIAL.name}"
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    if text_frame is None:
        body = parse_xml("""
          <p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
                xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
            <p:nvSpPr>
              <p:cNvPr id="2" name="Notes Placeholder 1"/>
              <p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>
              <p:nvPr><p:ph type="body" idx="3" sz="quarter"/></p:nvPr>
            </p:nvSpPr>
            <p:spPr/>
            <p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody>
          </p:sp>
        """)
        notes_slide.element.cSld.spTree.insert_element_before(body, "p:extLst")
        text_frame = notes_slide.notes_text_frame
    text_frame.text = notes


def add_footer(slide, number: int):
    add_text(slide, 0.42, 6.98, 3.0, 0.2,
             "第二课 · 电商智能客服项目实战", 10.5, FAINT)
    add_text(slide, 11.95, 6.98, 0.97, 0.2,
             f"{number:02d} / {TOTAL}", 10.5, FAINT, align=PP_ALIGN.RIGHT)


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    bg = add_rect(slide, 0, 0, W, H, BG, BG, radius=False)
    bg.line.fill.background()
    return slide


def content_slide(prs, number: int, section: str, title: str, subtitle: str,
                  headings: tuple[str, ...], form: str, extra_notes=""):
    slide = blank_slide(prs)
    add_text(slide, 0.42, 0.27, 5.1, 0.22, section, 10.5, NAVY, True)
    add_text(slide, 0.42, 0.55, 11.1, 0.52, title, 24, INK, True)
    add_text(slide, 0.42, 1.03, 11.1, 0.3, subtitle, 12.5, MUTED)
    add_line(slide, 11.92, 0.91, 12.92, 0.91, BLUE, 2.0)
    add_footer(slide, number)
    set_notes(slide, section.split(" · ")[0], headings, extra_notes)
    LEDGER.append({
        "slide": number,
        "section": section.split(" · ")[0],
        "title": title,
        "headings": list(headings),
        "form": form,
    })
    return slide


def chapter_slide(prs, number, chapter_no, section, title, subtitle, bullets, page_range):
    slide = blank_slide(prs)
    add_rect(slide, 0, 0, 0.17, H, NAVY, NAVY, radius=False)
    add_text(slide, 0.92, 2.32, 2.8, 1.18, chapter_no, 68, NAVY, True,
             font=MONO, valign=MSO_ANCHOR.MIDDLE)
    add_line(slide, 0.92, 4.46, 3.0, 4.46, BLUE, 2.0)
    add_text(slide, 0.92, 4.65, 2.92, 0.28, page_range, 11, FAINT, font=MONO)
    add_text(slide, 4.42, 1.45, 8.1, 1.2, title, 34, INK, True,
             valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 4.42, 2.78, 8.05, 0.7, subtitle, 15, MUTED)
    for idx, bullet in enumerate(bullets):
        add_rect(slide, 4.42, 3.78 + idx * 0.7, 0.10, 0.10, BLUE, BLUE, radius=False)
        add_text(slide, 4.66, 3.68 + idx * 0.7, 7.85, 0.42, bullet, 15.5, TEXT)
    set_notes(slide, section, ())
    LEDGER.append({"slide": number, "section": section, "title": title,
                   "headings": [], "form": "chapter"})
    return slide


def card(slide, x, y, w, h, title, body, accent=BLUE, fill=WHITE,
         title_size=16, body_size=14.5):
    add_rect(slide, x, y, w, h, fill, LINE)
    add_rect(slide, x, y, 0.08, h, accent, accent, radius=False)
    add_text(slide, x + 0.23, y + 0.18, w - 0.42, 0.35, title,
             title_size, INK, True)
    add_text(slide, x + 0.23, y + 0.72, w - 0.42, h - 0.9, body,
             body_size, TEXT)


def cards(slide, items, y=1.65, h=4.85):
    gap = 0.27
    total_w = 12.49
    w = (total_w - gap * (len(items) - 1)) / len(items)
    for idx, (title, body, accent, fill) in enumerate(items):
        card(slide, 0.42 + idx * (w + gap), y, w, h, title, body, accent, fill)


def table(slide, rows, widths, y=1.62, h=4.95, font_size=14):
    x = 0.42
    total_w = 12.49
    row_h = h / len(rows)
    for r, row in enumerate(rows):
        cy = y + r * row_h
        cx = x
        for c, value in enumerate(row):
            cw = total_w * widths[c] / sum(widths)
            fill = NAVY if r == 0 else (WHITE if r % 2 else SOFT_BLUE)
            color = WHITE if r == 0 else TEXT
            add_rect(slide, cx, cy, cw, row_h, fill, LINE, radius=False)
            add_text(slide, cx + 0.12, cy + 0.06, cw - 0.24, row_h - 0.12,
                     value, font_size if r else font_size - 0.5, color,
                     bold=(r == 0), valign=MSO_ANCHOR.MIDDLE)
            cx += cw


def flow(slide, steps, y=2.15, box_h=2.7):
    gap = 0.40
    total_w = 12.10
    w = (total_w - gap * (len(steps) - 1)) / len(steps)
    for idx, (title, body) in enumerate(steps):
        x = 0.52 + idx * (w + gap)
        add_text(slide, x, y - 0.43, 0.5, 0.28, f"{idx + 1:02d}", 11, BLUE, True, font=MONO)
        card(slide, x, y, w, box_h, title, body, BLUE, WHITE, 15, 13.5)
        if idx < len(steps) - 1:
            add_arrow(slide, x + w + 0.06, y + 1.17, 0.24)


def chat(slide, y, speaker, message, state, accent=BLUE):
    add_text(slide, 0.55, y, 0.72, 0.3, speaker, 11.5, accent, True)
    add_rect(slide, 1.37, y - 0.07, 6.7, 0.92, WHITE, LINE)
    add_text(slide, 1.58, y + 0.12, 6.27, 0.5, message, 15.5, TEXT, True,
             valign=MSO_ANCHOR.MIDDLE)
    add_rect(slide, 8.36, y - 0.07, 4.55, 0.92, PALE_BLUE, PALE_BLUE)
    add_text(slide, 8.58, y + 0.08, 4.08, 0.58, state, 13.5, NAVY,
             valign=MSO_ANCHOR.MIDDLE)


def update_existing_pages(prs):
    page_pattern = re.compile(r"^\s*\d{1,2}\s*/\s*42\s*$")
    for number in range(1, 33):
        slide = prs.slides[number - 1]
        for shape in slide.shapes:
            if not shape.has_text_frame or not page_pattern.match(shape.text):
                continue
            paragraph = shape.text_frame.paragraphs[0]
            if paragraph.runs:
                paragraph.runs[0].text = f"{number:02d} / {TOTAL}"
            else:
                paragraph.text = f"{number:02d} / {TOTAL}"
            for run in paragraph.runs:
                run.font.name = FONT
                run.font.size = Pt(10.5)
                run.font.color.rgb = rgb(FAINT)
        # Preserve chapter and interaction pages that intentionally had no footer.

    transition = prs.slides[32]
    replacements = {
        "第二课 · 收束": "第二课 · 前四节小结",
        "把用户的自然语言，接到确定的业务流程上。":
            "我们已经把售前、查单与售后，接进了确定的业务流程。",
    }
    for shape in transition.shapes:
        if not shape.has_text_frame:
            continue
        current = shape.text.strip()
        if current not in replacements:
            continue
        paragraph = shape.text_frame.paragraphs[0]
        if paragraph.runs:
            paragraph.runs[0].text = replacements[current]
        else:
            paragraph.text = replacements[current]
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = FONT
    add_text(transition, 2.1, 3.13, 9.1, 0.9,
             "接下来补齐推荐、流程协同、人工接管与个人项目交付。",
             17, "DCE9FF", align=PP_ALIGN.CENTER,
             valign=MSO_ANCHOR.MIDDLE)
    add_text(transition, 11.96, 6.98, 0.96, 0.2,
             f"33 / {TOTAL}", 10.5, "C5D5ED", align=PP_ALIGN.RIGHT)


def build():
    ROOT.mkdir(parents=True, exist_ok=True)
    prs = Presentation(SOURCE_DECK)
    assert len(prs.slides) == 33, "Expected the 33-slide light source deck"
    assert round(prs.slide_width / Inches(1), 3) == 13.333
    assert round(prs.slide_height / Inches(1), 3) == 7.5
    update_existing_pages(prs)

    chapter_slide(prs, 34, "05", "第五节", "商品推荐与闲聊兜底",
                  "补齐购买前服务，并为无法进入明确业务流程的表达设计出口。",
                  ["理解用户需求并筛选商品", "接住用户的后续调整", "处理寒暄、模糊表达和超范围请求"],
                  "P.34 – 40")

    s = content_slide(prs, 35, "第五节 · 商品推荐与闲聊兜底",
                      "商品推荐的需求收集",
                      "用户：想买个通勤用的耳机，预算三百以内。",
                      ("一、商品推荐与订单查询有什么不同", "三、推荐流程需要收集哪些信息"),
                      "cards")
    cards(s, [
        ("已经明确", "商品：耳机\n预算：300 元以内\n场景：通勤", BLUE, WHITE),
        ("关键追问", "更看重降噪，还是长时间佩戴舒适？\n\n每次只问最影响选择的问题。", AMBER, PALE_AMBER),
        ("形成查询条件", "保留预算与场景，补充核心偏好。\n\n不是重新盘问已经知道的信息。", GREEN, PALE_GREEN),
    ], h=4.65)

    s = content_slide(prs, 36, "第五节 · 商品推荐与闲聊兜底",
                      "模拟商品目录与筛选结果",
                      "筛选条件：300 元以内、地铁通勤、需要主动降噪。",
                      ("二、先准备一份模拟商品数据", "六、推荐结果不能脱离真实商品数据"),
                      "table")
    table(s, [
        ["商品", "价格", "佩戴方式", "主动降噪", "单次续航"],
        ["通勤降噪款", "269 元", "入耳式", "支持", "7 小时"],
        ["舒适轻听款", "199 元", "半入耳式", "不支持", "8 小时"],
        ["深度降噪款", "329 元", "入耳式", "支持", "6 小时"],
        ["运动稳固款", "239 元", "耳挂式", "不支持", "9 小时"],
    ], [2.7, 1.35, 2.15, 2.0, 1.7], h=4.25, font_size=13.5)
    add_rect(s, 0.42, 6.04, 12.49, 0.52, PALE_BLUE, PALE_BLUE)
    add_text(s, 0.66, 6.15, 12.0, 0.28,
             "当前符合全部条件：通勤降噪款。商品、价格与参数均为教学模拟。",
             13.5, NAVY, True)

    s = content_slide(prs, 37, "第五节 · 商品推荐与闲聊兜底",
                      "商品推荐工作流",
                      "需求理解、商品查询、条件筛选与推荐解释各司其职。",
                      ("四、展开商品推荐工作流", "六、推荐结果不能脱离真实商品数据"),
                      "flow")
    flow(s, [
        ("理解需求", "提取预算与场景，判断是否需要追问。"),
        ("查询商品", "从商品数据获取价格、参数和库存。"),
        ("筛选比较", "先满足硬条件，再比较偏好匹配。"),
        ("解释推荐", "说明推荐理由，同时说明关键差异。"),
    ])
    add_rect(s, 0.52, 5.35, 12.1, 0.76, PALE_BLUE, PALE_BLUE)
    add_text(s, 0.78, 5.55, 11.55, 0.36,
             "“269 元、支持主动降噪，符合你的地铁通勤需求。”",
             16, NAVY, True, align=PP_ALIGN.CENTER)

    s = content_slide(prs, 38, "第五节 · 商品推荐与闲聊兜底",
                      "推荐后的条件更新",
                      "用户：有没有便宜一点的？",
                      ("五、用户继续追问时怎样处理",), "comparison",
                      "当前商品表没有更便宜且支持主动降噪的型号。应说明没有匹配项，询问用户是否愿意调整条件。")
    cards(s, [
        ("继续保留", "通勤场景\n300 元以内\n需要主动降噪\n当前候选商品列表", BLUE, WHITE),
        ("本轮更新", "新增偏好：价格更低。\n\n重新查询或排序，不重复询问已知条件。", AMBER, PALE_AMBER),
        ("没有匹配项", "明确说明当前无结果。\n\n询问是否调整预算或降噪条件，不能静默推荐不符合要求的商品。", RED, PALE_RED),
    ], h=4.65)

    s = content_slide(prs, 39, "第五节 · 商品推荐与闲聊兜底",
                      "闲聊、模糊表达与超范围请求",
                      "先结合上下文判断，再决定回应、澄清、引导或转人工。",
                      ("七、什么情况下进入闲聊兜底", "八、闲聊兜底与转人工不是一回事"),
                      "table")
    table(s, [
        ["用户表达", "判断", "处理方式"],
        ["你好。", "简单寒暄", "简短回应，并说明服务范围"],
        ["这个怎么办？", "对象或诉求不明确", "先查上下文，必要时澄清"],
        ["与电商服务无关的问题", "超出业务范围", "说明边界，引导回到业务"],
        ["直接找人工。", "明确要求人工", "进入人工处理路径"],
    ], [2.8, 3.0, 4.1], h=4.8, font_size=13.5)

    s = content_slide(prs, 40, "第五节 · 商品推荐与闲聊兜底",
                      "推荐与兜底的测试重点",
                      "检查真实数据、用户条件与下一步处理是否一致。",
                      ("九、用几组对话检查推荐与兜底", "十、本节完成了什么"),
                      "table")
    table(s, [
        ["测试情况", "重点检查"],
        ["只说“推荐个耳机”", "追问最影响结果的条件，不一次盘问全部字段"],
        ["说“便宜一点”或商品缺货", "保留原有条件，更新查询与推荐结果"],
        ["价格、库存或商品能力没有资料", "不编造；说明当前无法确认"],
        ["说“这个不太行”", "结合上一轮对象理解，必要时再澄清"],
    ], [3.8, 6.1], h=4.8, font_size=13.5)

    chapter_slide(prs, 41, "06", "第六节", "多流程协同与上下文",
                  "用户不会按流程图聊天；产品必须在任务切换时保留已有信息。",
                  ["区分继续、切换与返回", "保存对象、条件和未完成操作", "处理多任务与业务状态"],
                  "P.41 – 46")

    s = content_slide(prs, 42, "第六节 · 多流程协同与上下文",
                      "跨流程对话：推荐 → 查单 → 返回推荐",
                      "一次会话中，多条业务流程会来回切换。",
                      ("二、先看一段跨流程对话", "六、继续、切换和返回，需要分别处理"),
                      "dialogue")
    chat(s, 1.82, "用户", "想买个通勤耳机，预算三百以内。", "进入推荐；保存条件与候选商品")
    chat(s, 3.18, "用户", "对了，我上周买的那个怎么还没到？", "切到订单查询；保留刚才的推荐", AMBER)
    chat(s, 4.54, "用户", "刚才推荐的第一个，适合跑步吗？", "返回推荐；识别的是商品，不是订单", GREEN)

    s = content_slide(prs, 43, "第六节 · 多流程协同与上下文",
                      "会话中需要保存的业务信息",
                      "聊天原文与整理后的业务状态，共同支持后续处理。",
                      ("三、系统需要记住哪些信息", "四、为什么不能只把整段聊天记录交给模型"),
                      "table")
    table(s, [
        ["信息类型", "业务举例", "用途"],
        ["当前任务与进度", "订单查询，等待用户选择", "判断下一句话接在哪里"],
        ["最新条件", "预算从 300 改为 400", "使用修改后的条件"],
        ["对话中的对象", "推荐列表、候选订单", "理解“第一个”“刚才那个”"],
        ["未完成操作", "售后待确认、材料未补齐", "恢复暂停的处理进度"],
    ], [2.6, 4.1, 3.2], h=4.8, font_size=13)

    s = content_slide(prs, 44, "第六节 · 多流程协同与上下文",
                      "主流程每轮的处理步骤",
                      "主流程负责衔接，业务子流程负责完成具体任务。",
                      ("五、主流程每一轮需要做什么",), "flow")
    flow(s, [
        ("读取状态", "当前任务、对象与未完成操作。"),
        ("理解本轮", "判断继续、修改、切换或返回。"),
        ("调用流程", "传入已有信息，处理具体任务。"),
        ("更新结果", "保存最新状态，并明确下一步。"),
    ])
    add_rect(s, 0.52, 5.36, 12.1, 0.76, PALE_BLUE, PALE_BLUE)
    add_text(s, 0.75, 5.55, 11.6, 0.4,
             "子流程返回：处理结果、更新的信息、正在等待什么、任务是否完成。",
             14.5, NAVY, True, align=PP_ALIGN.CENTER)

    s = content_slide(prs, 45, "第六节 · 多流程协同与上下文",
                      "中途切换与幂等保护",
                      "提交结果尚未确认时，保留原请求并核实状态。",
                      ("七、中途切换时，哪些操作可以暂停", "十、用三段对话检查流程协同"),
                      "comparison")
    cards(s, [
        ("可以暂停的对话步骤", "收集推荐条件\n选择候选订单\n解释政策\n提交前补充信息", BLUE, WHITE),
        ("需要持续跟踪的操作", "申请正在提交\n接口返回超时\n状态尚未确认", AMBER, PALE_AMBER),
        ("保护原则", "保留请求标识并查询结果。\n\n不能因为超时，就直接创建第二份申请。", RED, PALE_RED),
    ], h=4.65)

    s = content_slide(prs, 46, "第六节 · 多流程协同与上下文",
                      "一句话多任务的顺序与系统组织",
                      "用户：查一下什么时候到，到了以后不合适能退吗？",
                      ("八、用户一句话里可能包含多个任务", "九、要不要为每条流程设计一个 Agent"),
                      "flow")
    flow(s, [
        ("先查订单", "定位具体商品，取得订单状态。"),
        ("再查政策", "使用对应商品，匹配适用要求。"),
        ("分别回复", "说明物流结果与售后条件。"),
    ], y=2.2)
    add_rect(s, 0.52, 5.34, 12.1, 0.82, PALE_BLUE, PALE_BLUE)
    add_text(s, 0.78, 5.53, 11.55, 0.46,
             "本项目：主流程负责识别与调度，业务子流程按明确步骤执行。",
             15.5, NAVY, True, align=PP_ALIGN.CENTER)

    chapter_slide(prs, 47, "07", "第七节", "人工接管与客服工作台",
                  "转人工不是失败后的补丁，而是完整服务流程的一部分。",
                  ["明确什么时候转人工", "让人工客服直接接着处理", "记录人工修正与最终结果"],
                  "P.47 – 52")

    s = content_slide(prs, 48, "第七节 · 人工接管与客服工作台",
                      "转人工的触发条件",
                      "根据用户诉求、可用依据、业务权限和系统状态判断。",
                      ("一、哪些问题应该转人工",), "table")
    table(s, [
        ["情况", "典型表现"],
        ["用户要求或沟通受阻", "明确要求人工；多次澄清仍无法确认诉求"],
        ["依据缺失或冲突", "找不到适用政策；业务系统状态不一致"],
        ["权限或特殊风险", "特殊补偿、重大投诉、账户安全问题"],
        ["系统无法继续", "接口持续失败；提交结果无法确认"],
    ], [3.8, 6.1], h=4.8, font_size=13.5)

    s = content_slide(prs, 49, "第七节 · 人工接管与客服工作台",
                      "转接状态与用户告知",
                      "让用户知道是否在排队、信息是否已交接，以及下一步是什么。",
                      ("二、转人工之前要告诉用户什么",), "flow",
                      "本案例未提供排队估时，不承诺具体等待分钟数。")
    flow(s, [
        ("准备转接", "整理订单、诉求与未完成操作。"),
        ("正在排队", "展示真实队列状态，保留会话信息。"),
        ("人工接入", "说明已经完成交接，继续处理问题。"),
    ], y=2.05)
    add_rect(s, 0.52, 5.15, 12.1, 1.02, PALE_RED, PALE_RED)
    add_text(s, 0.78, 5.33, 11.55, 0.3,
             "暂时无人接入：提供留言、服务单或其他联系渠道。", 14.5, RED, True)
    add_text(s, 0.78, 5.72, 11.55, 0.25,
             "预计等待时间仅在排队系统真实提供时展示。", 12.5, MUTED)

    s = content_slide(prs, 50, "第七节 · 人工接管与客服工作台",
                      "人工交接摘要示例",
                      "DEMO-101 · AirLite 白色耳机 · 9 月 4 日签收。",
                      ("三、人工客服需要收到哪些信息", "四、不能只把聊天记录扔给人工"),
                      "table", "签收日期统一为 9 月 4 日；完整聊天记录同时保留。")
    table(s, [
        ["字段", "本次交接内容"],
        ["当前诉求与故障", "换货；右耳无声音"],
        ["已经完成的步骤", "已确认订单；重新连接后仍未恢复"],
        ["仍缺少的信息", "故障材料尚未提供"],
        ["业务操作状态", "尚未提交售后申请"],
        ["转接原因与下一步", "用户要求人工；核实材料与后续受理方式"],
    ], [3.2, 6.7], h=4.9, font_size=12.8)

    s = content_slide(prs, 51, "第七节 · 人工接管与客服工作台",
                      "客服工作台页面结构",
                      "人工同时查看对话、业务信息，并执行有权限的操作。",
                      ("五、客服工作台应该怎样组织",), "wireframe")
    card(s, 0.42, 1.58, 12.49, 0.78, "用户与当前任务",
         "身份、问题类型、等待状态、重复咨询", BLUE, PALE_BLUE, 14, 12.5)
    card(s, 0.42, 2.58, 7.65, 2.65, "完整对话",
         "用户与 AI 的历史消息\n标出用户确认、业务操作与错误\n保留原文，方便人工核对上下文", BLUE, WHITE, 16, 14)
    card(s, 8.30, 2.58, 4.61, 2.65, "业务信息与摘要",
         "订单与申请状态\n已尝试操作\n仍未解决的问题", AMBER, PALE_AMBER, 16, 14)
    card(s, 0.42, 5.46, 12.49, 0.78, "人工操作区",
         "回复用户 · 查询 · 登记服务单 · 备注 · 结束会话", RED, PALE_RED, 14, 12.5)

    s = content_slide(prs, 52, "第七节 · 人工接管与客服工作台",
                      "AI 辅助、人工修正与结果回写",
                      "建议、人工决定与最终记录需要分别保存。",
                      ("六、AI 在人工工作台里还能做什么", "七、人工修改的信息怎样返回系统", "九、怎样检查人工接管是否设计完整"),
                      "cards")
    cards(s, [
        ("AI 辅助", "整理对话与字段\n查找适用政策\n建议下一步\n拟写回复", BLUE, WHITE),
        ("人工处理", "核对建议与依据\n确认必要信息\n执行有权限的操作\n记录最终结果", AMBER, PALE_AMBER),
        ("记录修正", "原判断：退货\n人工修正：价格保护\n\n补充测试案例，调整分类与知识。", GREEN, PALE_GREEN),
    ], h=4.65)

    chapter_slide(prs, 53, "08", "第八节", "个人项目与课后作业",
                  "把公共项目的方法迁移到自己的业务场景，并交付可运行证据。",
                  ["选择自己的业务场景", "用 Coze 跑通一条核心流程", "提交实际测试与演示，点评后完善"],
                  "P.53 – 59")

    s = content_slide(prs, 54, "第八节 · 个人项目与课后作业",
                      "个人项目的场景选择",
                      "使用自己的业务规则、数据和处理步骤，而不是只换名称。",
                      ("一、把今天的公共项目换成自己的业务场景",), "table")
    table(s, [
        ["场景", "可以完成的一条核心流程"],
        ["电商售后", "确认订单、收集故障、登记售后申请"],
        ["教育培训", "查询报名、核实调课要求、登记申请"],
        ["企业软件", "了解故障、查询解决方法、创建服务单"],
        ["家电服务", "确认设备、收集故障、登记维修需求"],
        ["物流服务", "查询运单、说明状态、登记异常跟进"],
        ["酒店住宿", "查询预订、核实变更条件、登记或转交"],
    ], [2.8, 7.1], h=5.05, font_size=12.5)

    s = content_slide(prs, 55, "第八节 · 个人项目与课后作业",
                      "个人项目的最小完成范围",
                      "家电维修示例：从故障描述到维修需求登记。",
                      ("二、这次只完成一条核心流程", "三、先准备材料，再去 Coze 搭建"),
                      "flow")
    flow(s, [
        ("确认设备", "将设备与故障对应到模拟数据。"),
        ("了解问题", "查询排查资料，记录已尝试操作。"),
        ("整理需求", "收集预约信息，并让用户确认。"),
        ("模拟登记", "创建测试记录，返回等待联系状态。"),
    ])
    add_rect(s, 0.52, 5.35, 12.1, 0.92, PALE_RED, PALE_RED)
    add_text(s, 0.76, 5.53, 11.6, 0.46,
             "最终要能运行：只回答“怎样报修”，还没有完成这条流程。",
             14.5, RED, True, align=PP_ALIGN.CENTER)

    s = content_slide(prs, 56, "第八节 · 个人项目与课后作业",
                      "项目作业的三项交付",
                      "三项都围绕同一个个人项目，缺一不可。",
                      ("四、最终提交什么，以及分几步完成",), "cards")
    cards(s, [
        ("个人项目方案", "用户与问题\n业务规则与边界\n核心流程图\n资料与模拟数据", BLUE, WHITE),
        ("可运行的 Coze 项目", "至少一条核心流程\n实际查询或处理数据\n完成必要信息收集\n返回明确结果", AMBER, PALE_AMBER),
        ("实际测试与演示", "三组真实运行记录\n3—5 分钟演示录屏\n说明已实现、模拟与未实现内容", GREEN, PALE_GREEN),
    ], h=4.65)

    s = content_slide(prs, 57, "第八节 · 个人项目与课后作业",
                      "三组实际测试的要求",
                      "每组记录：初始条件、输入、预期、实际结果、问题与修正。",
                      ("四、最终提交什么，以及分几步完成",), "table")
    table(s, [
        ["测试类型", "要看到的行为"],
        ["正常完成", "信息齐全，核心流程走到明确结束状态"],
        ["需求变化", "用户补充或修改信息，后续处理同步更新"],
        ["无法完成", "查询失败或条件不满足，说明状态与下一步"],
    ], [3.1, 6.8], y=1.85, h=3.55, font_size=14)
    add_rect(s, 0.42, 5.74, 12.49, 0.72, PALE_RED, PALE_RED)
    add_text(s, 0.68, 5.92, 12.0, 0.36,
             "最终版提交实际运行结果；预期答案与“待验证”不能替代测试。",
             14, RED, True, align=PP_ALIGN.CENTER)

    s = content_slide(prs, 58, "第八节 · 个人项目与课后作业",
                      "初版、点评与最终版",
                      "分阶段完成，最终交付能够运行和复核的项目证据。",
                      ("四、最终提交什么，以及分几步完成", "五、下次课怎么点评"),
                      "timeline")
    flow(s, [
        ("点评课前", "提交项目方案和 Coze 初版；带着实际测试结果或具体卡点。"),
        ("作业点评", "检查业务、流程与运行结果；确定优先修改项。"),
        ("点评之后", "跑通核心流程，补齐三组测试，并提交演示录屏。"),
    ], y=2.05, box_h=2.85)
    add_text(s, 0.58, 5.63, 12.0, 0.58,
             "老师提供数据、规则和搭建指引，并结合具体卡点答疑；提交时间与渠道另行通知。",
             13.5, MUTED, align=PP_ALIGN.CENTER)

    s = content_slide(prs, 59, "第八节 · 个人项目与课后作业",
                      "把公共项目的方法，迁移成自己的项目证据",
                      "现在，在聊天区写下你的个人项目方向。",
                      ("五、下次课怎么点评", "六、整课收尾"), "statement")
    add_text(s, 0.92, 1.88, 11.5, 1.15,
             "我想为哪类用户，解决什么具体问题；\n这次先把流程做到哪一步。",
             28, INK, True, align=PP_ALIGN.CENTER,
             valign=MSO_ANCHOR.MIDDLE)
    add_rect(s, 1.28, 3.55, 10.77, 1.22, PALE_BLUE, PALE_BLUE)
    add_text(s, 1.62, 3.82, 10.1, 0.68,
             "例：为购买家电的用户提供维修服务，\n完成从故障描述到维修需求登记的流程。",
             17, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, 1.2, 5.35, 10.95, 0.56,
             "课后交付：项目方案 · Coze 核心流程 · 实际测试 · 演示录屏",
             15, MUTED, align=PP_ALIGN.CENTER)
    add_line(s, 4.8, 6.22, 8.53, 6.22, BLUE, 2.2)

    assert len(prs.slides) == TOTAL
    prs.core_properties.title = "第二课｜电商智能客服项目实战｜浅色完整版"
    prs.core_properties.subject = "第二阶段公共项目一，八节完整授课版"
    prs.core_properties.comments = "基于用户提供的 33 页浅色底稿续写；业务数据均为教学模拟。"
    prs.save(OUTPUT)

    LEDGER_PATH.write_text(json.dumps({
        "source_deck": str(SOURCE_DECK),
        "source_sha256": hashlib.sha256(SOURCE_DECK.read_bytes()).hexdigest(),
        "lecture": str(LECTURE),
        "material": str(MATERIAL),
        "output": str(OUTPUT),
        "total_slides": TOTAL,
        "slides": LEDGER,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT)


if __name__ == "__main__":
    build()
