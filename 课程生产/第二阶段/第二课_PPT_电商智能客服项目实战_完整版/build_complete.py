#!/usr/bin/env python3
"""Append the approved final four sections to the delivered 28-slide deck."""
from pathlib import Path
import importlib.util
import hashlib
import json
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
BASE = ROOT.parent
OLD_DIR = BASE / '第二课_PPT_电商智能客服项目实战_前四节版'
OLD = OLD_DIR / '第二课_电商智能客服项目实战_前四节_正式版PPT_v1.0.pptx'
SOURCE = BASE / '04_第二课_电商智能客服项目实战_老师版完整讲义_v1.0.md'
MATERIAL = BASE / '05_第二课_项目案例与作业材料包_v1.0.md'
OUT = ROOT / '第二课_电商智能客服项目实战_完整PPT_v1.1.pptx'
spec = importlib.util.spec_from_file_location('lesson2_base', OLD_DIR / 'build_lesson2_customer_service.py')
old = importlib.util.module_from_spec(spec)
spec.loader.exec_module(old)
m, dk = old.m, old.dk
m.SLIDE_COUNT = 54
m.PARTS.update({29:('第五节','商品推荐与闲聊兜底'),36:('第六节','多流程衔接与上下文管理'),42:('第七节','人工接管与客服工作台'),48:('第八节','个人项目选题与课后作业')})
m.OUT = OUT
LEDGER = []

def section_text(section, *heads):
    text = SOURCE.read_text(encoding='utf-8')
    start = text.index('# '+section+'：')
    end = text.find('\n# ', start+3)
    body = text[start:end if end >= 0 else len(text)]
    if not heads:
        return body
    pieces = []
    for head in heads:
        pos = body.index('## '+head)
        stop = body.find('\n## ', pos+3)
        pieces.append(body[pos:stop if stop >= 0 else len(body)])
    return '\n\n'.join(pieces)

def note(slide, section, *heads, extra=''):
    detail = section_text(section, *heads)
    # Reconcile the two isolated examples with the established teaching dataset.
    if section == '第七节':
        detail = detail.replace('2026 年 9 月 5 日', '2026 年 9 月 4 日')
    dk.speaker_notes(slide, detail+'\n\n'+extra+'\n\n讲义来源：'+str(SOURCE)+'\n配套案例：'+str(MATERIAL))

def page(prs, number, title, subtitle, section, heads=(), form='comparison'):
    s = m.base(prs, number, title, subtitle)
    # Stronger, explicit Chinese chapter labels on appended slides.
    for sh in s.shapes:
        if sh.has_text_frame and sh.text.startswith(section+'｜'):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(16)
    LEDGER.append({'slide':number,'title':title,'section':section,'form':form,'headings':list(heads)})
    note(s, section, *heads)
    return s

def divider(prs, number, cn, num, title, items):
    s = dk.add_slide(prs)
    dk.slide_background(s, m.BG)
    dk.box(s, .28, 0, .16, m.H, fill=m.TEAL)
    m.txt(s,.95,.72,5,.45,cn,24,m.TEAL,True)
    m.txt(s,.95,1.6,1.6,1.2,num,66,m.TEAL,True,wrap=False)
    m.txt(s,2.9,1.65,9,1.35,title,35,m.TEXT,True)
    dk.box(s,2.9,3.16,2.2,.05,fill=m.CORAL)
    for i,item in enumerate(items):
        m.txt(s,2.9,3.7+i*.7,.45,.4,str(i+1).zfill(2),13,m.TEAL,True)
        m.txt(s,3.55,3.68+i*.7,8.2,.55,item,19,m.MUTE)
    m.txt(s,11.7,6.65,.7,.35,str(number),12,m.MUTE)
    dk.speaker_notes(s,section_text(cn).split('\n## ')[0]+'\n\n本节内容：'+'；'.join(items))
    LEDGER.append({'slide':number,'title':title.replace('\n',''),'section':cn,'form':'divider','headings':[]})

def columns(s, items, y=2.2, h=3.8):
    gap=.28
    w=(11.52-gap*(len(items)-1))/len(items)
    for i,(title,body,color) in enumerate(items):
        m.card(s,.82+i*(w+gap),y,w,h,title,body,None,color,20,17)

def tab(s, rows, widths, size=17, row_h=.72, y=2.12, highlight=None):
    dk.table(s,.82,y,11.52,rows,col_w=widths,header=True,highlight=highlight,
             size=size,row_h=row_h,head_c=m.TEAL,body_c=m.TEXT,rule_c=m.LINE,
             hi_fill=m.PANEL2,hi_c=m.CORAL,font='Hiragino Sans GB')

def line(s,y,label,body,color=None):
    color = color or m.TEAL
    m.txt(s,.94,y,2.4,.5,label,21,color,True)
    m.txt(s,3.6,y,8.6,.65,body,19,m.TEXT)

def flow(s, labels, y=3.0):
    gap=.25
    w=(11.5-gap*(len(labels)-1))/len(labels)
    for i,(title,body) in enumerate(labels):
        x=.9+i*(w+gap)
        m.txt(s,x,y-.6,w,.35,f'{i+1:02d}',14,m.TEAL,True)
        dk.box(s,x,y,w,.05,fill=m.TEAL)
        m.txt(s,x,y+.3,w,.55,title,20,m.TEXT,True)
        m.txt(s,x,y+1.05,w,1.4,body,16,m.MUTE)
        if i<len(labels)-1:
            dk.arrow(s,x+w+.025,y+.01,.18,.15,color=m.TEAL)

def chatrow(s,y,who,body,meaning,color=None):
    color=color or m.TEAL
    m.txt(s,.95,y,.7,.35,who,14,color,True)
    m.txt(s,1.85,y,6.3,.82,body,20,m.TEXT,True)
    m.txt(s,8.6,y,3.6,.86,meaning,16,m.MUTE)
    dk.box(s,.95,y+1.05,11.2,.013,fill=m.LINE)

def update_opening(prs):
    # Work from the existing PPTX rather than recreating its first 28 slides.
    cover=prs.slides[0]
    for sh in cover.shapes:
        if sh.has_text_frame and sh.text=='第二课｜从业务背景到订单查询与售后受理':
            sh.text_frame.paragraphs[0].runs[0].text='第二课｜电商客服方案、业务流程与个人项目实践'
    s=prs.slides[1]
    for sh in list(s.shapes):
        el=sh._element
        el.getparent().remove(el)
    m.chrome(s,2)
    m.heading(s,'课程结构','前四节展开核心业务，后四节补齐产品协同与个人项目')
    titles=['业务背景与项目任务','整体方案与主工作流','订单查询工作流','售后咨询与处理','商品推荐与闲聊兜底','多流程衔接与上下文','人工接管与客服工作台','个人项目与课后作业']
    cns=['第一节','第二节','第三节','第四节','第五节','第六节','第七节','第八节']
    for i,(cn,title) in enumerate(zip(cns,titles)):
        col,row=divmod(i,4)
        x=.92+col*5.93
        y=2.2+row*1.0
        m.txt(s,x,y,1.1,.44,cn,18,m.TEAL,True)
        m.txt(s,x+1.25,y,4.48,.55,title,20,m.TEXT,True)
    dk.speaker_notes(s,'本课共八节。课堂以方案讲解、流程拆解和案例推演为主。课后完成自己的项目，Coze 搭建是正式交付。点评前交初版，点评后交可运行最终版。')
    for idx,cn in [(2,'第一节'),(8,'第二节'),(17,'第三节'),(24,'第四节')]:
        for sh in prs.slides[idx].shapes:
            if sh.has_text_frame and sh.text.startswith('SECTION '):
                r=sh.text_frame.paragraphs[0].runs[0]
                r.text=cn
                r.font.size=Pt(19)
                sh.height=Inches(.42)
    for i,s in enumerate(prs.slides,1):
        for sh in s.shapes:
            if abs(sh.top/Inches(1)-7.11)<.006 and abs(sh.height/Inches(1)-.04)<.006:
                sh.width=Inches(11.82*i/m.SLIDE_COUNT)

def refine_text(prs):
    """Keep approved layouts; use paragraph spacing instead of empty text lines."""
    for n,s in enumerate(prs.slides,1):
        for sh in s.shapes:
            if n==15 and abs(sh.top/Inches(1)-4.18)<.02 and abs(sh.height/Inches(1)-.78)<.02:
                sh.height=Inches(.94)
            if not sh.has_text_frame:
                continue
            if n==15 and sh.text=='→ DEMO-102':
                sh.text_frame.paragraphs[0].runs[0].text='DEMO-102'
            if n==27 and sh.text.startswith('知识：排查'):
                sh.text_frame.paragraphs[0].runs[0].text='知识：排查、受理流程\n与特殊情况\n规则：信息完整性、\n重复申请与操作权限'
            if n==51 and sh.text.startswith('三组实际测试记录'):
                sh.text_frame.paragraphs[0].runs[0].text='三组实际测试记录\n3—5 分钟演示录屏\n说明已实现、模拟\n与未实现的内容'
            for p in list(sh.text_frame.paragraphs):
                if '\n\n' not in p.text:
                    continue
                assert len(p.runs)==1, 'Review mixed-format paragraph before splitting'
                chunks=p.text.split('\n\n')
                template=deepcopy(p._p)
                p.runs[0].text=chunks[0]
                anchor=p._p
                for chunk in chunks[1:]:
                    q=deepcopy(template)
                    anchor.addnext(q)
                    from pptx.text.text import _Paragraph
                    para=_Paragraph(q,p._parent)
                    para.runs[0].text=chunk
                    para.space_before=Pt((p.runs[0].font.size.pt or 17)*1.12)
                    anchor=q

def build():
    ROOT.mkdir(parents=True,exist_ok=True)
    prs=Presentation(OLD)
    assert len(prs.slides)==28
    update_opening(prs)

    divider(prs,29,'第五节','05','商品推荐与\n闲聊兜底',['理解用户需求并筛选商品','接住用户的后续调整','处理寒暄、模糊表达和超范围问题'])
    s=page(prs,30,'商品推荐的需求收集','用户：“想买个通勤用的耳机，预算三百以内。”','第五节',('一、商品推荐与订单查询有什么不同','三、推荐流程需要收集哪些信息'))
    columns(s,[('已经明确','商品：耳机\n预算：300 元以内\n场景：通勤',m.TEAL),('关键追问','更看重降噪，\n还是长时间佩戴舒适？\n\n每次问最影响选择的问题',m.CORAL),('形成查询条件','用户：主要坐地铁\n更希望降噪好一点\n\n保留预算，补充核心偏好',m.YELLOW)])

    s=page(prs,31,'模拟商品目录与筛选结果','条件：300 元以内、地铁通勤、需要主动降噪','第五节',('二、先准备一份模拟商品数据','六、推荐结果不能脱离真实商品数据'),form='table')
    tab(s,[['商品','价格','佩戴方式','主动降噪','单次续航'],['通勤降噪款','269 元','入耳式','支持','7 小时'],['舒适轻听款','199 元','半入耳式','不支持','8 小时'],['深度降噪款','329 元','入耳式','支持','6 小时'],['运动稳固款','239 元','耳挂式','不支持','9 小时']],[3.65,1.72,2.35,2.1,1.7],highlight=0)
    m.txt(s,.96,6.18,11.3,.48,'当前符合条件：通勤降噪款。商品与参数均为教学模拟。',19,m.TEAL,True)

    s=page(prs,32,'商品推荐工作流','需求理解、商品查询与推荐解释的分工','第五节',('四、展开商品推荐工作流','六、推荐结果不能脱离真实商品数据'),form='flow')
    flow(s,[('理解需求','提取预算、场景\n判断是否需要追问'),('查询商品','从商品系统获取\n价格、参数和库存'),('筛选比较','先满足硬条件\n再比较偏好匹配'),('解释推荐','说明推荐理由\n同时说明差异')])
    m.txt(s,.97,6.1,11.2,.6,'“269 元、支持主动降噪，符合你的地铁通勤需求。”',20,m.TEXT,True)

    s=page(prs,33,'推荐后的条件更新','用户：“有没有便宜一点的？”','第五节',('五、用户继续追问时怎样处理',),form='comparison')
    columns(s,[('继续保留','通勤场景\n300 元以内\n需要主动降噪\n当前候选商品列表',m.TEAL),('本轮更新','希望价格更低\n\n重新查询或排序\n不重新问已经知道的条件',m.CORAL)])
    note(s,'第五节','五、用户继续追问时怎样处理',extra='配套材料补充：当前商品表没有更便宜且支持主动降噪的型号。说明没有匹配项，询问是否愿意调整条件，不能静默推荐不降噪的款式。')

    s=page(prs,34,'闲聊与模糊表达的处理','先结合上下文理解，再决定如何继续','第五节',('七、什么情况下进入闲聊兜底','八、闲聊兜底与转人工不是一回事'),form='table')
    tab(s,[['用户表达','判断','处理方式'],['“你好。”','简单寒暄','简短回应，说明服务范围'],['“这个怎么办？”','对象或诉求不明确','先查上下文，必要时澄清'],['与电商服务无关的问题','超出业务范围','说明范围，引导回到业务'],['“直接找人工。”','明确要求人工','进入人工处理路径']],[3.35,3.15,5.02],row_h=.85)

    s=page(prs,35,'商品推荐与兜底的检查重点','把真实数据、用户条件和下一步处理对应起来','第五节',('九、用几组对话检查推荐与兜底','十、本节完成了什么'),form='table')
    tab(s,[['测试情况','检查重点'],['只说“推荐个耳机”','追问影响结果的条件，不一次盘问全部字段'],['说“便宜一点”，或商品缺货','保留原有条件，更新查询与推荐结果'],['价格或商品能力没有资料','不编造价格、库存和功能'],['说“这个不太行”','结合上一轮对象理解，必要时澄清']],[4.0,7.52],row_h=.85)

    divider(prs,36,'第六节','06','多流程衔接与\n上下文管理',['区分继续、切换与返回','保存对象、条件和未完成操作','处理多任务与业务状态'])
    s=page(prs,37,'跨流程对话的处理','一次会话中，商品推荐和订单查询可以来回切换','第六节',('二、先看一段跨流程对话','六、继续、切换和返回，需要分别处理'),form='dialogue')
    chatrow(s,2.2,'用户','想买个通勤耳机，预算三百以内。','进入推荐\n保存条件与候选商品')
    chatrow(s,3.55,'用户','对了，我上周买的那个怎么还没到？','切到订单查询\n保留刚才的推荐',m.CORAL)
    chatrow(s,4.9,'用户','刚才推荐的第一个，适合跑步吗？','回到商品推荐\n识别的是商品，不是订单',m.YELLOW)

    s=page(prs,38,'会话中需要保存的业务信息','聊天原文与整理后的业务状态共同支持后续处理','第六节',('三、系统需要记住哪些信息','四、为什么不能只把整段聊天记录交给模型'),form='table')
    tab(s,[['信息类型','示例','用途'],['当前任务与进度','订单查询，等待选择','判断下一句话接在哪里'],['最新条件','通勤，预算从 300 改为 400','使用修改后的条件'],['对话中的对象','推荐列表、候选订单列表','理解“第一个”“刚才那个”'],['未完成操作','售后待确认、材料未补齐','恢复暂停的处理进度']],[3.0,4.7,3.82],size=16.5,row_h=.85)

    s=page(prs,39,'主流程每轮的处理步骤','主流程负责衔接，子流程负责具体业务','第六节',('五、主流程每一轮需要做什么',),form='flow')
    flow(s,[('读取状态','当前任务\n对象与未完成操作'),('理解本轮','继续、修改\n切换或返回'),('调用流程','传入已有信息\n处理具体任务'),('更新结果','保存最新状态\n明确下一步')])
    m.txt(s,.96,6.15,11.3,.5,'子流程返回：处理结果、更新的信息、等待什么、是否完成。',19,m.TEAL,True)

    s=page(prs,40,'业务操作中途切换的处理','提交结果尚未确认时，保留原请求并核实状态','第六节',('七、中途切换时，哪些操作可以暂停','十、用三段对话检查流程协同'),form='comparison')
    columns(s,[('可以暂停的对话步骤','收集推荐条件\n选择候选订单\n解释政策\n提交前补充信息',m.TEAL),('需要跟踪的业务操作','申请正在提交\n接口返回超时\n\n保留请求标识，查询结果\n不直接创建第二份申请',m.CORAL)])

    s=page(prs,41,'多任务的顺序与系统组织','用户：“查一下什么时候到，到了以后不合适能退吗？”','第六节',('八、用户一句话里可能包含多个任务','九、要不要为每条流程设计一个 Agent'),form='flow')
    flow(s,[('先查订单','定位具体商品\n取得订单状态'),('再查政策','使用对应商品\n匹配适用要求'),('分别回复','物流查询结果\n适用售后说明')],y=2.85)
    m.txt(s,.96,5.85,11.3,.95,'本项目组织方式：主流程识别与调度，业务子流程按明确步骤执行。',20,m.TEXT,True)

    divider(prs,42,'第七节','07','人工接管与\n客服工作台',['明确何时转人工','让人工直接接着处理','记录人工修正与最终结果'])
    s=page(prs,43,'转人工的触发条件','根据用户诉求、可用依据、业务权限和系统状态判断','第七节',('一、哪些问题应该转人工',),form='table')
    tab(s,[['情况','典型表现'],['用户要求或沟通受阻','明确找人工；多次澄清仍不清楚'],['依据缺失或冲突','找不到适用政策；业务系统状态不一致'],['权限或特殊风险','特殊补偿、重大投诉、账户安全问题'],['系统无法继续','接口持续失败；提交结果无法确认']],[4.2,7.32],row_h=.87)

    s=page(prs,44,'转接状态与用户告知','让用户知道是否在排队、信息是否已交接、下一步是什么','第七节',('二、转人工之前要告诉用户什么',),form='flow')
    flow(s,[('准备转接','整理订单、诉求\n与未完成操作'),('正在排队','展示真实队列状态\n保留会话信息'),('人工接入','说明已完成交接\n继续处理问题')],y=2.65)
    m.txt(s,.96,5.6,11.1,.6,'暂时无人接入：提供留言、服务单或其他联系渠道。',20,m.CORAL,True)
    m.txt(s,.96,6.25,11.1,.4,'预计等待时间仅在排队系统提供时展示。',16,m.MUTE)
    note(s,'第七节','二、转人工之前要告诉用户什么',extra='材料包统一口径：本案例未提供排队估时，不使用原讲义示例中的3—5分钟承诺。')

    s=page(prs,45,'人工交接信息示例','订单编号为 DEMO-101 的 AirLite 白色耳机，9 月 4 日签收','第七节',('三、人工客服需要收到哪些信息','四、不能只把聊天记录扔给人工'),form='table')
    tab(s,[['字段','本次交接内容'],['当前诉求与故障','换货；右耳无声音'],['已完成的步骤','已确认订单；重新连接后仍未恢复'],['缺少的信息','故障材料尚未提供'],['业务操作状态','尚未提交售后申请'],['转接原因与下一步','用户要求人工；核实材料与后续受理方式']],[3.7,7.82],size=17,row_h=.71)
    note(s,'第七节','三、人工客服需要收到哪些信息','四、不能只把聊天记录扔给人工',extra='签收日期统一采用前文与材料包的9月4日。完整聊天记录同时保留，摘要用于快速定位。')

    s=page(prs,46,'客服工作台的页面结构','人工同时查看对话、业务信息，并执行有权限的操作','第七节',('五、客服工作台应该怎样组织',),form='wireframe')
    m.card(s,.85,2.03,11.55,.78,'用户与当前任务','',None,m.TEAL,18,17)
    m.txt(s,5.0,2.22,7,.36,'身份、问题类型、等待时长、重复咨询',16,m.MUTE)
    m.card(s,.85,3.0,7.05,2.45,'完整对话','用户与 AI 的历史消息\n标出用户确认、业务操作与错误',None,m.TEAL,20,18)
    m.card(s,8.13,3.0,4.27,2.45,'业务信息与摘要','订单与申请状态\n已尝试操作、未解决问题',None,m.YELLOW,19,17)
    m.card(s,.85,5.64,11.55,.95,'人工操作区','',None,m.CORAL,18,17)
    m.txt(s,4.0,5.94,8.0,.45,'回复用户、查询、登记服务单、备注、结束会话',16,m.TEXT)

    s=page(prs,47,'人工处理中的 AI 辅助与反馈','建议、人工决定与最终记录分别保存','第七节',('六、AI 在人工工作台里还能做什么','七、人工修改的信息怎样返回系统','九、怎样检查人工接管是否设计完整'),form='comparison')
    columns(s,[('AI 辅助','整理对话与字段\n查找适用政策\n建议下一步\n拟写回复',m.TEAL),('人工处理','核对建议与依据\n确认必要信息\n执行有权限的操作\n记录最终结果',m.CORAL),('记录修正','原判断：退货\n人工修正：价格保护\n\n补充测试案例\n调整分类与知识',m.YELLOW)])

    divider(prs,48,'第八节','08','个人项目选题与\n课后作业',['选择自己的业务场景','用 Coze 跑通一条核心流程','提交实际测试与演示，点评后完善'])
    s=page(prs,49,'个人项目的场景选择','使用自己的业务规则、数据与处理步骤','第八节',('一、把今天的公共项目换成自己的业务场景',),form='table')
    tab(s,[['场景','可以完成的核心流程'],['电商售后','确认订单、收集故障、登记售后申请'],['教育培训','查询报名、核实调课要求、登记申请'],['企业软件','了解故障、查询解决方法、创建服务单'],['家电服务','确认设备、收集故障、登记维修需求'],['物流服务','查询运单、说明状态、登记异常跟进'],['酒店住宿','查询预订、核实变更条件、登记或转交']],[3.1,8.42],size=16.5,row_h=.60)

    s=page(prs,50,'个人项目的完成范围','家电维修示例：从故障描述到维修需求登记','第八节',('二、这次只完成一条核心流程','三、先准备材料，再去 Coze 搭建'),form='flow')
    flow(s,[('确认设备','设备与故障\n对应到模拟数据'),('了解问题','查询排查资料\n记录已尝试操作'),('整理需求','收集预约信息\n用户确认'),('模拟登记','创建测试记录\n返回等待联系状态')],y=2.8)
    m.txt(s,.96,5.85,11.25,.55,'最终要能运行：仅回答“怎样报修”，还没有完成这条流程。',20,m.CORAL,True)
    m.txt(s,.96,6.48,11.25,.35,'准备材料：业务资料、模拟数据、流程图。',17,m.MUTE)

    s=page(prs,51,'项目作业的三项交付','三项均需完成，围绕同一个个人项目','第八节',('四、最终提交什么，以及分几步完成',),form='comparison')
    columns(s,[('个人项目方案','用户与问题\n业务规则与范围\n核心流程图\n资料与模拟数据',m.TEAL),('可运行的 Coze 项目','至少一条核心流程\n实际查询或处理数据\n完成必要的信息收集\n返回明确结果',m.CORAL),('实际测试与演示','三组实际测试记录\n3—5 分钟演示录屏\n\n说明已实现、模拟\n与未实现的内容',m.YELLOW)])

    s=page(prs,52,'三组实际测试的要求','每组记录：初始条件、输入、预期、实际结果、问题与修正','第八节',('四、最终提交什么，以及分几步完成',),form='table')
    tab(s,[['测试类型','要看到的行为'],['正常完成','信息齐全，核心流程走到明确结束状态'],['需求变化','用户补充或修改信息，后续处理同步更新'],['无法完成','查询失败或条件不满足，说明状态与下一步']],[3.4,8.12],row_h=1.03)
    m.txt(s,.96,6.5,11.25,.4,'最终版提交实际运行结果；预期答案与“待验证”不能替代测试。',17,m.CORAL,True)

    s=page(prs,53,'初版、点评与最终版','分阶段完成，最终需要交付可运行项目','第八节',('四、最终提交什么，以及分几步完成','五、下次课怎么点评'),form='timeline')
    flow(s,[('点评课前','项目方案与 Coze 初版\n已有测试结果\n未搭通则带具体卡点'),('作业点评','检查业务与流程\n查看实际结果\n定位优先修改项'),('点评之后','修正并跑通核心流程\n补齐三组实际测试\n提交演示录屏')],y=2.65)
    m.txt(s,.96,6.0,11.25,.8,'老师提供数据、规则和搭建指引，结合卡点答疑。\n提交时间与渠道另行通知。',17,m.MUTE)

    s=page(prs,54,'个人项目选题练习','在聊天区写下你的项目方向','第八节',('五、下次课怎么点评','六、整课收尾'),form='statement')
    m.txt(s,.96,2.22,11.3,1.5,'我想为哪类用户，解决什么具体问题；\n这次先把流程做到哪一步。',30,m.TEXT,True)
    m.txt(s,.96,4.12,11.2,1.15,'例：为购买家电的用户提供维修服务，\n完成从故障描述到维修需求登记的流程。',22,m.TEAL,True)
    m.txt(s,.96,6.0,11.2,.66,'课后：项目方案、Coze 核心流程、实际测试与演示。',19,m.TEXT,True)

    m.txt(prs.slides[32],.96,6.25,11.3,.55,'当前没有更便宜且支持主动降噪的款式，先询问是否调整条件。',17,m.CORAL,True)
    for sh in prs.slides[37].shapes:
        if sh.has_table:
            sh.table.cell(0,1).text_frame.paragraphs[0].runs[0].text='业务举例'
    refine_text(prs)
    assert len(prs.slides)==54
    dk.declare_delivery(OUT,'presented')
    dk.lint_layout(prs,strict=True)
    prs.save(OUT)
    (ROOT/'.append-ledger.json').write_text(json.dumps({'source':str(SOURCE),'base':str(OLD),'base_sha256':hashlib.sha256(OLD.read_bytes()).hexdigest(),'slides':LEDGER},ensure_ascii=False,indent=2),encoding='utf-8')
    print(OUT)

if __name__=='__main__':
    build()
