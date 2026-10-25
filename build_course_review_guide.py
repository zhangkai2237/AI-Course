from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = "/Users/keivn/Project/AI-Course/AI_Agent项目作业课堂讲评手册.docx"

FONT_CN = "Arial Unicode MS"
FONT_EN = "Arial Unicode MS"
BLACK = "000000"
BLUE = "1F4E78"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
BORDER = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, bold=None, color=BLACK, italic=None):
    run.font.name = FONT_CN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_CN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_CN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_CN)
    r_fonts.set(qn("w:hAnsi"), FONT_CN)
    r_fonts.set(qn("w:eastAsia"), FONT_CN)
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_para(doc, text="", bold_lead=None, style=None, space_after=5, left=0, first_line=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.left_indent = Cm(left)
    p.paragraph_format.first_line_indent = Cm(first_line)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.18
    for r in p.runs:
        set_run_font(r)
    if not p.runs:
        r = p.add_run(text)
        set_run_font(r)
    else:
        p.runs[0].text = text
    return p


def add_number(doc, text, index):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    set_run_font(p.add_run(f"{index}. {text}"))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.5, repeat_header=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    if repeat_header:
        set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(text))
        set_run_font(r, size=font_size, bold=True, color=WHITE)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[idx])
            if ridx % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    r1 = p.add_run(label)
    set_run_font(r1, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2)
    return p


def page_break(doc):
    doc.add_page_break()


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_CN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_CN)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.25

    title = doc.styles["Title"]
    title.font.name = FONT_CN
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT_CN)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_CN)
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(12)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    sizes = {1: 18, 2: 14, 3: 11.5}
    spaces = {1: (14, 8), 2: (11, 5), 3: (8, 3)}
    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_CN)
        style.font.size = Pt(sizes[level])
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(spaces[level][0])
        style.paragraph_format.space_after = Pt(spaces[level][1])
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_CN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_CN)
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(BLACK)

    # Header and footer
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("AI Agent 项目作业课堂讲评手册")
    set_run_font(hr, size=8.5, color="666666")

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(0)
    fr = fp.add_run("第 ")
    set_run_font(fr, size=8.5, color="666666")
    add_page_field(fp)
    fr2 = fp.add_run(" 页")
    set_run_font(fr2, size=8.5, color="666666")


def add_project_header(doc, index, title, url, category, one_liner):
    page_break(doc)
    add_heading(doc, f"项目 {index}  {title}", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    add_hyperlink(p, "飞书项目链接", url)
    add_label_paragraph(doc, "项目类型  ", category)
    add_label_paragraph(doc, "一句话理解  ", one_liner)


def add_project_sections(doc, data):
    add_heading(doc, "项目是做什么的", 2)
    add_para(doc, data["overview"])
    add_heading(doc, "典型使用方式", 3)
    for index, item in enumerate(data["flow"], 1):
        add_number(doc, item, index)
    add_heading(doc, "AI 负责什么", 3)
    for item in data["ai"]:
        add_bullet(doc, item)
    add_heading(doc, "AI 不负责什么", 3)
    for item in data["not_ai"]:
        add_bullet(doc, item)

    add_heading(doc, "课堂点评", 2)
    add_label_paragraph(doc, "总体判断  ", data["judgement"])
    add_heading(doc, "值得肯定", 3)
    for item in data["strengths"]:
        add_bullet(doc, item)
    add_heading(doc, "需要重点讨论", 3)
    for item in data["issues"]:
        add_bullet(doc, item)
    add_heading(doc, "建议的下一步", 3)
    for item in data["next"]:
        add_bullet(doc, item)
    add_heading(doc, "课堂上可以直接说", 3)
    add_para(doc, data["script"])
    add_heading(doc, "建议追问", 3)
    for item in data["questions"]:
        add_bullet(doc, item)
    add_heading(doc, "评价时重点观察", 3)
    add_table(doc, ["观察项", "判断要点"], data["watch"], widths=[1.65, 5.2], font_size=9.3)


doc = Document()
configure_document(doc)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(95)
r = p.add_run("AI Agent 项目作业课堂讲评手册")
set_run_font(r, size=27, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(20)
r = p.add_run("学生汇报结构  教师点评方法  六个项目讲评参考")
set_run_font(r, size=13, color="444444")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(170)
r = p.add_run("适用场景  第三课项目作业课堂讲评")
set_run_font(r, size=10.5, color="666666")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("编制日期  2026 年 10 月")
set_run_font(r, size=10.5, color="666666")

page_break(doc)
add_heading(doc, "使用说明", 1)
add_para(doc, "本手册用于组织六个 AI Agent 产品作业的课堂汇报与讲评。建议先让学生用统一结构介绍项目，再由教师围绕产品价值、MVP 范围、Agent 必要性、实现路径和验证证据进行点评。这样能够避免课堂时间被完整 PRD 复述占用，也能够检验学生是否真正理解自己的产品决策。")
add_para(doc, "本轮作业普遍具备较完整的 PRD 结构。课堂点评应明确区分两类质量：一类是设计完整度，即流程、权限、安全、异常和评测是否覆盖；另一类是产品可信度，即痛点是否有基线、首期是否足够小、数据与工具是否真实存在、指标是否可归因、是否已经通过原型或案例验证。文档写得全面并不等于产品已经成立。")

add_heading(doc, "手册内容", 2)
for item in [
    "课堂组织建议与时间安排",
    "学生统一介绍结构和汇报要求",
    "教师点评框架与评分参考",
    "六个项目的基本认知和详细点评",
    "项目横向比较与课程总结",
    "课堂记录表和课前项目介绍卡",
]:
    add_bullet(doc, item)

add_heading(doc, "课堂使用原则", 2)
for item in [
    "不要求学生逐章介绍 PRD，只讲核心用户、核心任务、核心流程和验证计划。",
    "每个项目先肯定产品决策，再指出一个最关键问题，避免将点评变成零散挑错。",
    "优先追问事实、基线、数据来源和取舍，不只讨论模型、提示词或架构名词。",
    "所有数值先确认是规划目标还是已有结果，不把目标值当成验证结果。",
    "把安全与权限作为上线门槛，把业务指标作为价值验证，两者不能互相抵消。",
]:
    add_bullet(doc, item)

add_heading(doc, "六个项目快速定位", 2)
add_table(
    doc,
    ["项目", "主要用户", "解决的问题", "本质类型"],
    [
        ["电子流助手", "企业员工和审批人", "填表 提交 审核内部流程", "企业流程助手"],
        ["教育课程助手", "教育平台人工客服", "课程筛选 比较 推荐解释", "内部决策支持助手"],
        ["撮合平台助手", "中英文站内部员工", "查规则 查状态 找资源 提问题", "内部业务助手"],
        ["Nexa", "海外消费者和人工客服", "客服问答 实时查询 转人工", "消费者客服和工作台"],
        ["算量问答助手", "算量软件用户", "报错 操作路径 术语解释", "软件使用支持助手"],
        ["物联网售前工作台", "销售 技术支持 方案人员", "客户交流准备 答疑 纪要 跟进", "售前协作助手"],
    ],
    widths=[1.35, 1.55, 2.5, 1.45],
    font_size=8.8,
)

page_break(doc)
add_heading(doc, "课堂组织建议", 1)
add_heading(doc, "为什么需要学生先介绍", 2)
add_para(doc, "建议每个项目都由学生先进行简短介绍。学生能否在有限时间内说明用户、场景、取舍和验证方法，本身就是产品能力的一部分。统一汇报结构还能够减少因文档长度、排版和模板差异带来的评价偏差。")
add_para(doc, "学生介绍的目的不是完整复述作业，而是证明自己理解项目。尤其要观察学生能否回答为什么需要 Agent、首期明确不做什么，以及最可能不成立的假设是什么。")

add_heading(doc, "九十分钟课堂建议", 2)
add_table(
    doc,
    ["环节", "每组时间", "要求"],
    [
        ["学生介绍", "4 分钟", "按照统一五页结构 不得逐章复述 PRD"],
        ["教师点评", "4 分钟", "一句判断 两个亮点 一个关键问题 一个行动建议"],
        ["提问讨论", "3 分钟", "教师优先 同学补充 聚焦假设与验证"],
        ["切换", "1 分钟", "记录本组一句结论"],
        ["六组小计", "72 分钟", "每组 12 分钟"],
        ["横向总结", "15 分钟", "比较六种 Agent 场景与共性问题"],
    ],
    widths=[1.4, 1.25, 4.6],
)

add_heading(doc, "六十分钟课堂建议", 2)
add_table(
    doc,
    ["环节", "每组时间", "说明"],
    [
        ["学生介绍", "3 分钟", "每页约 30 至 40 秒"],
        ["教师点评", "4 分钟", "只讲最重要的两个优点和一个问题"],
        ["提问", "2 分钟", "保留一个教师追问和一个同学追问"],
        ["六组小计", "54 分钟", "最后 6 分钟做总结"],
    ],
    widths=[1.4, 1.25, 4.6],
)

add_heading(doc, "教师点评的推荐顺序", 2)
for index, item in enumerate([
    "先复述你对项目的理解，确认双方讨论的是同一个产品。",
    "指出两个值得保留的产品决策，不只表扬文档完整。",
    "指出一个最影响产品成立的关键问题，不同时罗列十个小问题。",
    "给出一个下一版本可以执行的改进动作。",
    "提出一个必须由学生现场回答的问题。",
], 1):
    add_number(doc, item, index)

page_break(doc)
add_heading(doc, "学生统一介绍结构", 1)
add_para(doc, "建议要求每组最多五页，汇报时间三至四分钟。页面内容可以使用流程图、原型和简表，但不得将完整 PRD 直接缩小后放入幻灯片。")

slides = [
    ["第一页", "项目定位", "用户是谁 使用场景是什么 当前怎么解决 最大痛点是什么 一句话定位是什么"],
    ["第二页", "为什么需要 Agent", "普通搜索 表单 规则引擎或 FAQ 为什么不够 Agent 做什么 人和确定性系统做什么"],
    ["第三页", "一个核心流程", "从用户输入到结果输出 标出知识库 Tool 数据 权限 人工确认和失败降级"],
    ["第四页", "MVP 范围", "首期必须做三至五项 首期明确不做什么 如果资源减半保留什么"],
    ["第五页", "验证计划", "北极星指标 当前基线 样本和测试方法 最大不确定假设 下一步验证"],
]
add_table(doc, ["页码", "主题", "必须回答"], slides, widths=[0.85, 1.45, 5.0], font_size=9.2)

add_heading(doc, "第一页项目定位", 2)
add_para(doc, "推荐学生用下面的句式开场：")
add_para(doc, "面向____用户，在____场景下，当前主要依靠____解决，导致____问题。我们的产品通过____，帮助用户实现____。", first_line=0.4)
add_para(doc, "教师重点观察学生是否描述了具体用户和触发场景，而不是只说面向企业或提升效率。")

add_heading(doc, "第二页为什么需要 Agent", 2)
add_para(doc, "学生需要说明任务是否包含开放式自然语言、多轮澄清、动态路由、外部工具调用或复杂上下文。如果输入高度结构化、步骤完全固定，就应解释为什么不用普通表单、搜索或规则系统。")
add_table(
    doc,
    ["由谁负责", "适合承担的内容"],
    [
        ["Agent", "理解表达 抽取信息 拆解任务 选择受控工具 组织有依据的结果"],
        ["确定性系统", "身份 权限 数据范围 业务规则 真实写入和最终状态"],
        ["人工", "高风险判断 对外承诺 审批 决策和异常接管"],
    ],
    widths=[1.45, 5.8],
)

add_heading(doc, "第三页一个核心流程", 2)
add_para(doc, "只讲一个最典型任务。建议统一画成以下结构：")
add_para(doc, "用户输入 → 意图与信息识别 → 补充缺失信息 → 调用知识库或业务 Tool → 结果验证 → 用户确认或转人工 → 最终结果")
add_para(doc, "必须标出数据来源、工具调用、人工介入点和失败出口。不能只画模型输入和模型输出。")

add_heading(doc, "第四页 MVP 范围", 2)
add_para(doc, "首期范围必须能够在有限时间内验证一个关键假设。建议控制在三至五项核心能力，并同时列出明确不做的能力。学生还需要回答资源减少一半时的取舍。")

add_heading(doc, "第五页验证计划", 2)
for item in [
    "北极星指标是什么，为什么它能代表用户价值。",
    "当前基线是多少，若没有基线准备如何采集。",
    "计划使用多少真实用户、历史案例或测试样本。",
    "哪些指标是上线硬门槛，哪些只是观察指标。",
    "项目当前最可能被证明不成立的假设是什么。",
]:
    add_bullet(doc, item)

add_heading(doc, "统一项目介绍卡", 2)
add_table(
    doc,
    ["字段", "学生填写内容"],
    [
        ["目标用户", "谁在什么工作情境下使用"],
        ["核心痛点", "当前替代方案及可量化问题"],
        ["一句话定位", "产品帮助谁完成什么"],
        ["Agent 必要性", "为什么普通系统不足"],
        ["核心任务", "首期最重要的一个任务"],
        ["关键数据和 Tool", "真实来源 接口 可用状态和责任人"],
        ["首期边界", "明确不做的能力"],
        ["北极星指标", "定义 分母 分子和基线"],
        ["最大假设", "最可能使项目不成立的条件"],
        ["下一步验证", "样本 用户 周期和判定标准"],
    ],
    widths=[1.55, 5.7],
)

page_break(doc)
add_heading(doc, "教师点评框架", 1)
add_heading(doc, "一句判断 两个亮点 一个问题 一个动作 一个追问", 2)
add_para(doc, "建议每次点评都保持完整闭环。先判断项目属于什么类型以及当前成熟度，再肯定产品决策，随后指出最关键的风险，最后给出可执行的下一步。")
add_table(
    doc,
    ["步骤", "目的", "示例"],
    [
        ["一句判断", "建立共同认知", "这是一个范围清楚但数据准备尚未验证的内部推荐助手"],
        ["两个亮点", "指出值得保留的设计", "用户边界清楚 非 Agent 基线合理"],
        ["一个问题", "指出最影响产品成立的风险", "课程事实更新责任尚未明确"],
        ["一个动作", "让学生知道下一版怎么改", "用历史案例做人工 规则 Agent 三组盲测"],
        ["一个追问", "检验学生是否真正理解", "如果数据过期 谁发现 谁确认 谁更新"],
    ],
    widths=[1.2, 1.7, 4.35],
    font_size=9.2,
)

add_heading(doc, "评分参考", 2)
add_table(
    doc,
    ["维度", "分值", "判断重点"],
    [
        ["问题与价值", "15", "痛点是否真实 高频 且值得解决"],
        ["场景与 MVP", "20", "用户是否明确 首期是否收敛 边界是否清楚"],
        ["Agent 必要性", "15", "是否真的需要理解 规划 多轮和工具调用"],
        ["流程与可实现性", "20", "数据 Tool 权限 状态 异常是否能够落地"],
        ["安全与人机协作", "10", "高风险动作 人工确认 越权与降级"],
        ["指标与验证", "15", "指标是否可计算 可归因 是否有基线和样本"],
        ["表达与取舍", "5", "是否能在有限时间讲清核心决策"],
    ],
    widths=[1.55, 0.75, 4.95],
)

add_heading(doc, "常见点评误区", 2)
for item in [
    "把章节数量和文字长度当成产品成熟度。",
    "看到具体阈值就默认其有数据依据。",
    "只讨论模型和提示词，不追问数据、接口、权限与运营责任。",
    "只看正常流程，不讨论零结果、接口失败、人工不可用和上下文断裂。",
    "用成交率、端到端时长等受多因素影响的指标直接评价 Agent。",
    "一次指出太多问题，导致学生不知道下一版最该改什么。",
]:
    add_bullet(doc, item)


projects = [
    {
        "index": 1,
        "title": "电子流流程 AI 辅助助手",
        "url": "https://qi9bp2sfmq4.feishu.cn/wiki/SwZDwzGvLiZFJake9Skc5GMbnRf",
        "category": "企业流程助手",
        "one_liner": "把员工的一句话转成内部流程表单，并在不替代人工决策的前提下辅助审批。",
        "overview": "该项目服务企业内部报修、技术变更报审和工单申请等电子流。员工不必先找到正确表单，可以直接描述问题，由 Agent 判断流程类型、提取字段、补充缺失信息、检查业务规则和重复单，再由员工确认提交。审批阶段，Agent 为审批人生成单据摘要、风险提示和有来源的意见草稿。",
        "flow": [
            "员工用自然语言描述报修或申请需求。",
            "Agent 识别流程类型并抽取设备 位置 现象 紧急程度等字段。",
            "Agent 追问关键缺失信息并完成规则校验。",
            "系统检测相似单并提示关联 合并或继续提交。",
            "Agent 推荐审批路径 员工确认后提交。",
            "审批人按需生成审核辅助信息并作出最终决定。",
        ],
        "ai": [
            "识别报修 报审 工单等流程意图。",
            "将自然语言转换为可编辑表单字段。",
            "检查缺项 超范围和可能重复的单据。",
            "生成带来源的单据摘要和审核意见草稿。",
        ],
        "not_ai": [
            "不自动提交表单。",
            "不自动审批或替审批人决定通过和驳回。",
            "没有制度和规则来源时不输出结论性审核意见。",
            "不在首期建设复杂流程自定义编排。",
        ],
        "judgement": "流程闭环和责任边界比较成熟，但首期覆盖流程和能力较多，北极星指标的归因还需收紧。",
        "strengths": [
            "产品定位合理 将 AI 作为既有电子流的辅助层 而不是重新建设流程引擎。",
            "关键动作必须人工确认 AI 不自动提交 不自动审批 责任边界清楚。",
            "从发起 补全 校验 防重 路径推荐到审批辅助和归档形成完整闭环。",
            "填单时长 首次填写成功率 预填覆盖率和意见采纳率等指标较具体。",
        ],
        "issues": [
            "一期同时覆盖三类电子流和多项能力 更像完整产品版本而不是最小验证。",
            "端到端处理时长受到审批人响应和组织 SLA 影响 难以全部归因于 AI。",
            "意图置信度和防重相似度等阈值需要说明验证样本与确定依据。",
            "原型 同义词词典和审核意见来源库仍是落地前关键依赖。",
        ],
        "next": [
            "首期只选择一个高频报修流程。",
            "优先验证自然语言转表单 字段校验和相似单检测。",
            "使用历史工单和真实用户建立人工流程基线。",
            "将端到端指标拆成 AI 可直接影响的填写时间 错误率和一次成功率。",
        ],
        "script": "这份方案最值得肯定的是没有让 AI 越过业务责任边界，而是把 AI 放在理解 补全 校验和建议的位置。但现在一期承担的任务较多。我建议先选择一个高频电子流形成最小闭环，并把端到端时长拆成 AI 可以直接影响的指标，否则审批 SLA 的变化会干扰结果。",
        "questions": [
            "如果首期只能保留三个能力 会保留哪三个。",
            "端到端时长下降中如何排除审批 SLA 的影响。",
            "防重相似度阈值通过什么样本确定。",
            "首期准备使用多少历史电子流进行验证。",
        ],
        "watch": [
            ["MVP 取舍", "是否能从三类流程中选择一个最高价值流程"],
            ["指标归因", "是否将组织效率和 AI 效果分开统计"],
            ["数据准备", "规则库 历史单 设备台账是否真实可用"],
            ["安全边界", "关键动作是否始终由人工确认"],
        ],
    },
    {
        "index": 2,
        "title": "教育电商售前知识与课程推荐助手",
        "url": "https://qi9bp2sfmq4.feishu.cn/wiki/UiVZwPdLCiNpC8khpF2cY2tTndh",
        "category": "内部决策支持助手",
        "one_liner": "客服输入学生需求，Agent 帮助筛选和比较雅思课程，并解释推荐依据。",
        "overview": "该项目供教育平台内部人工客服使用。客服在与学生或家长沟通后，将基础分数、目标分数、考试时间、预算、地区、授课方式和班型偏好录入系统。Agent 从平台课程中筛选候选课程，说明匹配点、不匹配点、资料缺口和沟通建议，同时查询售前相关的平台规则。",
        "flow": [
            "客服录入已经确认的学生需求。",
            "Agent 检查完成首轮推荐所需的必要字段。",
            "系统先按硬条件筛选 再按软条件排序候选课程。",
            "系统补齐课程与机构事实并检查有效期。",
            "Agent 生成推荐理由 候选差异 沟通建议和待确认项。",
            "客服判断是否采用并负责最终对外表达。",
        ],
        "ai": [
            "理解客服录入的自然语言需求并标准化字段。",
            "在结构化和半结构化课程资料中检索 比较和排序。",
            "解释课程与当前学生需求的匹配关系。",
            "生成供客服修改的沟通建议和建议追问。",
        ],
        "not_ai": [
            "首期不直接面向学生和家长。",
            "不自动下单 收款 退款或开通会员。",
            "不保证提分 考试结果 班期或机构服务效果。",
            "不在缺少来源或资料冲突时给出确定性结论。",
        ],
        "judgement": "六个项目中 MVP 范围最收敛，但产品成败主要取决于课程数据质量和推荐验证方法。",
        "strengths": [
            "目标用户清楚 首期只面向内部客服 不直接开放 C 端。",
            "首期只验证雅思课程 避免一开始覆盖全部教育品类。",
            "明确标准课程表和推荐手册是非 Agent 的最低成本基线。",
            "有效推荐任务完成率比搜索量或成交率更适合作为低样本 MVP 指标。",
        ],
        "issues": [
            "价格 班期 优惠和课程状态容易变化 数据过期会直接破坏推荐可信度。",
            "课程事实责任人 数据有效期和更新流程仍需冻结。",
            "需要证明 Agent 相比结构化筛选器和规则系统的增量价值。",
            "推荐理由可用率和客服反馈带有主观性 需要统一判定标准和盲测。",
        ],
        "next": [
            "冻结二十至三十门真实雅思课程和字段有效期。",
            "选择二十至五十个历史咨询案例进行盲测。",
            "比较人工查表 规则筛选和 Agent 推荐三种方式。",
            "重点统计耗时 候选可用率 事实错误率和解释可用率。",
        ],
        "script": "这个项目最好的地方是范围收得比较稳 内部客服 雅思课程 推荐辅助而不是自动销售。下一步不应继续扩充意图标签 而应冻结真实课程数据 历史咨询案例和审核责任人，并用对照实验回答 Agent 相比普通筛选器增加了什么价值。",
        "questions": [
            "当前有多少门课程的数据可以直接用于推荐。",
            "价格 班期和优惠分别多长时间算过期。",
            "推荐正确的标准由谁制定。",
            "如何避免对不同机构产生隐性推荐偏差。",
        ],
        "watch": [
            ["范围收敛", "是否坚持内部客服和雅思单品类"],
            ["数据治理", "课程事实的 Owner 有效期和审核机制"],
            ["Agent 增量", "是否优于普通筛选和固定规则"],
            ["验证方法", "是否使用真实案例和盲评"],
        ],
    },
    {
        "index": 3,
        "title": "撮合平台智能 AI 业务助手",
        "url": "https://qi9bp2sfmq4.feishu.cn/wiki/KP3aw1DL0iDwRlk7scScIyEEnae",
        "category": "企业内部业务助手",
        "one_liner": "为中英文站内部员工提供查规则 查状态 找资源和整理业务问题的统一入口。",
        "overview": "该项目服务撮合平台中文站和英文站的内部 BD 推广 运营 客服 产品和技术人员。员工可以用自然语言查询平台规则 查询权限范围内的业务状态 获取站内资源推荐，或将异常整理成问题卡片。Agent 负责理解和路由，站点隔离 权限判断和真实业务数据由后端及业务 Tool 控制。",
        "flow": [
            "内部员工在所属站点工作台输入问题。",
            "系统取得身份 岗位 站点和数据范围。",
            "Agent 将请求识别为知识问答 私有查询 资源推荐或问题整理。",
            "系统按站点和权限调用知识库 搜索或业务 Tool。",
            "结果经过对象 字段 来源和 Allowed Facts 校验。",
            "系统返回答案 推荐卡片 问题卡片或受控阻断说明。",
        ],
        "ai": [
            "识别开放式自然语言和业务意图。",
            "抽取业务对象 条件和上下文指代。",
            "将复合请求拆成受控任务并选择预设工具。",
            "基于允许事实组织可追溯结果。",
        ],
        "not_ai": [
            "不自行决定身份 岗位 站点 数据范围和字段权限。",
            "不跨站查询或共享私有数据。",
            "首期不修改订单 会员 资金 CRM 等业务数据。",
            "Tool 未返回或未通过事实校验的信息不由模型补全。",
        ],
        "judgement": "工程治理和安全设计最完整，但当前更像完整目标架构，尚未充分收敛为可快速验证的产品 MVP。",
        "strengths": [
            "双站知识 数据 权限 会话 日志和模型链路相互隔离。",
            "明确 Agent 负责理解编排 确定性系统负责权限和执行。",
            "系统处理结果与业务目标结果分开统计 避免把正确拒绝算成用户成功。",
            "Allowed Facts Fail Closed Tool 白名单和字段级权限接近生产要求。",
        ],
        "issues": [
            "首期覆盖两个站点 多类岗位和四类任务 核心用户价值不够突出。",
            "效率 问题量和重复咨询比例基线尚未形成。",
            "真实接口 字段 Owner 和期限仍需冻结。",
            "治理体系很完整 但首屏体验和一个具体任务的端到端旅程不够突出。",
        ],
        "next": [
            "先选择一个站点 一个岗位和两个高频意图。",
            "只接入少量真实知识和业务 Tool 做纵向切片。",
            "制作从输入 权限校验 Tool 到结果呈现的可演示旅程。",
            "在扩展双站前完成真实接口和待确认事项冻结。",
        ],
        "script": "这份方案的安全治理水平很高 特别是站点隔离 字段权限和事实准入。但产品设计当前有些被治理框架淹没。建议先找出一个用户 一条高频任务和一条完整链路，用真实接口和真实数据证明价值，再扩展到双站和更多岗位。",
        "questions": [
            "当前最高频的角色和任务分别是什么。",
            "为什么两个站点必须同时进入一期。",
            "如果四周内交付 最小版本包含什么。",
            "哪些接口已经存在 哪些仍是设计假设。",
        ],
        "watch": [
            ["产品聚焦", "能否从治理框架回到一个具体用户任务"],
            ["接口真实性", "真实 Tool 字段 Owner 和可用状态"],
            ["双站必要性", "是否必须同时验证两个站点"],
            ["安全治理", "站点 权限 字段和审计是否保持确定性控制"],
        ],
    },
    {
        "index": 4,
        "title": "Nexa 海外独立站 AI 客服与人工协同系统",
        "url": "https://qi9bp2sfmq4.feishu.cn/wiki/EeH0wkAGvi5pEOkuzOKcRvWenNf",
        "category": "消费者客服和人工工作台",
        "one_liner": "海外消费者先由 AI 接待，复杂问题带着会话和业务上下文转给人工客服。",
        "overview": "该项目面向美国市场泳衣独立站，包含消费者侧 AI 客服和人工客服工作台。消费者可以咨询商品 尺码 推荐 FAQ 订单 物流和基础售后资格。需要实时数据时系统调用库存 订单和物流接口。AI 无法解决或用户要求人工时，系统生成 Handoff Summary 并进入客服队列，人工接管后 AI 停止自动对外回复。",
        "flow": [
            "消费者在网站聊天入口提出问题。",
            "Agent 区分公开知识 商品推荐和私密业务查询。",
            "订单和物流问题触发登录和本人归属校验。",
            "系统调用知识库 Inventory API 或 Order API。",
            "无法确认或需要高权限处理时生成唯一 Handoff。",
            "人工客服接单 查看摘要和上下文 使用 Copilot 建议并自行发送回复。",
        ],
        "ai": [
            "理解商品 尺码 订单 物流和售后等自然语言问题。",
            "在公开知识和实时 Tool 之间进行路由。",
            "生成转人工摘要和内部建议回复。",
            "识别越权 高风险投诉和需要人工判断的情况。",
        ],
        "not_ai": [
            "不查询其他消费者的订单和私密信息。",
            "不在 Tool 失败时编造订单 库存 物流或预计时间。",
            "不自动创建退款 换货或赔偿等高权限业务动作。",
            "人工接管后不继续自动向消费者回复。",
        ],
        "judgement": "功能规格和测试最接近开发交付，但消费者客服与人工工作台同时建设，所有能力列为 P0 导致范围过大。",
        "strengths": [
            "功能需求包含用户故事 输入 流程 权限 验收和埋点 可开发性较强。",
            "Human Handoff 包含摘要 唯一任务 队列状态和人工接管后的行为边界。",
            "人工工作台区分建议回复和真实发送 业务数据保持只读。",
            "二十个测试用例覆盖正常 异常 越权 接口超时 注入和上下文恢复。",
        ],
        "issues": [
            "商品 尺码 推荐 FAQ 订单 物流 售后 Handoff Workspace 和 Copilot 全部为 P0。",
            "消费者客服和人工工作台本质上是两个相互依赖的产品。",
            "自助有效解决率的分母排除项较多 需要防止指标被优化。",
            "人工在线时段 接单 SLA 接口稳定性和英语质量是重要运营依赖。",
        ],
        "next": [
            "P0 只保留 FAQ 订单 物流和 Human Handoff。",
            "商品咨询和基础售后放入 P1 尺码与个性化推荐后置。",
            "分别验证静态知识 实时查询和转人工三条链路。",
            "补充人工客服容量 排队时间和无人在线策略。",
        ],
        "script": "这份方案的开发规格 异常处理和测试用例都很扎实，特别是人工接管后的状态转换。但目前 P0 范围过大，相当于同时建设用户端客服和客服工作台。建议先跑通 FAQ 订单物流和转人工三条主链路，再加入尺码和个性化推荐。",
        "questions": [
            "人工客服在线时段和接单 SLA 是什么。",
            "无人工在线时用户会看到什么。",
            "订单 物流和库存接口是否真实可用。",
            "自助解决率的分母如何防止被人为缩小。",
            "英文回复质量和消费者隐私如何验证。",
        ],
        "watch": [
            ["P0 范围", "是否拆分消费者侧和人工工作台的交付优先级"],
            ["实时 Tool", "订单 物流 库存接口是否真实且稳定"],
            ["人工运营", "队列 SLA 在线时段和无人接单策略"],
            ["测试质量", "是否覆盖越权 接口失败和状态恢复"],
        ],
    },
    {
        "index": 5,
        "title": "算量问答助手",
        "url": "https://qi9bp2sfmq4.feishu.cn/wiki/MCtnwAuuNihRvfk4B64cKhLQnig",
        "category": "软件使用支持助手",
        "one_liner": "嵌入算量软件的智能帮助中心，优先处理报错 操作路径和术语解释。",
        "overview": "该项目面向项目算量员 工程部管理人员和项目管理人员，解决算量软件使用中的报错 操作路径和术语理解问题。当前用户主要查操作手册 帮助文档 视频或在微信群中向实施 运营 测试和产品求助。Agent 希望在客户端内提供即时 可追溯的自助问答，无法处理时再转人工。",
        "flow": [
            "用户在算量软件中提出报错 操作或术语问题。",
            "系统获取软件版本 当前页面 用户角色和项目上下文。",
            "Agent 识别问题类型并检索对应知识。",
            "系统返回适用版本的操作步骤和来源。",
            "用户反馈是否有帮助。",
            "仍无法解决时生成摘要并转给人工支持。",
        ],
        "ai": [
            "理解用户对报错 操作路径和术语的自然语言描述。",
            "检索操作手册 帮助文档和经过审核的历史问题。",
            "根据当前页面和软件版本组织操作步骤。",
            "整理转人工时的问题背景 已尝试步骤和日志线索。",
        ],
        "not_ai": [
            "不在缺少版本上下文时猜测按钮和操作路径。",
            "不把高风险工程结论当成普通软件操作问题。",
            "无法判断报错原因时不反复让用户无目的尝试。",
            "当前定位不是自动完成工程算量。",
        ],
        "judgement": "问题真实，但定位和核心流程尚未完全写实，当前更像算量软件使用帮助助手而不是专业算量 Agent。",
        "strengths": [
            "微信群重复答疑 跨团队流转和非工作时间值班是具体且常见的痛点。",
            "提出回答需要与知识来源一致 具备可追溯意识。",
            "已经设置解决率 准确率 转人工率和响应时间等指标。",
        ],
        "issues": [
            "项目名称容易让人理解为能直接参与工程算量 与当前能力存在预期偏差。",
            "场景确定 决策流程和核心工具契约的实质内容仍偏少。",
            "有效问题解决率和满意度都依赖有帮助反馈 指标区分度不足。",
            "操作和报错场景准确率目标仅为百分之八十 可能不足以上线。",
            "软件版本 错误码 客户端日志和知识更新责任尚需明确。",
        ],
        "next": [
            "首期只做错误码解释和操作路径导航两个任务。",
            "自动携带软件版本 当前页面 用户角色和项目类型。",
            "建立至少五十条黄金测试集并按版本维护。",
            "回答必须展示知识来源和适用版本 无依据时转人工。",
        ],
        "script": "项目的问题是真实的，但需要先把名字和能力边界说清楚。它目前解决的是算量软件的使用支持，而不是直接替用户算量。建议从报错解释和操作路径两个高频任务切入，把软件版本 错误码和知识来源纳入回答。",
        "questions": [
            "用户最常见的前二十个问题是什么。",
            "系统能否自动取得软件版本 当前页面和错误码。",
            "不同版本操作路径冲突时如何回答。",
            "知识由谁更新 多久更新一次。",
            "准确率达到多少才允许正式上线。",
        ],
        "watch": [
            ["定位一致性", "名称是否与实际能力匹配"],
            ["上下文", "版本 页面 项目和错误码能否自动获得"],
            ["知识治理", "文档版本 来源和更新责任"],
            ["完成度", "关键场景和决策流程是否从标题变成可执行内容"],
        ],
    },
    {
        "index": 6,
        "title": "物联网集成项目售前协作工作台",
        "url": "https://qi9bp2sfmq4.feishu.cn/wiki/WeGYw8kXUiJNuCkc55Uc368gnVe",
        "category": "B2B 售前协作助手",
        "one_liner": "围绕一次客户技术交流，帮助销售准备材料 理解问题 整理纪要并推动后续行动。",
        "overview": "该项目面向物联网集成公司的销售 技术支持和解决方案人员。产品以项目内三栏工作区承载客户背景 交流资料 讲解准备 预备问答 会议纪要和待办事项，右侧常驻 AI 助手。长期希望覆盖信息采集 方案制作 报价和移交，但当前 PRD 主要定义技术交流工作台。",
        "flow": [
            "会前汇总客户背景并明确交流目标。",
            "检索适用的历史资料并检查版本和对客范围。",
            "生成讲解重点 预备问题和需要向客户确认的条件。",
            "会中按需查询技术资料并识别需要专业核实的问题。",
            "会后整理纪要 未决事项和候选待办。",
            "销售或专业人员确认后保存版本 创建待办或提交审核。",
        ],
        "ai": [
            "整理分散的客户背景和项目资料。",
            "检索并比较技术资料的版本 条件和适用范围。",
            "生成讲解草稿 预备问答和客户澄清问题。",
            "将真实会议记录整理为纪要 未决项和候选待办。",
        ],
        "not_ai": [
            "不把 AI 建议写成已发生的客户事实。",
            "不在缺少现场条件时承诺标准能力一定适用。",
            "不自动对外联系专业人员或提交审核。",
            "不把复制 采纳 保存版本 创建待办和审核视为同一动作。",
        ],
        "judgement": "业务洞察和界面原型最有场景感，但整体愿景和一期任务仍偏多，需要进一步收紧专业责任范围。",
        "strengths": [
            "九类售前痛点来自具体工作环节 不只是泛泛描述效率问题。",
            "左侧导航 中间业务内容 右侧 AI 助手的三栏原型易于理解。",
            "事实 销售理解 AI 建议 未知和待确认信息区分清楚。",
            "资料版本 适用条件 对客范围 专业审核和正式保存状态设计较成熟。",
        ],
        "issues": [
            "长期愿景覆盖五个业务模块 当前技术交流内部仍有九类任务。",
            "现场答疑 识图 内容生成 纪要 待办和审核同时进入首期 范围较宽。",
            "不同风险能力统一采用约百分之八十门槛 不够体现对客技术结论的高风险。",
            "会议现场流畅回答可能放大错误资料和过度承诺风险。",
        ],
        "next": [
            "首期只保留会前背景与目标 资料查找核实 会后纪要与待办。",
            "现场识图答疑 完整方案生成 报价和复杂写入后置。",
            "选择五至十场真实客户交流进行前后对照。",
            "对资料错误 技术结论错误和纪要结构问题设置不同门槛。",
        ],
        "script": "这份方案的行业理解和产品原型都比较强，能够看出是在真实售前流程上设计的。但一期内部仍包含九类任务。建议先抓住会前准备 资料核实和会后纪要三个高频闭环，暂时不要让 Agent 在会议现场承担过多专业判断。",
        "questions": [
            "九类任务中频率和价值最高的是哪三类。",
            "哪些技术资料已经完成版本化和权限治理。",
            "如何保证项目 A 的资料不会进入项目 B。",
            "现场回答错误和纪要错误是否应该使用不同门槛。",
            "是否已经选择真实客户交流进行原型测试。",
        ],
        "watch": [
            ["业务洞察", "痛点是否来自真实销售和技术支持流程"],
            ["首期取舍", "是否从九类任务收敛到三个闭环"],
            ["资料治理", "版本 权限 适用范围和审核状态"],
            ["风险分级", "对客技术结论是否使用更高门槛"],
        ],
    },
]

for project in projects:
    add_project_header(
        doc,
        project["index"],
        project["title"],
        project["url"],
        project["category"],
        project["one_liner"],
    )
    add_project_sections(doc, project)

page_break(doc)
add_heading(doc, "项目横向比较", 1)
add_para(doc, "六个项目并不是同一种 Agent。课堂上不宜只比较谁的文档更长，而应比较不同场景下的产品聚焦 数据依赖 工具风险和验证路径。")
add_table(
    doc,
    ["观察角度", "相对突出项目", "原因"],
    [
        ["MVP 范围最收敛", "教育课程助手", "内部客服 单一品类 明确非 Agent 基线"],
        ["功能规格和测试最完整", "Nexa", "功能需求 验收 埋点和二十个测试用例较完整"],
        ["安全与工程治理最成熟", "撮合平台助手", "站点隔离 字段权限 Allowed Facts 和 Fail Closed"],
        ["业务洞察和原型最强", "物联网售前工作台", "真实售前痛点和三栏工作台原型较清楚"],
        ["业务流程闭环最清楚", "电子流助手", "从自然语言发起到审批归档链路完整"],
        ["最需要补充核心设计", "算量问答助手", "定位 场景 工具契约和决策流程仍需写实"],
    ],
    widths=[1.85, 1.75, 3.65],
    font_size=9.1,
)

add_heading(doc, "六句话记住六个项目", 2)
for index, item in enumerate([
    "电子流助手把一句话变成内部流程表单并辅助审批。",
    "教育课程助手根据学生需求帮客服筛选和解释课程。",
    "撮合平台助手让内部员工统一查规则 查状态 找资源。",
    "Nexa 让 AI 先接待海外消费者 复杂问题带上下文转人工。",
    "算量问答助手解决算量软件的报错 操作和使用问题。",
    "物联网售前工作台帮助销售准备客户交流 查资料 做纪要和跟进。",
], 1):
    add_number(doc, item, index)

add_heading(doc, "全班共性问题", 2)
for item in [
    "多数项目的指标是规划目标 还不是运行结果。",
    "文档覆盖面普遍较全 但 MVP 仍有继续收敛空间。",
    "不少阈值很具体 但尚未说明样本和验证依据。",
    "Agent 价值需要与搜索 表单 规则引擎和人工流程进行对照。",
    "知识和 Tool 的 Owner 有效期 权限和失败策略是落地关键。",
    "生产化治理不能代替产品价值验证 产品价值也不能抵消安全红线。",
]:
    add_bullet(doc, item)

add_heading(doc, "课程总结话术", 2)
add_para(doc, "大家这次普遍已经能够写出结构完整的 Agent PRD。下一阶段要训练的不是继续增加章节，而是做取舍和验证。优秀的 Agent 产品方案不是能力最多 架构最复杂，而是能够清楚回答最重要的用户问题是什么 为什么必须使用 Agent 首期只验证什么 以及用什么证据证明它有效。")
add_para(doc, "在下一版作业中，建议每组都减少至少一项首期能力，补充一项真实基线，冻结一组真实数据或接口，并用一个可演示的核心旅程说明产品如何工作。")

page_break(doc)
add_heading(doc, "课堂讲评记录表", 1)
add_para(doc, "每个项目讲评时可复制使用以下记录结构。")
add_table(
    doc,
    ["记录项", "课堂记录"],
    [
        ["项目名称", ""],
        ["一句话理解", ""],
        ["最值得肯定的决策", ""],
        ["第二个亮点", ""],
        ["最关键的问题", ""],
        ["建议的下一步动作", ""],
        ["现场追问", ""],
        ["学生回答摘要", ""],
        ["最终课堂结论", ""],
    ],
    widths=[2.0, 5.25],
)

add_heading(doc, "教师快速检查清单", 2)
for item in [
    "我是否先确认了自己对项目的理解。",
    "我是否肯定了具体产品决策 而不是笼统表扬。",
    "我是否只聚焦一个最重要的问题。",
    "我是否区分规划目标与已有结果。",
    "我是否追问了数据 Tool Owner 基线或样本。",
    "我是否给出了学生下一版可以执行的动作。",
]:
    add_bullet(doc, "□ " + item)

add_heading(doc, "学生课前提交检查清单", 2)
for item in [
    "五页以内 三至四分钟内能够讲完。",
    "只介绍一个核心用户旅程。",
    "明确说明为什么需要 Agent。",
    "明确列出首期不做什么。",
    "至少说明一个真实数据源或 Tool。",
    "区分目标值和已经测得的结果。",
    "能够回答如果资源减少一半会砍掉什么。",
    "能够回答最可能不成立的假设是什么。",
]:
    add_bullet(doc, "□ " + item)

# Document properties
doc.core_properties.title = "AI Agent 项目作业课堂讲评手册"
doc.core_properties.subject = "第三课项目作业课堂介绍与点评参考"
doc.core_properties.author = "课程讲评资料"
doc.core_properties.keywords = "AI Agent PRD 课堂点评 项目作业"

doc.save(OUTPUT)
print(OUTPUT)
